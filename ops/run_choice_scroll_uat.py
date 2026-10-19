"""Exercise real native scrollbar dragging and report downloads on synthetic QA."""
import argparse
import io
import json
from pathlib import Path
import sys

from run_form_contract_uat import configure


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--directory', type=Path, required=True)
    parser.add_argument('--packages-dir', type=Path, required=True)
    parser.add_argument('--port', type=int, default=8840)
    args = parser.parse_args()
    root, folder = Path(__file__).resolve().parents[1], args.directory.resolve()
    if folder == root or root in folder.parents:
        parser.error('External synthetic runtime required')
    configure(root, folder)
    sys.path.insert(0, str(args.packages_dir.resolve()))
    from playwright.sync_api import sync_playwright
    from pptx import Presentation

    password = next(line.split(': ', 1)[1] for line in (folder/'DEMO_ACCOUNTS.txt').read_text(encoding='utf-8').splitlines() if line.startswith('Mật khẩu'))
    output = folder/'choice-scroll-uat'
    output.mkdir(exist_ok=True)
    report = {'synthetic_only':True, 'checks':[], 'failures':[], 'geometry':[]}

    def check(value, label):
        report['checks'].append(label)
        if not value:
            report['failures'].append(label)

    with sync_playwright() as pw:
        # Playwright hides native scrollbars by default, which would turn a
        # purported drag regression into a click on blank padding.
        browser = pw.chromium.launch(channel='msedge', headless=True,
                                     ignore_default_args=['--hide-scrollbars'])
        context = browser.new_context(service_workers='block', viewport={'width':1756,'height':1050})
        page = context.new_page()
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        base = f'http://127.0.0.1:{args.port}'
        page.goto(base+'/portal/login/')
        page.locator('[name=username]').fill('demo-hr')
        page.locator('[name=password]').fill(password)
        page.locator('main form button[type=submit]').click()
        page.wait_for_load_state('networkidle')
        for width in (1756,1024,390,360):
            page.set_viewport_size({'width':width,'height':1050 if width>600 else 844})
            for theme in ('light','dark'):
                page.goto(base+'/portal/tasks/create/', wait_until='networkidle')
                page.evaluate('(t)=>document.documentElement.dataset.theme=t', theme)
                for index in (0,1):
                    search = page.locator('.ui-choice-search').nth(index)
                    source = page.locator('#'+search.get_attribute('data-ui-choice-source'))
                    search.click()
                    popup = page.locator('.ui-choice-popover:popover-open')
                    prefix = f'{width}/{theme}/choice{index}'
                    bounds = popup.bounding_box()
                    geometry = popup.evaluate('p=>({top:p.scrollTop,client:p.clientHeight,scroll:p.scrollHeight,gutter:p.offsetWidth-p.clientWidth})')
                    check(geometry['gutter']>2, prefix+' visible native scrollbar gutter')
                    original = source.input_value()
                    thumb_y = bounds['y']+min(30,geometry['client']**2/geometry['scroll']/2)
                    gutter_x = bounds['x']+bounds['width']-geometry['gutter']/2
                    page.mouse.move(gutter_x,thumb_y)
                    page.mouse.down()
                    page.mouse.move(gutter_x,thumb_y+min(155,geometry['client']/2),steps=12)
                    page.mouse.up()
                    check(popup.count()==1, prefix+' scrollbar drag keeps popup open')
                    if not popup.count():
                        continue
                    after = popup.evaluate('p=>p.scrollTop')
                    check(after>geometry['top'], prefix+' native drag actually scrolls')
                    check(source.input_value()==original, prefix+' drag does not change submitted ID')
                    report['geometry'].append({'case':prefix,'before':geometry,'after_scroll_top':after})
                    # Full-page capture can resize the viewport and legitimately
                    # close an anchored popup. Capture the actual viewport.
                    page.screenshot(path=str(output/f'{width}-{theme}-choice{index}.png'))
                    popup.hover()
                    page.mouse.wheel(0,-180)
                    page.wait_for_timeout(100)
                    check(popup.evaluate('p=>p.scrollTop')<after, prefix+' wheel scroll remains usable')
                    option = popup.locator('[role=option]').last
                    chosen_text = option.text_content()
                    option.click()
                    check(page.locator('.ui-choice-popover:popover-open').count()==0, prefix+' choosing after drag closes once')
                    check(search.input_value()==chosen_text and source.input_value()!=original, prefix+' selected value and label match')
                    search.click()
                    page.keyboard.press('Escape')
                    check(search.get_attribute('aria-expanded')=='false' and search.evaluate('x=>x===document.activeElement'), prefix+' Escape restores focus without reopening')
                    search.click()
                    page.locator('h1').click()
                    check(search.get_attribute('aria-expanded')=='false', prefix+' outside click closes')
                    search.click()
                    page.keyboard.press('Tab')
                    check(search.get_attribute('aria-expanded')=='false', prefix+' Tab closes')
        check(not errors, 'no JavaScript page errors')
        page.goto(base+'/portal/reports/?from_date=2026-01-01&to_date=2026-10-19', wait_until='networkidle')
        link = page.get_by_role('link', name='Tạo slide', exact=True)
        check(link.count()==1, 'report has slide export action')
        with page.expect_download() as download:
            link.click()
        file = download.value
        check(file.failure() is None and file.suggested_filename.endswith('.pptx'), 'browser downloads PPTX successfully')
        target = output/'synthetic-browser-report.pptx'
        file.save_as(target)
        deck = Presentation(io.BytesIO(target.read_bytes()))
        check(len(deck.slides)>=13, 'download uses redesigned complete section deck')
        check(sum(shape.has_chart for slide in deck.slides for shape in slide.shapes)==2, 'download has two editable charts')
        context.close()
        browser.close()
    (output/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'checks':len(report['checks']),'failures':report['failures']}))
    return bool(report['failures'])


if __name__ == '__main__':
    raise SystemExit(main())
