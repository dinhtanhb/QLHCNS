from apps.common.number_format import vi_number
import io
from calendar import monthrange
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from django.db.models import Avg, Count, DecimalField, ExpressionWrapper, F, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.activity.models import Activity, ActivityProposal
from apps.asset.models import Asset, AssetDisposalRequest, AssetMaintenance, InventorySession
from apps.attendance.models import MonthlyAttendanceSummary
from apps.contract.models import Contract
from apps.document.models import Document, OfficialDispatch, PersonnelDecision
from apps.hrm.models import Employee, EmployeePromotion, EmployeeTransfer
from apps.kpi.models import CompanyKPI, DepartmentKPI, EmployeeKPI
from apps.leave.models import LeaveRequest
from apps.lifecycle.models import EmployeeLifecycleCase, LifecycleChecklistItem
from apps.office.models import RoomBooking, SupplyRequest, VehicleBooking
from apps.payroll.models import PayrollPeriod, Payslip
from apps.planning.models import AnnualBudget, Event, HeadcountPlan, TrainingCourse, TrainingEnrollment
from apps.procurement.models import ProcurementItem, ProcurementRequest
from apps.project.models import Project
from apps.recruitment.models import Candidate, Interview, JobPosting, Offer, RecruitmentRequest
from apps.task.models import Task


BLUE = "1D4ED8"
DARK = "0F172A"
LIGHT = "EFF6FF"
MUTED = "64748B"


def _month_starts(from_date, to_date):
    cursor = from_date.replace(day=1)
    months = []
    while cursor <= to_date:
        months.append(cursor)
        cursor = date(cursor.year + (cursor.month == 12), 1 if cursor.month == 12 else cursor.month + 1, 1)
    return months[-12:]


def _metric(label, value, *, unit="", tone="blue"):
    if value is None:
        display = "Chưa có dữ liệu"
    elif unit == "%":
        display = vi_number(value, 1) + "%"
    elif unit == "triệu đồng":
        display = vi_number(value, 1) + " triệu đồng"
    else:
        display = vi_number(int(value)) + (f" {unit}" if unit else "")
    return {"label": label, "value": value, "display": display, "unit": unit, "tone": tone}


def build_report_snapshot(from_date, to_date):
    """Build one aggregate-only snapshot used by HTML and all export formats."""
    today = timezone.localdate()
    report_from = timezone.make_aware(datetime.combine(from_date, time.min))
    report_until = timezone.make_aware(datetime.combine(to_date + timedelta(days=1), time.min))
    months = _month_starts(from_date, to_date)
    month_rows = []
    for month in months:
        end = date(month.year, month.month, monthrange(month.year, month.month)[1])
        month_rows.append({
            "label": f"T{month.month}/{month.year}",
            "new_hires": Employee.objects.filter(join_date__range=(month, end)).count(),
            "departures": Employee.objects.filter(leave_date__range=(month, end), employment_status="Resigned").count(),
            "transfers": EmployeeTransfer.objects.filter(effective_date__range=(month, end)).count(),
            "promotions": EmployeePromotion.objects.filter(effective_date__range=(month, end)).count(),
        })

    total_active = Employee.objects.filter(employment_status="Active").count()
    total_probation = Employee.objects.filter(employment_status="Probation").count()
    new_hires = Employee.objects.filter(join_date__range=(from_date, to_date)).count()
    departures = Employee.objects.filter(leave_date__range=(from_date, to_date), employment_status="Resigned").count()
    transfers = EmployeeTransfer.objects.filter(effective_date__range=(from_date, to_date)).count()
    promotions = EmployeePromotion.objects.filter(effective_date__range=(from_date, to_date)).count()
    turnover = (departures / max(total_active + departures, 1)) * 100

    candidate_counts = {row["status"]: row["total"] for row in Candidate.objects.values("status").annotate(total=Count("pk"))}
    candidates_processing = sum(candidate_counts.get(status, 0) for status in ("New", "Screening", "Interviewing", "Offered"))
    recruitment_funnel = [
        {"label": "Đề xuất đã duyệt", "value": RecruitmentRequest.objects.filter(status="Approved").count()},
        {"label": "Tin đang tuyển", "value": JobPosting.objects.filter(status="Open").count()},
        {"label": "Ứng viên đang xử lý", "value": candidates_processing},
        {"label": "Phỏng vấn chờ kết quả", "value": Interview.objects.filter(result="Pending").count()},
        {"label": "Offer đã chấp nhận", "value": Offer.objects.filter(status="Accepted", offer_date__range=(from_date, to_date)).count()},
        {"label": "Đã tuyển", "value": candidate_counts.get("Hired", 0)},
    ]
    funnel_max = max((item["value"] for item in recruitment_funnel), default=0) or 1
    for item in recruitment_funnel:
        item["width"] = max(4, round(item["value"] / funnel_max * 100)) if item["value"] else 0

    pending_timesheets = MonthlyAttendanceSummary.objects.filter(status="Pending").count()
    pending_leave = LeaveRequest.objects.filter(status="Pending").count()
    leave_in_period = LeaveRequest.objects.filter(status="Approved", from_date__lte=to_date, to_date__gte=from_date).count()

    pending_dispatches = OfficialDispatch.objects.filter(status__in=("Pending", "Processing")).count()
    overdue_dispatches = OfficialDispatch.objects.filter(status__in=("Pending", "Processing"), due_date__lt=today).count()
    pending_rooms = RoomBooking.objects.filter(status="Pending").count()
    pending_vehicles = VehicleBooking.objects.filter(status="Pending").count()
    pending_supplies = SupplyRequest.objects.filter(status="Pending").count()

    asset_total = Asset.objects.count()
    asset_in_use = Asset.objects.filter(status="InUse").count()
    asset_damaged = Asset.objects.filter(condition__in=("Damaged", "Lost")).count()
    maintenance_cost = AssetMaintenance.objects.filter(maintenance_date__range=(from_date, to_date)).aggregate(total=Sum("cost"))["total"] or Decimal("0")
    pending_disposals = AssetDisposalRequest.objects.filter(status="Pending").count()
    active_inventories = InventorySession.objects.filter(status="Draft").count()

    pending_procurement = ProcurementRequest.objects.filter(status="Pending").count()
    purchase_amount_expr = ExpressionWrapper(F("estimated_cost") * F("quantity"), output_field=DecimalField(max_digits=24, decimal_places=2))
    purchase_amount = ProcurementItem.objects.filter(
        procurement_request__request_date__range=(from_date, to_date),
        procurement_request__status__in=("Approved", "Purchased"),
    ).aggregate(total=Sum(purchase_amount_expr))["total"] or Decimal("0")

    open_tasks = Task.objects.filter(status__in=("Todo", "InProgress")).count()
    overdue_tasks = Task.objects.filter(status__in=("Todo", "InProgress"), due_date__lt=today).count()
    completed_tasks = Task.objects.filter(status="Done", completed_date__range=(from_date, to_date)).count()
    kpi_avg = None
    kpi_values = []
    for model in (CompanyKPI, DepartmentKPI, EmployeeKPI):
        value = model.objects.filter(kpi_period__year=to_date.year).aggregate(avg=Avg("achievement"))["avg"]
        if value is not None:
            kpi_values.append(value)
    if kpi_values:
        kpi_avg = sum(kpi_values) / len(kpi_values)

    contract_limit = today + timedelta(days=60)
    expiring_contracts = Contract.objects.filter(status="Active", end_date__range=(today, contract_limit)).count()
    open_payroll_periods = PayrollPeriod.objects.filter(status="Open").count()
    issued_payslips = Payslip.objects.filter(payroll__payroll_period__end_date__range=(from_date, to_date)).count()
    pending_documents = Document.objects.filter(status="PendingApproval").count()
    pending_personnel_decisions = PersonnelDecision.objects.filter(status="PendingApproval").count()
    issued_personnel_decisions = PersonnelDecision.objects.filter(
        status="Issued",
        issued_at__gte=report_from,
        issued_at__lt=report_until,
    ).count()
    onboarding_cases = EmployeeLifecycleCase.objects.filter(case_type="Onboarding", status="Open").count()
    offboarding_cases = EmployeeLifecycleCase.objects.filter(case_type="Offboarding", status="Open").count()
    overdue_checklist_items = LifecycleChecklistItem.objects.filter(status="Pending", due_date__lt=today).count()

    planned_headcount = HeadcountPlan.objects.filter(year=to_date.year).aggregate(total=Sum("planned_headcount"))["total"] or 0
    headcount_gap = int(planned_headcount) - total_active if planned_headcount else None
    allocated_budget = AnnualBudget.objects.filter(year=to_date.year).aggregate(total=Sum("allocated_amount"))["total"] or Decimal("0")
    training_courses = TrainingCourse.objects.filter(start_date__range=(from_date, to_date)).exclude(status="Cancelled").count()
    training_enrollments = TrainingEnrollment.objects.filter(course__start_date__range=(from_date, to_date)).exclude(status="Cancelled").count()
    events = Event.objects.filter(start_time__gte=report_from, start_time__lt=report_until).count()

    active_projects = Project.objects.filter(is_active=True).count()
    active_activities = Activity.objects.filter(start_date__lte=to_date, end_date__gte=from_date).exclude(status="Cancelled").count()
    pending_activity_proposals = ActivityProposal.objects.filter(status__in=("Submitted", "ChangesRequested")).count()

    sections = [
        {"key": "people", "title": "Nhân sự và biến động", "description": "Quy mô hiện tại và các thay đổi trong kỳ", "href": "portal:report_personnel_changes", "metrics": [
            _metric("Đang làm việc", total_active), _metric("Thử việc", total_probation), _metric("Tuyển mới", new_hires, tone="green"),
            _metric("Nghỉ việc", departures, tone="red"), _metric("Điều chuyển", transfers), _metric("Bổ nhiệm", promotions), _metric("Tỷ lệ nghỉ việc", turnover, unit="%", tone="red"),
        ]},
        {"key": "recruitment", "title": "Tuyển dụng", "description": "Nhu cầu tuyển, ứng viên, phỏng vấn và offer", "href": None, "metrics": [
            _metric("Đề xuất chờ duyệt", RecruitmentRequest.objects.filter(status="Pending").count(), tone="amber"),
            _metric("Tin đang tuyển", JobPosting.objects.filter(status="Open").count()),
            _metric("Ứng viên đang xử lý", candidates_processing),
            _metric("Phỏng vấn chờ kết quả", Interview.objects.filter(result="Pending").count(), tone="amber"),
            _metric("Đã tuyển", candidate_counts.get("Hired", 0), tone="green"),
        ]},
        {"key": "attendance", "title": "Chấm công và nghỉ phép", "description": "Tình trạng duyệt bảng công và đơn nghỉ", "href": "portal:report_timesheet", "metrics": [
            _metric("Bảng công chờ duyệt", pending_timesheets, tone="amber"), _metric("Đơn nghỉ chờ duyệt", pending_leave, tone="amber"),
            _metric("Đơn nghỉ đã duyệt trong kỳ", leave_in_period),
        ]},
        {"key": "office", "title": "Hành chính", "description": "Công văn, phòng họp, xe và văn phòng phẩm", "href": None, "metrics": [
            _metric("Công văn đang xử lý", pending_dispatches), _metric("Công văn quá hạn", overdue_dispatches, tone="red"),
            _metric("Phòng họp chờ duyệt", pending_rooms, tone="amber"), _metric("Xe chờ duyệt", pending_vehicles, tone="amber"),
            _metric("VPP chờ duyệt", pending_supplies, tone="amber"),
        ]},
        {"key": "assets", "title": "Tài sản và mua sắm", "description": "Tình trạng tài sản, kiểm kê và đề nghị mua sắm", "href": None, "metrics": [
            _metric("Tổng tài sản", asset_total), _metric("Đang sử dụng", asset_in_use), _metric("Hư hỏng hoặc thất lạc", asset_damaged, tone="red"),
            _metric("Kiểm kê đang mở", active_inventories), _metric("Thanh lý chờ duyệt", pending_disposals, tone="amber"),
            _metric("Mua sắm chờ duyệt", pending_procurement, tone="amber"), _metric("Chi phí bảo trì", maintenance_cost / Decimal("1000000"), unit="triệu đồng"),
            _metric("Giá trị mua sắm trong kỳ", purchase_amount / Decimal("1000000"), unit="triệu đồng"),
        ]},
        {"key": "work", "title": "Công việc và KPI", "description": "Tiến độ công việc và mức hoàn thành KPI", "href": None, "metrics": [
            _metric("Công việc đang mở", open_tasks), _metric("Công việc quá hạn", overdue_tasks, tone="red"),
            _metric("Hoàn thành trong kỳ", completed_tasks, tone="green"), _metric("KPI bình quân", kpi_avg, unit="%"),
        ]},
        {"key": "contracts", "title": "Hợp đồng và lương", "description": "Chỉ hiển thị trạng thái, không hiển thị số tiền lương", "href": None, "metrics": [
            _metric("Hợp đồng hết hạn trong 60 ngày", expiring_contracts, tone="amber"), _metric("Kỳ lương đang mở", open_payroll_periods),
            _metric("Phiếu lương đã phát hành trong kỳ", issued_payslips),
        ]},
        {"key": "documents", "title": "Hồ sơ và quy trình nhân sự", "description": "Tài liệu, quyết định và checklist tiếp nhận hoặc nghỉ việc", "href": None, "metrics": [
            _metric("Tài liệu chờ phê duyệt", pending_documents, tone="amber"),
            _metric("Quyết định nhân sự chờ duyệt", pending_personnel_decisions, tone="amber"),
            _metric("Quyết định đã ban hành trong kỳ", issued_personnel_decisions),
            _metric("Hồ sơ tiếp nhận đang mở", onboarding_cases), _metric("Hồ sơ nghỉ việc đang mở", offboarding_cases),
            _metric("Checklist quá hạn", overdue_checklist_items, tone="red"),
        ]},
        {"key": "planning", "title": "Kế hoạch, ngân sách và đào tạo", "description": "Định biên, ngân sách năm, đào tạo và sự kiện", "href": "portal:report_headcount", "metrics": [
            _metric("Định biên năm", planned_headcount), _metric("Chênh lệch định biên", headcount_gap, tone="amber"),
            _metric("Ngân sách được duyệt", allocated_budget / Decimal("1000000"), unit="triệu đồng"),
            _metric("Khóa đào tạo trong kỳ", training_courses), _metric("Lượt học viên", training_enrollments), _metric("Sự kiện trong kỳ", events),
        ]},
        {"key": "projects", "title": "Dự án và hoạt động", "description": "Dự án, hoạt động và hồ sơ đề xuất", "href": None, "metrics": [
            _metric("Dự án đang hoạt động", active_projects), _metric("Hoạt động trong kỳ", active_activities),
            _metric("Hồ sơ hoạt động cần xử lý", pending_activity_proposals, tone="amber"),
        ]},
    ]
    movement_total = max((max(row["new_hires"], row["departures"], row["transfers"], row["promotions"]) for row in month_rows), default=0) or 1
    for row in month_rows:
        for key in ("new_hires", "departures", "transfers", "promotions"):
            row[f"{key}_height"] = max(3, round(row[key] / movement_total * 100)) if row[key] else 0
    return {
        "from_date": from_date,
        "to_date": to_date,
        "generated_at": timezone.localtime(),
        "sections": sections,
        "month_rows": month_rows,
        "recruitment_funnel": recruitment_funnel,
        "headline": {
            "active": total_active, "new_hires": new_hires, "departures": departures,
            "candidates": candidates_processing, "overdue_tasks": overdue_tasks, "overdue_dispatches": overdue_dispatches,
        },
    }


def _safe_excel_text(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def export_snapshot_xlsx(snapshot):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "BÁO CÁO ĐIỀU HÀNH IHRMS"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=DARK)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:D2")
    ws["A2"] = f"Kỳ báo cáo: {snapshot['from_date']:%d/%m/%Y} - {snapshot['to_date']:%d/%m/%Y}"
    ws["A2"].font = Font(name="Arial", italic=True, color=MUTED)
    row = 4
    thin = Side(style="thin", color="CBD5E1")
    for section in snapshot["sections"]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row, 1, section["title"])
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        row += 1
        for metric in section["metrics"]:
            ws.cell(row, 1, _safe_excel_text(metric["label"]))
            ws.cell(row, 2, _safe_excel_text(metric["display"]))
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            for col in range(1, 5):
                ws.cell(row, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 38
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 20
    ws.freeze_panes = "A4"

    movement = wb.create_sheet("Biến động nhân sự")
    movement.append(["Tháng", "Tuyển mới", "Nghỉ việc", "Điều chuyển", "Bổ nhiệm"])
    for item in snapshot["month_rows"]:
        movement.append([item["label"], item["new_hires"], item["departures"], item["transfers"], item["promotions"]])
    for cell in movement[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
    if snapshot["month_rows"]:
        chart = LineChart()
        chart.title = "Biến động nhân sự theo tháng"
        chart.y_axis.title = "Số lượng"
        chart.x_axis.title = "Tháng"
        chart.add_data(Reference(movement, min_col=2, max_col=5, min_row=1, max_row=movement.max_row), titles_from_data=True)
        chart.set_categories(Reference(movement, min_col=1, min_row=2, max_row=movement.max_row))
        chart.height = 8
        chart.width = 17
        movement.add_chart(chart, "G2")
    for index, width in enumerate((16, 14, 14, 14, 14), 1):
        movement.column_dimensions[get_column_letter(index)].width = width

    recruitment = wb.create_sheet("Tuyển dụng")
    recruitment.append(["Giai đoạn", "Số lượng"])
    for item in snapshot["recruitment_funnel"]:
        recruitment.append([item["label"], item["value"]])
    for cell in recruitment[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Phễu tuyển dụng"
    chart.add_data(Reference(recruitment, min_col=2, min_row=1, max_row=recruitment.max_row), titles_from_data=True)
    chart.set_categories(Reference(recruitment, min_col=1, min_row=2, max_row=recruitment.max_row))
    chart.height = 8
    chart.width = 16
    recruitment.add_chart(chart, "D2")
    recruitment.column_dimensions["A"].width = 32
    recruitment.column_dimensions["B"].width = 14

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_snapshot_docx(snapshot):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.65)
    title = document.add_paragraph(style="Title")
    title.add_run("Báo cáo điều hành IHRMS")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.runs[0].font.name = "Arial"
    title.runs[0].font.size = Pt(20)
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    period = document.add_paragraph(f"Kỳ báo cáo từ {snapshot['from_date']:%d/%m/%Y} đến {snapshot['to_date']:%d/%m/%Y}")
    period.runs[0].font.name = "Arial"
    period.runs[0].font.size = Pt(10)

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    for section_data in snapshot["sections"]:
        heading = document.add_heading(section_data["title"], level=1)
        heading.style.font.name = "Arial"
        heading.style.font.color.rgb = RGBColor(0, 0, 0)
        document.add_paragraph(section_data["description"])
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width = Inches(4.8)
        table.columns[1].width = Inches(3.0)
        for index, text in enumerate(("Chỉ số", "Giá trị")):
            cell = table.rows[0].cells[index]
            cell.text = text
            shade(cell, BLUE)
            for run in cell.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
        for index, metric in enumerate(section_data["metrics"]):
            cells = table.add_row().cells
            cells[0].text = metric["label"]
            cells[1].text = metric["display"]
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in cell.paragraphs[0].runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
            if index % 2:
                for cell in cells:
                    shade(cell, LIGHT)
        document.add_paragraph()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_snapshot_pptx(snapshot):
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(15, 23, 42)
    cover.shapes.title.text = "Báo cáo điều hành IHRMS"
    cover.placeholders[1].text = f"Kỳ báo cáo {snapshot['from_date']:%d/%m/%Y} đến {snapshot['to_date']:%d/%m/%Y}"
    for shape in cover.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(255, 255, 255)

    summary = deck.slides.add_slide(deck.slide_layouts[5])
    summary.shapes.title.text = "Tổng quan"
    metrics = [
        ("Nhân sự đang làm việc", snapshot["headline"]["active"]),
        ("Tuyển mới trong kỳ", snapshot["headline"]["new_hires"]),
        ("Nghỉ việc trong kỳ", snapshot["headline"]["departures"]),
        ("Ứng viên đang xử lý", snapshot["headline"]["candidates"]),
        ("Công việc quá hạn", snapshot["headline"]["overdue_tasks"]),
        ("Công văn quá hạn", snapshot["headline"]["overdue_dispatches"]),
    ]
    table = summary.shapes.add_table(3, 4, Inches(0.8), Inches(1.5), Inches(11.8), Inches(4.8)).table
    for index, (label, value) in enumerate(metrics):
        row, col = divmod(index, 2)
        label_cell = table.cell(row, col * 2)
        value_cell = table.cell(row, col * 2 + 1)
        label_cell.text, value_cell.text = label, str(value)
        for cell in (label_cell, value_cell):
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(17)
                    run.font.color.rgb = RGBColor(15, 23, 42)
        value_cell.text_frame.paragraphs[0].runs[0].font.bold = True
        value_cell.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    movement = deck.slides.add_slide(deck.slide_layouts[5])
    movement.shapes.title.text = "Biến động nhân sự theo tháng"
    chart_data = ChartData()
    chart_data.categories = [item["label"] for item in snapshot["month_rows"]]
    for label, key in (("Tuyển mới", "new_hires"), ("Nghỉ việc", "departures"), ("Điều chuyển", "transfers"), ("Bổ nhiệm", "promotions")):
        chart_data.add_series(label, [item[key] for item in snapshot["month_rows"]])
    chart = movement.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.4), Inches(12), Inches(5.4), chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.value_axis.has_major_gridlines = True

    recruitment = deck.slides.add_slide(deck.slide_layouts[5])
    recruitment.shapes.title.text = "Tuyển dụng"
    recruit_data = ChartData()
    recruit_data.categories = [item["label"] for item in snapshot["recruitment_funnel"]]
    recruit_data.add_series("Số lượng", [item["value"] for item in snapshot["recruitment_funnel"]])
    chart = recruitment.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.8), Inches(1.4), Inches(11.8), Inches(5.4), recruit_data).chart
    chart.has_legend = False

    for section_data in snapshot["sections"][2:]:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = section_data["title"]
        rows = len(section_data["metrics"]) + 1
        table = slide.shapes.add_table(rows, 2, Inches(1.1), Inches(1.45), Inches(11.1), Inches(5.3)).table
        table.columns[0].width = Inches(7.5)
        table.columns[1].width = Inches(3.6)
        for col, text in enumerate(("Chỉ số", "Giá trị")):
            table.cell(0, col).text = text
        for row, metric in enumerate(section_data["metrics"], 1):
            table.cell(row, 0).text = metric["label"]
            table.cell(row, 1).text = metric["display"]
        for row in range(rows):
            for col in range(2):
                cell = table.cell(row, col)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(29, 78, 216) if row == 0 else RGBColor(248, 250, 252)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(15)
                        run.font.color.rgb = RGBColor(255, 255, 255) if row == 0 else RGBColor(15, 23, 42)
                        run.font.bold = row == 0
    for slide in list(deck.slides)[1:]:
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()
