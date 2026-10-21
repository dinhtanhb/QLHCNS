"""Editable executive slides from the existing aggregate report snapshot.

Presentation only: no queries, permissions, or recalculation of report metrics.
"""
import io
import math

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


# Palette Chàm Đông (indigo), đồng bộ với khối light của ui_tokens.css.
# Đổi màu ở ui_tokens.css thì sửa cả đây; test_report_palette.py giữ hai nơi khớp nhau.
NAVY = "101828"
TEAL = "3A4FD0"
MINT = "EEF1FD"
PAPER = "E7ECF4"
WHITE = "FFFFFF"
MUTED = "475467"
LINE = "DFE4EE"
RED = "B42318"
AMBER = "B54708"
BLUE = "175CD3"
FONT = "Arial"


def _rgb(color):
    return RGBColor.from_string(color)


class _Slides:
    def __init__(self, snapshot):
        self.snapshot = snapshot
        self.deck = Presentation()
        self.deck.slide_width = Inches(13.333333)
        self.deck.slide_height = Inches(7.5)

    def text(self, slide, name, text, x, y, w, h, *, size=20,
             color=NAVY, bold=False, align=PP_ALIGN.LEFT):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        self.style_text(shape, text, size=size, color=color, bold=bold, align=align)
        return shape

    @staticmethod
    def style_text(shape, text, *, size, color, bold=False, align=PP_ALIGN.LEFT):
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(str(text).split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_after = Pt(0)
            paragraph.font.name = FONT
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = _rgb(color)
            paragraph.text = line

    @staticmethod
    def rect(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(color)
        shape.line.fill.background()
        # Override the default Office shape style's inherited shadow.
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        style = shape._element.find(qn("p:style"))
        if style is not None:
            for tag in ("a:effectRef", "a:lnRef"):
                ref = style.find(qn(tag))
                if ref is not None:
                    ref.set("idx", "0")
        return shape

    def slide(self, title, description="", *, dark=False):
        slide = self.deck.slides.add_slide(self.deck.slide_layouts[5])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(NAVY if dark else WHITE)
        shape = slide.shapes.title
        shape.left, shape.top = Inches(.65), Inches(.48)
        shape.width, shape.height = Inches(12), Inches(.64)
        self.style_text(shape, title, size=30 if len(title) > 42 else 34,
                        color=WHITE if dark else NAVY, bold=True)
        if description:
            self.text(slide, "section-description", description, .67, 1.17, 11.9, .48,
                      size=16, color=MINT if dark else MUTED)
        return slide

    def metric(self, slide, metric, name, x, y, w, *, size=44, dark=False, label_size=18):
        value = str(metric["display"])
        # Keep the complete formatted value, including its unit and sign.
        # Wide currency values wrap in a larger field rather than being clipped.
        number_size = min(size, 29 if len(value) > 17 else 36 if len(value) > 10 else size)
        color = WHITE if dark else NAVY
        if metric.get("value") not in (None, 0) and not dark:
            color = {"red": RED, "amber": AMBER, "green": TEAL}.get(metric.get("tone"), NAVY)
        lines = max(1, math.ceil(len(value)*number_size*.55/(w*72)))
        value_height = max(.65, number_size/72*1.22*lines)
        self.text(slide, name+":value", value, x, y, w, value_height,
                  size=number_size, color=color, bold=True)
        self.text(slide, name+":label", metric["label"], x, y+value_height+.14, w, .66,
                  size=label_size, color=MINT if dark else MUTED)

    def grid(self, slide, metrics, key, *, x=.7, y=1.95, w=11.9, columns=3, dark=False):
        rows = max(1, math.ceil(len(metrics) / columns))
        row_height = min(2.15, 4.65 / rows)
        cell_width = w / columns
        for index, metric in enumerate(metrics):
            row, col = divmod(index, columns)
            # A compact row retains an 18pt label and a complete value.
            top = y + row * row_height
            if row_height < 1.6:
                self.text(slide, f"metric:{key}:{index}:label", metric["label"],
                          x+col*cell_width, top, cell_width-.35, .58, size=18,
                          color=MINT if dark else MUTED)
                self.text(slide, f"metric:{key}:{index}:value", metric["display"],
                          x+col*cell_width, top+.62, cell_width-.35, .48, size=27,
                          color=WHITE if dark else NAVY, bold=True)
            else:
                self.metric(slide, metric, f"metric:{key}:{index}",
                            x+col*cell_width, top, cell_width-.4, size=46, dark=dark)
            if row < rows-1:
                self.rect(slide, x+col*cell_width, top+row_height-.04, cell_width-.4, .012,
                          MUTED if dark else LINE)

    def hero(self, slide, section):
        metrics, key = section["metrics"], section["key"]
        if not metrics:
            self.text(slide, "empty-section", "Chưa có chỉ số trong nhóm này.", .7, 2.2, 11, 1)
            return
        self.rect(slide, .65, 1.9, 3.65, 4.6, PAPER)
        self.metric(slide, metrics[0], f"metric:{key}:0", .95, 2.55, 3.05, size=78, label_size=23)
        # The remaining measures stay distinct, never added into a new total.
        for index, metric in enumerate(metrics[1:], 1):
            row, col = divmod(index-1, 2)
            top = 2.0 + row * min(1.6, 4.65/max(1, math.ceil((len(metrics)-1)/2)))
            self.metric(slide, metric, f"metric:{key}:{index}", 4.8+col*3.9, top, 3.5, size=39)

    def chart(self, slide, data, kind, name, x, y, w, h, *, colors=(TEAL, RED, AMBER, BLUE)):
        shape = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
        shape.name = name
        chart = shape.chart
        chart.has_title = False
        chart.font.name = FONT
        chart.font.size = Pt(14)
        chart.has_legend = len(chart.series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(15)
        chart.value_axis.minimum_scale = 0
        maximum = max((float(value) for series in data for value in series.values), default=0)
        rough_step = max(1, maximum/5)
        magnitude = 10**math.floor(math.log10(rough_step))
        step = next((n*magnitude for n in (1,2,5,10) if n*magnitude>=rough_step),10*magnitude)
        chart.value_axis.major_unit = step
        chart.value_axis.maximum_scale = max(1, math.ceil(maximum/step)*step)
        chart.value_axis.tick_labels.font.size = Pt(14)
        chart.category_axis.tick_labels.font.size = Pt(14)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(LINE)
        for index, series in enumerate(chart.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(colors[index % len(colors)])
            series.format.line.fill.background()
        return chart

    def build(self):
        s = self.snapshot
        cover = self.slide("Báo cáo điều hành IHRMS", dark=True)
        title = cover.shapes.title
        title.left, title.top, title.width, title.height = map(Inches, (.75, 1.8, 8.1, 2.55))
        self.style_text(title, "Báo cáo điều hành IHRMS", size=52, color=WHITE, bold=True)
        self.rect(cover, 9.35, 0, 3.983333, 7.5, MINT)
        years = str(s["to_date"].year) if s["from_date"].year == s["to_date"].year else f'{s["from_date"].year}\n{s["to_date"].year}'
        self.text(cover, "report-years", years, 9.8, 2.35, 3.05, 2.55, size=66, color=TEAL, bold=True)
        self.text(cover, "period", f'{s["from_date"]:%d/%m/%Y} — {s["to_date"]:%d/%m/%Y}',
                  .8, 5.6, 8, .7, size=27, color=MINT)

        overview = self.slide("Tổng quan điều hành", "Số liệu tổng hợp theo phạm vi báo cáo; các chỉ số không cộng gộp với nhau.")
        headlines = [("Nhân sự đang làm việc", "active"), ("Tuyển mới trong kỳ", "new_hires"),
                     ("Nghỉ việc trong kỳ", "departures"), ("Ứng viên đang xử lý", "candidates"),
                     ("Công việc quá hạn", "overdue_tasks"), ("Công văn quá hạn", "overdue_dispatches")]
        self.grid(overview, [{"label":label,"display":str(s["headline"][key]),"value":s["headline"][key],
                              "tone":"red" if key in ("departures","overdue_tasks","overdue_dispatches") else "green" if key=="new_hires" else "blue"}
                             for label,key in headlines], "headline")

        movement = self.slide("Biến động nhân sự theo tháng", "Tuyển mới, nghỉ việc, điều chuyển và bổ nhiệm")
        months = s["month_rows"]
        movement_notes = []
        if months:
            data = ChartData()
            data.categories = [row["label"] for row in months]
            series = (("Tuyển mới","new_hires"),("Nghỉ việc","departures"),
                      ("Điều chuyển","transfers"),("Bổ nhiệm","promotions"))
            for label,key in series:
                data.add_series(label,[row[key] for row in months])
            chart = self.chart(movement, data, XL_CHART_TYPE.COLUMN_CLUSTERED,
                               "movement-chart", .65, 1.85, 12, 4.75)
            if not any(row[key] for row in months for _,key in series):
                chart.value_axis.maximum_scale = 1
                movement_notes.append("Các chỉ số trên biểu đồ đều bằng 0.")
        else:
            self.text(movement,"empty-movement","Không có chuỗi số liệu theo tháng.",.8,3,11,1,size=26,color=MUTED)
        period_months = (s["to_date"].year-s["from_date"].year)*12+s["to_date"].month-s["from_date"].month+1
        if months and period_months > len(months):
            movement_notes.append(f"Biểu đồ hiển thị {len(months)} tháng gần nhất; chỉ số tổng hợp vẫn theo kỳ báo cáo.")
        if movement_notes:
            self.text(movement,"movement-scope"," ".join(movement_notes),.8,6.62,11.8,.46,size=13,color=MUTED)

        for section in s["sections"]:
            key = section["key"]
            metrics = section["metrics"]
            # More than eight measures continue on another page, never shrink
            # the entire section to unreadable type or drop excess measures.
            page_size = 6 if any(len(str(m["display"])) > 20 for m in metrics) else 8
            pages = [metrics[i:i+page_size] for i in range(0,len(metrics),page_size)] or [[]]
            for page_index, page_metrics in enumerate(pages):
                current = {**section, "metrics":page_metrics}
                dark = key == "contracts"
                title = section["title"] + (f" — {page_index+1}/{len(pages)}" if len(pages)>1 else "")
                slide = self.slide(title, section.get("description", ""), dark=dark)
                if key == "recruitment" and page_index == 0:
                    funnel = s["recruitment_funnel"]
                    if funnel:
                        data = ChartData();data.categories = [row["label"] for row in funnel]
                        data.add_series("Số lượng",[row["value"] for row in funnel])
                        chart = self.chart(slide,data,XL_CHART_TYPE.BAR_CLUSTERED,"recruitment-chart",.6,1.8,8,4.7)
                        chart.plots[0].has_data_labels = True
                        chart.plots[0].data_labels.font.size = Pt(18)
                        if not any(row["value"] for row in funnel):chart.value_axis.maximum_scale = 1
                    # Keep all section metrics as well as the separate chart.
                    for index,metric in enumerate(page_metrics):
                        y = 1.85+index*(4.7/max(len(page_metrics),1))
                        self.text(slide,f"metric:{key}:{index}:label",metric["label"],9,y,3.55,.37,size=16,color=MUTED)
                        self.text(slide,f"metric:{key}:{index}:value",metric["display"],9,y+.44,3.55,.43,size=25,color=TEAL,bold=True)
                    self.text(slide,"recruitment-scope","Các nhóm có phạm vi khác nhau; không suy tỷ lệ chuyển đổi từ biểu đồ này.",.8,6.6,11.8,.45,size=14,color=MUTED)
                elif key in ("people","projects") and page_index == 0:
                    self.hero(slide,current)
                elif key == "office" and len(page_metrics)>2 and page_index == 0:
                    self.rect(slide,.65,1.9,5.5,4.6,PAPER)
                    for index,metric in enumerate(page_metrics[:2]):
                        self.metric(slide,metric,f"metric:{key}:{index}",1,2.1+index*2.1,4.75,size=64)
                    for index,metric in enumerate(page_metrics[2:],2):
                        self.metric(slide,metric,f"metric:{key}:{index}",7,2.0+(index-2)*1.5,5.2,size=38)
                elif key == "planning" and len(page_metrics)==6:
                    self.rect(slide,.65,1.9,6,2.9,PAPER)
                    self.metric(slide,page_metrics[2],f"metric:{key}:2",1,2.55,5.25,size=56)
                    for index in (0,1):self.metric(slide,page_metrics[index],f"metric:{key}:{index}",7.2+index*2.85,2.1,2.5,size=39)
                    for index,metric in enumerate(page_metrics[3:],3):self.metric(slide,metric,f"metric:{key}:{index}",.8+(index-3)*4.1,5.15,3.7,size=38)
                else:
                    self.grid(slide,page_metrics,key,columns=2 if len(page_metrics)>6 else 3,dark=dark)
                slide.notes_slide.notes_text_frame.text = (
                    f"{section['title']}\n" + "\n".join(f"{m['label']}: {m['display']}" for m in page_metrics)
                    + "\nSố liệu lấy nguyên từ snapshot báo cáo. Không suy diễn xu hướng hoặc cộng các chỉ số khác đơn vị."
                )

        count = len(self.deck.slides)
        for index,slide in enumerate(self.deck.slides,1):
            dark = slide.background.fill.fore_color.rgb == _rgb(NAVY)
            footer = f'IHRMS · {s["from_date"]:%d/%m/%Y}–{s["to_date"]:%d/%m/%Y} · {index:02}/{count:02}'
            self.text(slide,"report-footer",footer,.7,7.16,11.9,.24,size=11,color=MINT if dark else MUTED)
        self.deck.core_properties.title = "Báo cáo điều hành IHRMS"
        self.deck.core_properties.subject = "Số liệu tổng hợp theo kỳ báo cáo"
        output = io.BytesIO()
        self.deck.save(output)
        return output.getvalue()


def build_report_presentation(snapshot):
    return _Slides(snapshot).build()
