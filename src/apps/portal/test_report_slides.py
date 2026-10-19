"""Presentation contracts: complete evidence, editable charts and safe layout."""
import copy
import io
from datetime import date

from django.test import SimpleTestCase
from pptx import Presentation

from .reporting_dashboard import export_snapshot_pptx


def sample_snapshot():
    metrics = [{"label":f"Chỉ số nghiệp vụ {i}", "value":i, "display":str(i), "tone":"blue"} for i in range(8)]
    return {
        "from_date":date(2026,1,1), "to_date":date(2026,10,19),
        "headline":dict(active=37,new_hires=32,departures=0,candidates=2,overdue_tasks=4,overdue_dispatches=1),
        "month_rows":[dict(label=f"T{i}/2026",new_hires=i,departures=0,transfers=1,promotions=0) for i in range(1,11)],
        "recruitment_funnel":[dict(label="Ứng viên đang xử lý",value=2),dict(label="Đã tuyển",value=0)],
        "sections":[
            dict(key="people",title="Nhân sự và biến động",metrics=[dict(label="Thử việc",value=3,display="3")]),
            dict(key="recruitment",title="Tuyển dụng",metrics=[dict(label="Đề xuất chờ duyệt",value=4,display="4")]),
            dict(key="assets",title="Tài sản và mua sắm",metrics=metrics),
        ],
    }


class ReportSlideTests(SimpleTestCase):
    def deck(self,snapshot):
        return Presentation(io.BytesIO(export_snapshot_pptx(snapshot)))

    def test_every_section_measure_is_visible_and_snapshot_is_not_mutated(self):
        snapshot=sample_snapshot();before=copy.deepcopy(snapshot);deck=self.deck(snapshot)
        visible=[shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
        for section in snapshot['sections']:
            for metric in section['metrics']:
                self.assertIn(metric['label'],visible)
                self.assertIn(metric['display'],visible)
        self.assertEqual(snapshot,before)
        self.assertEqual(deck.slides[0].shapes.title.text,"Báo cáo điều hành IHRMS")

    def test_charts_are_editable_and_preserve_every_month_and_series_value(self):
        snapshot=sample_snapshot();deck=self.deck(snapshot)
        charts={shape.name:shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart}
        movement=charts['movement-chart']
        self.assertEqual([c.label for c in movement.plots[0].categories],[r['label'] for r in snapshot['month_rows']])
        for series,key in zip(movement.series,('new_hires','departures','transfers','promotions')):
            self.assertEqual(list(series.values),[r[key] for r in snapshot['month_rows']])
        self.assertTrue(movement.part.chart_workbook.xlsx_part.blob)
        self.assertEqual(list(charts['recruitment-chart'].series[0].values),[2,0])

    def test_long_values_and_extra_measures_continue_without_off_slide_objects(self):
        snapshot=sample_snapshot();metrics=snapshot['sections'][2]['metrics']
        metrics[0].update(display="999.999.999.999.999,9 triệu đồng")
        metrics.extend(dict(label=f"Chỉ số bổ sung {i}",value=i,display=str(i)) for i in range(8,17))
        deck=self.deck(snapshot)
        self.assertGreaterEqual(sum(slide.shapes.title.text.startswith('Tài sản và mua sắm') for slide in deck.slides),3)
        texts=[shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
        for metric in metrics:self.assertIn(metric['label'],texts);self.assertIn(metric['display'],texts)
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                self.assertLessEqual(shape.left+shape.width,deck.slide_width+2)
                self.assertLessEqual(shape.top+shape.height,deck.slide_height+2)
            fields={shape.name:shape for shape in slide.shapes if shape.name.startswith('metric:')}
            for name,value in fields.items():
                if not name.endswith(':value'):continue
                label=fields[name[:-5]+'label']
                self.assertTrue(value.top+value.height<=label.top or label.top+label.height<=value.top,
                                f"Overlapping value/label on {slide.shapes.title.text}: {name}")

    def test_zero_and_unknown_values_are_distinct_and_long_period_is_explicit(self):
        snapshot=sample_snapshot();snapshot['from_date']=date(2024,1,1)
        for row in snapshot['month_rows']:
            for key in ('new_hires','departures','transfers','promotions'):row[key]=0
        snapshot['sections'][2]['metrics'][0].update(value=None,display="Chưa có dữ liệu")
        deck=self.deck(snapshot)
        text='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Chưa có dữ liệu',text)
        self.assertIn('Các chỉ số trên biểu đồ đều bằng 0.',text)
        self.assertIn('10 tháng gần nhất',text)
