import io
from datetime import date

from django.test import TestCase
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from apps.portal.reporting_dashboard import (
    build_report_snapshot,
    export_snapshot_docx,
    export_snapshot_pptx,
    export_snapshot_xlsx,
)
from apps.security.models import SecurityEvent, User


class ReportDashboardTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser(
            username="report-dashboard-admin",
            email="report-dashboard@example.test",
            password="test-password",
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_report_hub_renders_all_available_sections_and_export_actions(self):
        response = self.client.get(reverse("portal:report_hub"), {
            "from_date": "2026-01-01", "to_date": "2026-12-31",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nhân sự và biến động")
        self.assertContains(response, "Tuyển dụng")
        self.assertContains(response, "Hành chính")
        self.assertContains(response, "Tài sản và mua sắm")
        self.assertContains(response, "Công việc và KPI")
        self.assertContains(response, "Hợp đồng và lương")
        self.assertContains(response, "Dự án và hoạt động")
        self.assertContains(response, "Xuất Excel")
        self.assertContains(response, "Xuất Word")
        self.assertContains(response, "Tạo slide")
        self.assertContains(response, reverse("portal:report_budget"))
        self.assertContains(response, "Xem ngân sách")
        self.assertContains(response, "Xem định biên")
        self.assertContains(response, "Xem số liệu theo tháng")
        self.assertContains(response, "css/report_dashboard.css")
        self.assertContains(response, 'class="movement-chart" role="region" tabindex="0"')

    def test_snapshot_exports_are_valid_office_files(self):
        snapshot = build_report_snapshot(date(2026, 1, 1), date(2026, 12, 31))

        excel_payload = export_snapshot_xlsx(snapshot)
        workbook = load_workbook(io.BytesIO(excel_payload))
        self.assertEqual(workbook.sheetnames, ["Tổng quan", "Biến động nhân sự", "Tuyển dụng"])
        self.assertEqual(workbook["Tổng quan"]["A1"].value, "BÁO CÁO ĐIỀU HÀNH IHRMS")

        word_payload = export_snapshot_docx(snapshot)
        document = Document(io.BytesIO(word_payload))
        self.assertEqual(document.paragraphs[0].text, "Báo cáo điều hành IHRMS")
        self.assertGreaterEqual(len(document.tables), 9)

        slide_payload = export_snapshot_pptx(snapshot)
        presentation = Presentation(io.BytesIO(slide_payload))
        self.assertGreaterEqual(len(presentation.slides), 10)
        self.assertEqual(presentation.slides[0].shapes.title.text, "Báo cáo điều hành IHRMS")

    def test_export_endpoints_return_expected_content_types_and_audit(self):
        expected = {
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        for export_format, content_type in expected.items():
            with self.subTest(export_format=export_format):
                response = self.client.get(
                    reverse("portal:report_hub_export", args=(export_format,)),
                    {"from_date": "2026-01-01", "to_date": "2026-12-31"},
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], content_type)
                self.assertIn(f".{export_format}", response["Content-Disposition"])
                self.assertGreater(len(response.content), 1000)
        self.assertEqual(SecurityEvent.objects.filter(action="data_export").count(), 3)

    def test_export_rejects_invalid_period_and_format(self):
        invalid_period = self.client.get(
            reverse("portal:report_hub_export", args=("xlsx",)),
            {"from_date": "2026-12-31", "to_date": "2026-01-01"},
        )
        self.assertEqual(invalid_period.status_code, 400)
        invalid_format = self.client.get(
            reverse("portal:report_hub_export", args=("pdf",)),
            {"from_date": "2026-01-01", "to_date": "2026-12-31"},
        )
        self.assertEqual(invalid_format.status_code, 404)

    def test_export_requires_report_permission(self):
        ordinary_user = User.objects.create_user(username="report-no-permission", password="test-password")
        self.client.force_login(ordinary_user)
        response = self.client.get(
            reverse("portal:report_hub_export", args=("xlsx",)),
            {"from_date": "2026-01-01", "to_date": "2026-12-31"},
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(SecurityEvent.objects.filter(user=ordinary_user, action="data_export").exists())
