import re
import logging
import datetime
from django.conf import settings
from django.core.paginator import Paginator
from django.core.exceptions import ValidationError
from django.db import models, transaction
from django.views.decorators.http import require_POST
from apps.security.permissions import get_role_code, has_permission
from django.db.models import Count, Prefetch, Sum, Q
from apps.recruitment.models import RecruitmentRequest
from apps.document.models import Document, OfficialDispatch, DispatchTracking, DispatchBook
from apps.task.models import Task
from apps.activity.permissions import can_view_activity, visible_activities_for
from apps.common.uploads import validate_upload
from apps.m365.graph_client import GraphAPIError
from apps.hrm.models import Employee, EmployeeProfileChangeRequest
from apps.asset.models import Asset
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from apps.attendance.models import Attendance
from apps.leave.models import LeaveRequest, LeaveType, LeaveApproval
from apps.payroll.models import Payroll
from apps.notification.services import notify, notify_with_email_draft
from apps.notification.models import Notification, NotificationEmailOutbox
from decimal import Decimal, InvalidOperation
import calendar as calendar_module
from django.db.models import Sum
from django.http import HttpResponse, Http404, HttpResponseForbidden, FileResponse
from django.views.decorators.http import require_GET
from apps.project.models import Project
from apps.attendance.models import TimesheetEntry, WorkingCalendar, Holiday, MonthlyAttendanceSummary
from apps.leave.models import LeaveBalance
from django.utils import timezone
from apps.attendance.models import ProjectTimeAllocation  # nếu chưa import ở bước trước, thêm vào
from apps.hrm.models import EmployeeTransfer, EmployeePromotion
from apps.office.models import MeetingRoom, Vehicle, RoomBooking, VehicleBooking, SupplyItem, SupplyRequest, SupplyRequestItem
from apps.planning.models import HeadcountPlan, TrainingEnrollment, AnnualBudget, Event
from apps.organization.models import Department
from apps.hrm.models import EmployeeReward, EmployeeDiscipline
from apps.contract.models import ContractType, Contract, ContractAppendix
from apps.recruitment.models import Candidate, Interview, Offer
from apps.organization.models import Position, JobTitle
from apps.payroll.models import EmployeeSalary
from apps.payroll.models import (
    EmployeeSalary, InsuranceRateConfig, PersonalIncomeTaxBracket,
    TaxDeductionConfig, EmployeeDependent, PayrollPeriod, Payroll, Payslip,
)
from apps.asset.models import (
    AssetType, Asset, AssetAllocation,
    InventorySession, InventoryMember, InventoryLine,
    AssetDisposalRequest, AssetDisposalRequestMember, AssetDisposalItem, AssetDisposalMinutes,
)
from apps.procurement.models import ProcurementRequest, ProcurementItem, ProcurementApproval
from apps.common.params import parse_int

logger = logging.getLogger(__name__)


def service_worker(request):
    """Serve the portal-only worker with its scope explicitly limited to /portal/."""
    worker_path = settings.BASE_DIR / "static" / "js" / "service-worker.js"
    try:
        worker_source = worker_path.read_text(encoding="utf-8")
    except OSError as exc:
        raise Http404("Service worker file not found") from exc

    response = HttpResponse(worker_source, content_type="application/javascript; charset=utf-8")
    response["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response["Service-Worker-Allowed"] = "/portal/"
    return response

DIRECTOR_TITLE_KEYWORDS = ["Giám đốc"]

def _day_before(date_str_or_date):
    """Trả về ngày liền trước `date_str_or_date` (str 'YYYY-MM-DD' hoặc date) — dùng để tự động
    đóng `effective_to` của thế hệ cấu hình lương/thuế/BH cũ khi thêm thế hệ mới (mục C1)."""
    d = date_str_or_date
    if isinstance(d, str):
        d = datetime.date.fromisoformat(d)
    return d - datetime.timedelta(days=1)


def _paginate(request, queryset, per_page=20, param="page"):
    """Phân trang chung cho các danh sách dài — mục 3.F.

    Trả về đối tượng Page để dùng thẳng trong template (has_next, has_previous,
    number, paginator.num_pages...).
    """
    paginator = Paginator(queryset, per_page)
    page_number = request.GET.get(param) or 1
    try:
        return paginator.page(page_number)
    except Exception:
        return paginator.page(1)

def _friendly_error(request, exc, user_message, log_context=""):
    """Log lỗi kỹ thuật ở server, chỉ hiển thị thông báo thân thiện cho người dùng.

    Mục 3.D: trước đây rất nhiều `except Exception as e` show thẳng
    `f"...: {e}"` (lỗi kỹ thuật gốc) ra UI cho người dùng cuối.
    """
    logger.exception("Lỗi khi xử lý %s (user=%s): %s", log_context or request.path, request.user, exc)
    if isinstance(exc, ValidationError):
        # Lỗi validate dữ liệu (full_clean) vẫn nên cho người dùng biết cụ thể sai ở đâu.
        detail = "; ".join(exc.messages) if hasattr(exc, "messages") else str(exc)
        messages.error(request, f"{user_message} ({detail})")
    else:
        messages.error(request, user_message)

def _director_employees():
    return Employee.objects.filter(position__position_name__icontains="Giám đốc", employment_status="Active")


def _get_employee(request):
    return getattr(request.user, "employee", None)


def _role_code(request):
    return get_role_code(request.user)


def _can_approve(request):
    return has_permission(request.user, "approve_leave")

def _employees_with_permission(permission):
    """Các nhân viên có tài khoản mang vai trò được cấp `permission` (theo PERMISSION_MATRIX).

    Dùng để tìm người cần nhận thông báo/duyệt cho các nghiệp vụ không có một
    người duyệt cố định (đặt phòng, đặt xe, VPP, mua sắm...).
    """
    from apps.security.permissions import PERMISSION_MATRIX
    role_codes = PERMISSION_MATRIX.get(permission, ())
    if not role_codes:
        return Employee.objects.none()
    return Employee.objects.filter(
        user_account__role__role_code__in=role_codes, employment_status="Active",
    ).distinct()

def _leave_approvers_for(employee):
    """Những nhân viên có quyền duyệt đơn nghỉ phép của `employee` (HR/Director toàn công ty + Trưởng phòng của họ)."""
    q = models.Q(user_account__role__role_code__in=["HR_MANAGER", "DIRECTOR"])
    if employee.department_id:
        q |= models.Q(managed_departments=employee.department_id)
    return Employee.objects.filter(q).distinct()

def _approvable_leave_requests(request):
    """Các đơn nghỉ phép mà tài khoản đang đăng nhập THỰC SỰ có quyền duyệt."""
    qs = LeaveRequest.objects.filter(status="Pending").select_related("employee", "leave_type")
    if has_permission(request.user, "approve_leave_all_departments"):
        return qs
    if _role_code(request) == "DEPT_MANAGER":
        employee = _get_employee(request)
        if not employee:
            return qs.none()
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        return qs.filter(employee__department_id__in=dept_ids)
    return qs.none()


@login_required
def dashboard(request):
    employee = _get_employee(request)
    role_code = _role_code(request)
    is_super = request.user.is_superuser
    context = {"employee": employee}

    if has_permission(request.user, "view_own_lifecycle_items"):
        from apps.lifecycle.models import LifecycleChecklistItem
        overdue_items = LifecycleChecklistItem.objects.filter(
            case__status="Open", status="Pending", due_date__lt=timezone.localdate()
        )
        if has_permission(request.user, "manage_lifecycle_case"):
            context["lifecycle_overdue_count"] = overdue_items.count()
        elif employee:
            context["lifecycle_overdue_count"] = overdue_items.filter(assigned_to_user=request.user).count()
        else:
            context["lifecycle_overdue_count"] = 0

    # Shared homepage: show the current month's project activity progress only
    # to roles that can open the activity dashboard. The scoped queryset keeps
    # department managers within their own departments.
    if has_permission(request.user, "view_activity_dashboard"):
        today = timezone.localdate()
        month_start = datetime.date(today.year, today.month, 1)
        month_end = datetime.date(
            today.year, today.month,
            calendar_module.monthrange(today.year, today.month)[1],
        )
        task_prefetch = Prefetch(
            "tasks",
            queryset=Task.objects.order_by("created_date"),
            to_attr="dashboard_tracking_tasks",
        )
        activities = visible_activities_for(request.user).filter(
            Q(start_date__lte=month_end, end_date__gte=month_start)
            | Q(tasks__due_date__gte=month_start, tasks__due_date__lte=month_end)
        ).exclude(status="Cancelled").prefetch_related(task_prefetch).distinct()

        activity_rows = []
        for activity in activities:
            task = activity.dashboard_tracking_tasks[0] if activity.dashboard_tracking_tasks else None
            overdue = bool(
                task and task.status not in ("Done", "Cancelled")
                and task.due_date and task.due_date < today
            )
            activity_rows.append((task, overdue))
        tracked_tasks = [task for task, _ in activity_rows if task]
        context.update({
            "show_project_activity_summary": True,
            "activity_month": month_start.strftime("%Y-%m"),
            "activity_month_label": month_start.strftime("%m/%Y"),
            "activity_total": len(activity_rows),
            "activity_completed": sum(task.status == "Done" for task in tracked_tasks),
            "activity_overdue": sum(overdue for _, overdue in activity_rows),
            "activity_report_missing": sum(
                task.status != "Cancelled" and not task.report_file
                for task in tracked_tasks
            ),
            "activity_without_task": sum(task is None for task, _ in activity_rows),
        })

    # Nhân viên thường: các con số cá nhân
    if employee:
        context["my_pending_leaves"] = LeaveRequest.objects.filter(employee=employee, status="Pending").count()
        context["my_open_tasks"] = Task.objects.filter(assigned_to_employee=employee).exclude(status__in=["Done", "Cancelled"]).count()
        # Mục I8: chỉ hiện lương của kỳ đã "Locked" (đã chốt). Kỳ còn "Open" là số
        # Kế toán đang tính/sửa, show ra cho nhân viên sẽ gây hoang mang khi số đổi qua lại.
        latest_payroll = Payroll.objects.filter(
            employee=employee, payroll_period__status="Locked"
        ).select_related("payroll_period").order_by(
            "-payroll_period__payroll_year", "-payroll_period__payroll_month"
        ).first()
        context["latest_payroll"] = latest_payroll

    # HR Manager / Director / Super Admin: bức tranh toàn công ty
    if is_super or role_code in ("HR_MANAGER", "DIRECTOR"):
        context["is_management_view"] = True
        context["total_employees"] = Employee.objects.filter(employment_status="Active").count()
        context["pending_leave_approvals"] = LeaveRequest.objects.filter(status="Pending").count()
        context["pending_recruitment"] = RecruitmentRequest.objects.filter(status="Pending").count()
        context["pending_documents"] = Document.objects.filter(status="PendingApproval").count()
        context["open_tasks_company"] = Task.objects.exclude(status__in=["Done", "Cancelled"]).count()
        if has_permission(request.user, "view_asset"):
            context["assets_in_use"] = Asset.objects.filter(status="InUse").count()
        dept_counts = (
            Employee.objects.filter(employment_status="Active")
            .values("department__department_name")
            .annotate(total=Count("id"))
            .order_by("-total")
        )
        context["dept_counts"] = dept_counts

    # Trưởng phòng: bức tranh phòng ban mình quản lý
    elif role_code == "DEPT_MANAGER" and employee:
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        context["is_manager_view"] = True
        context["dept_employee_count"] = Employee.objects.filter(department_id__in=dept_ids, employment_status="Active").count()
        context["dept_pending_leaves"] = LeaveRequest.objects.filter(
            employee__department_id__in=dept_ids, status="Pending"
        ).count()
        context["dept_open_tasks"] = Task.objects.filter(
            department_id__in=dept_ids
        ).exclude(status__in=["Done", "Cancelled"]).count()
        if has_permission(request.user, "view_asset"):
            context["dept_assets_in_use"] = Asset.objects.filter(
                allocations__employee__department_id__in=dept_ids, allocations__return_date__isnull=True
            ).distinct().count()

    return render(request, "portal/dashboard.html", context)


@login_required
def profile(request):
    employee = _get_employee(request)
    context = {"employee": employee}
    if employee:
        context.update({
            "transfers": employee.transfers.select_related("from_department", "to_department"),
            "promotions": employee.promotions.select_related("old_position", "new_position"),
            "rewards": employee.rewards.all(),
            "disciplines": employee.disciplines.all(),
            "current_salary": employee.salaries.filter(status="Active").order_by("-effective_date").first(),
            "current_contract": employee.contracts.order_by("-start_date").first(),
        })
    return render(request, "portal/profile.html", context)


@login_required
def org_chart(request):
    employees = list(
        Employee.objects.filter(employment_status="Active")
        .select_related("department", "position")
        .only("id", "full_name", "employee_code", "department__department_name",
              "position__position_name", "manager_id")
        .order_by("full_name", "id")
    )
    nodes = {employee.pk: {"employee": employee, "children": [], "manager": None} for employee in employees}
    parents = {employee.pk: employee.manager_id for employee in employees}
    roots = []
    # Break any pre-existing manager cycle deterministically by making the
    # employee whose manager edge closes the cycle a root.
    for employee in employees:
        parent_id = parents[employee.pk]
        if parent_id not in nodes:
            roots.append(nodes[employee.pk])
            continue
        cursor, visited = parent_id, {employee.pk}
        while cursor in nodes and cursor not in visited:
            visited.add(cursor)
            cursor = parents.get(cursor)
        if cursor in visited:
            roots.append(nodes[employee.pk])
        else:
            nodes[employee.pk]["manager"] = nodes[parent_id]["employee"]
            nodes[parent_id]["children"].append(nodes[employee.pk])
    return render(request, "portal/org_chart.html", {
        "roots": roots,
        "employee_count": len(employees),
        "department_count": len({employee.department_id for employee in employees if employee.department_id}),
    })


@login_required
def profile_change_request_list(request):
    if not has_permission(request.user, "submit_profile_change_request"):
        return HttpResponseForbidden()
    employee = _get_employee(request)
    if not employee:
        return HttpResponseForbidden("Tài khoản chưa được gắn hồ sơ nhân viên.")
    requests_qs = EmployeeProfileChangeRequest.objects.select_related("employee", "decided_by")
    if has_permission(request.user, "approve_profile_change_request"):
        pending = requests_qs.filter(status="Pending")
        own = requests_qs.filter(employee=employee).exclude(status="Pending")
        items = list(pending) + list(own)
    else:
        items = requests_qs.filter(employee=employee)
    return render(request, "portal/profile_change_request_list.html", {"items": items})


@login_required
def profile_change_request_create(request):
    if not has_permission(request.user, "submit_profile_change_request"):
        return HttpResponseForbidden()
    employee = _get_employee(request)
    if not employee:
        return HttpResponseForbidden("Tài khoản chưa được gắn hồ sơ nhân viên.")
    if request.method == "POST":
        field_name = request.POST.get("field_name", "")
        if field_name not in EmployeeProfileChangeRequest.ALLOWED_FIELDS:
            return HttpResponseForbidden("Trường hồ sơ không được phép cập nhật.")
        new_value = request.POST.get("new_value", "").strip()
        if field_name == "mobile_phone" and not new_value:
            messages.error(request, "Số điện thoại không được để trống.")
            return redirect("portal:profile_change_request_create")
        item = EmployeeProfileChangeRequest(
            employee=employee,
            field_name=field_name,
            old_value=str(getattr(employee, field_name) or ""),
            submitted_snapshot=str(getattr(employee, field_name) or ""),
            new_value=new_value,
            reason=request.POST.get("reason", "").strip(),
            attachment=request.FILES.get("attachment"),
            created_by=request.user,
        )
        try:
            item.full_clean()
            item.save()
        except ValidationError as exc:
            messages.error(request, "; ".join(exc.messages))
            return redirect("portal:profile_change_request_create")
        messages.success(request, "Đã gửi yêu cầu cập nhật hồ sơ đến Nhân sự.")
        return redirect("portal:profile_change_request_list")
    return render(request, "portal/profile_change_request_form.html", {"employee": employee})


@login_required
@require_POST
def profile_change_request_decide(request, pk):
    if not has_permission(request.user, "approve_profile_change_request"):
        return HttpResponseForbidden()
    decision = request.POST.get("decision")
    status = {"approve": "Approved", "reject": "Rejected", "more_info": "NeedsMoreInfo"}.get(decision)
    if not status:
        return HttpResponseForbidden("Quyết định không hợp lệ.")
    with transaction.atomic():
        try:
            item = EmployeeProfileChangeRequest.objects.select_for_update().select_related("employee").get(
                pk=pk, status="Pending"
            )
        except EmployeeProfileChangeRequest.DoesNotExist:
            raise Http404
        employee = Employee.objects.select_for_update().get(pk=item.employee_id)
        if str(getattr(employee, item.field_name) or "") != item.submitted_snapshot:
            messages.error(request, "Giá trị hồ sơ đã thay đổi sau khi gửi. Không thể duyệt yêu cầu cũ; vui lòng yêu cầu nhân viên gửi lại.")
            return redirect("portal:profile_change_request_list")
        if status == "Approved":
            setattr(employee, item.field_name, item.new_value)
            employee.modified_by = request.user
            employee.save(update_fields=[item.field_name, "modified_by", "modified_date"])
        item.status = status
        item.decided_by = request.user
        item.decided_at = timezone.now()
        item.decision_note = request.POST.get("note", "").strip()[:500]
        item.modified_by = request.user
        item.save(update_fields=["status", "decided_by", "decided_at", "decision_note", "modified_by", "modified_date"])
        labels = {"Approved": "đã được duyệt", "Rejected": "bị từ chối", "NeedsMoreInfo": "cần bổ sung thông tin"}
        notify(employee, "profile-change", "Yêu cầu cập nhật hồ sơ đã được xử lý",
               f"Yêu cầu thay đổi hồ sơ của bạn {labels[status]}. Vui lòng xem chi tiết trong IHRMS.",
               link="/portal/profile/requests/", created_by=request.user)
    messages.success(request, "Đã ghi nhận quyết định xử lý yêu cầu.")
    return redirect("portal:profile_change_request_list")


@login_required
def profile_change_request_attachment(request, pk):
    employee = _get_employee(request)
    if not (employee and has_permission(request.user, "approve_profile_change_request")):
        if not employee:
            return HttpResponseForbidden()
        item = get_object_or_404(EmployeeProfileChangeRequest, pk=pk, employee=employee)
    else:
        item = get_object_or_404(EmployeeProfileChangeRequest, pk=pk)
    if not item.attachment:
        raise Http404
    response = FileResponse(item.attachment.open("rb"), as_attachment=True, filename=item.attachment.name.rsplit("/", 1)[-1])
    response["X-Content-Type-Options"] = "nosniff"
    return response


@login_required
def attendance_list(request):
    employee = _get_employee(request)
    records = Attendance.objects.filter(employee=employee).order_by("-attendance_date")[:31] if employee else []
    return render(request, "portal/attendance_list.html", {"records": records})


def _leave_balance(employee, year=None):
    """Số dư phép năm và nghỉ bù của nhân viên trong năm."""
    if not employee:
        return None
    year = year or timezone.localdate().year
    balance = LeaveBalance.objects.filter(employee=employee, year=year).first()
    if balance is None:
        balance = LeaveBalance(employee=employee, year=year)
    return balance


@login_required
def leave_list(request):
    employee = _get_employee(request)
    requests_qs = LeaveRequest.objects.filter(employee=employee).select_related("leave_type").order_by("-from_date") if employee else []
    return render(request, "portal/leave_list.html", {
        "requests": requests_qs,
        "balance": _leave_balance(employee),
        "current_year": timezone.localdate().year,
    })


@login_required
def leave_create(request):
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên, không thể tạo đơn.")
        return redirect("portal:leave_list")

    leave_types = LeaveType.objects.all()
    if request.method == "POST":
        leave_type_id = request.POST.get("leave_type")
        from_date = request.POST.get("from_date")
        to_date = request.POST.get("to_date")
        reason = request.POST.get("reason", "")
        try:
            leave_request = LeaveRequest(
                employee=employee,
                leave_type_id=leave_type_id,
                from_date=from_date,
                to_date=to_date,
                reason=reason,
                created_by=request.user,
            )
            leave_request.full_clean()
            leave_request.save()
            for approver in _leave_approvers_for(employee):
                notify(
                    approver, "leave", "Có đơn nghỉ phép chờ duyệt",
                    f"{employee.full_name} xin nghỉ {leave_request.from_date} - {leave_request.to_date}.",
                    link="/portal/leave/approvals/", created_by=request.user,
                )
            messages.success(request, "Đã gửi đơn xin nghỉ phép, chờ phê duyệt.")
            return redirect("portal:leave_list")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo đơn.")

    return render(request, "portal/leave_create.html", {
        "leave_types": leave_types,
        "balance": _leave_balance(employee),
        "current_year": timezone.localdate().year,
    })


@login_required
def leave_approval_list(request):
    """Danh sách đơn chờ duyệt — chỉ Trưởng phòng/HR Manager/Director/Superuser vào được."""
    if not _can_approve(request):
        messages.error(request, "Anh/chị không có quyền duyệt đơn nghỉ phép.")
        return redirect("portal:dashboard")

    qs = _approvable_leave_requests(request)

    return render(request, "portal/leave_approval_list.html", {"requests": qs.order_by("from_date")})


@login_required
@require_POST
def leave_approve(request, pk):
    if not _can_approve(request):
        messages.error(request, "Anh/chị không có quyền duyệt đơn nghỉ phép.")
        return redirect("portal:dashboard")

    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:leave_approval_list")

    approver = _get_employee(request)
    if not approver:
        messages.error(request, "Tài khoản của anh/chị chưa gắn hồ sơ nhân viên, không thể duyệt.")
        return redirect("portal:leave_approval_list")

    # I12 (đã sửa): bọc transaction.atomic() + select_for_update() ngay trên chính đơn nghỉ
    # đang duyệt — trước đây get_object_or_404 lấy đơn từ _approvable_leave_requests() (filter
    # status="Pending") KHÔNG khóa dòng, nên 2 request duyệt gửi gần như đồng thời đều có thể
    # đọc được status="Pending" trước khi bên kia kịp lưu → tạo 2 dòng LeaveApproval trùng cho
    # cùng 1 đơn, status cuối cùng phụ thuộc thứ tự save (race condition).
    with transaction.atomic():
        leave_request = get_object_or_404(
            _approvable_leave_requests(request).select_for_update(), pk=pk
        )

        # I2 (đã sửa): trước đây không có bước nào chặn người duyệt tự duyệt đơn nghỉ của
        # chính mình — Trưởng phòng (DEPT_MANAGER) vốn nằm trong phạm vi _approvable_leave_requests
        # của chính phòng mình nên tự duyệt được đơn của bản thân. HR_MANAGER/DIRECTOR (có quyền
        # approve_leave_all_departments) vẫn được duyệt đơn của mình vì không còn ai cấp trên hơn
        # để duyệt thay — chỉ chặn khi người duyệt là DEPT_MANAGER duyệt đơn của chính mình.
        if leave_request.employee_id == approver.pk and _role_code(request) == "DEPT_MANAGER":
            messages.error(request, "Anh/chị không thể tự duyệt đơn nghỉ phép của chính mình.")
            return redirect("portal:leave_approval_list")

        leave_request.status = decision
        leave_request.modified_by = request.user
        leave_request.save()
        LeaveApproval.objects.create(
            leave_request=leave_request,
            approver=approver,
            decision=decision,
            created_by=request.user,
        )
        label = "Đã duyệt" if decision == "Approved" else "Đã từ chối"
        notify_with_email_draft(
            leave_request.employee,
            "leave",
            f"Đơn nghỉ phép {label.lower()}",
            f"Đơn nghỉ {leave_request.from_date} - {leave_request.to_date} của bạn {label.lower()}.",
            link="/portal/leave/",
            email_subject="IHRMS: Có cập nhật về đơn nghỉ phép của bạn",
            email_body=(
                "Anh/chị có cập nhật mới về yêu cầu nghỉ phép trong IHRMS. "
                "Vui lòng đăng nhập cổng nhân viên để xem trạng thái và chi tiết."
            ),
            created_by=request.user,
            idempotency_key=f"leave-decision:{leave_request.pk}:{decision}",
        )
    return redirect("portal:leave_approval_list")


@login_required
def payslip_list(request):
    employee = _get_employee(request)
    # Mục I8: chỉ hiện phiếu lương của kỳ đã "Locked" — kỳ "Open" là số Kế toán
    # đang tính/chỉnh, chưa phải số cuối cùng, không nên cho nhân viên xem như đã chốt.
    payrolls = Payroll.objects.filter(
        employee=employee, payroll_period__status="Locked", payslip__isnull=False
    ).select_related("payroll_period", "payslip").order_by(
        "-payroll_period__payroll_year", "-payroll_period__payroll_month"
    ) if employee else []
    return render(request, "portal/payslip_list.html", {"payrolls": payrolls})


@login_required
@require_GET
def payslip_pdf(request, pk):
    """A staff member may download only their own issued payslip for a locked period."""
    employee = _get_employee(request)
    if employee is None:
        raise Http404
    payslip = get_object_or_404(
        Payslip.objects.select_related(
            "payroll__employee__department", "payroll__employee__position",
            "payroll__payroll_period",
        ),
        pk=pk, payroll__employee=employee, payroll__payroll_period__status="Locked",
    )
    from apps.payroll.payslip_pdf import render_payslip_pdf
    response = HttpResponse(render_payslip_pdf(payslip), content_type="application/pdf")
    period = payslip.payroll.payroll_period
    response["Content-Disposition"] = f'attachment; filename="phieu-luong-{period.payroll_year}-{period.payroll_month:02d}.pdf"'
    response["Cache-Control"] = "private, no-store"
    response["X-Content-Type-Options"] = "nosniff"
    return response


#=============================================
#Giao việc

from django.utils import timezone
from apps.hrm.models import Employee
from apps.task.models import Task, TaskComment



def _can_assign_task(request):
    return has_permission(request.user, "assign_task")


@login_required
def task_list(request):
    employee = _get_employee(request)
    is_project_overview = has_permission(request.user, "view_activity_dashboard")
    can_assign_task = _can_assign_task(request)
    can_view_task_scopes = can_assign_task or is_project_overview
    requested_scope = request.GET.get("scope")
    if requested_scope in {"assigned_to_me", "assigned_by_me"}:
        task_scope_name = requested_scope
    elif requested_scope == "activity_overview" and is_project_overview:
        task_scope_name = "activity_overview"
    else:
        task_scope_name = "activity_overview" if is_project_overview else "assigned_to_me"

    if task_scope_name == "assigned_by_me":
        # Người giao chỉ xem được những Task do chính họ giao; không mở rộng
        # sang toàn bộ công việc chỉ vì họ có quyền giao việc.
        task_scope = Q(assigned_by_employee=employee) if employee else Q(pk__in=[])
    elif task_scope_name == "activity_overview":
        # Quyền giám sát hoạt động chỉ áp dụng trong phạm vi riêng này,
        # không trộn Task của CBDA vào tab "Việc của tôi".
        task_scope = Q(activity__in=visible_activities_for(request.user))
    else:
        task_scope = Q(assigned_to_employee=employee) if employee else Q(pk__in=[])
    tasks = Task.objects.filter(task_scope).select_related(
        "assigned_to_employee", "assigned_by_employee", "activity"
    ).order_by("due_date", "created_date")
    return render(request, "portal/task_list.html", {
        "tasks": tasks,
        "can_assign_task": can_assign_task,
        "can_view_task_scopes": can_view_task_scopes,
        "is_project_overview": is_project_overview,
        "task_scope_name": task_scope_name,
    })


@login_required
def task_detail(request, pk):
    task = get_object_or_404(Task, pk=pk)
    employee = _get_employee(request)
    is_owner = bool(employee and task.assigned_to_employee_id == employee.id)
    is_assigner = bool(employee and task.assigned_by_employee_id == employee.id)

    if task.activity_id:
        can_view = bool(
            is_owner or is_assigner or request.user.is_superuser
            or can_view_activity(request.user, task.activity)
        )
    elif request.user.is_superuser or is_owner:
        can_view = True
    elif _role_code(request) == "HR_MANAGER":
        can_view = True
    elif _role_code(request) == "DEPT_MANAGER" and employee:
        managed_department_ids = employee.managed_departments.values_list("pk", flat=True)
        can_view = task.assigned_to_employee.department_id in managed_department_ids
    else:
        can_view = is_assigner

    if not can_view:
        messages.error(request, "Anh/chị không có quyền xem công việc này.")
        return redirect("portal:task_list")

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "update_status" and is_owner:
            new_status = request.POST.get("status")
            if new_status in dict(Task.STATUS_CHOICES):
                task.status = new_status
                task.completed_date = timezone.localdate() if new_status == "Done" else None
                try:
                    task.clean()
                    task.save()
                    status_label = dict(Task.STATUS_CHOICES).get(new_status, new_status)
                    notify(
                        task.assigned_by_employee, "task", "Công việc có cập nhật trạng thái",
                        f"{task.task_name}: {status_label}.",
                        link=f"/portal/tasks/{task.pk}/", created_by=request.user,
                    )
                    messages.success(request, "Đã cập nhật trạng thái công việc.")
                except ValidationError as exc:
                    messages.error(request, "; ".join(exc.messages))
        elif action == "upload_activity_report" and is_owner and task.activity_id:
            uploaded = request.FILES.get("report_file")
            if not uploaded:
                messages.error(request, "Chọn tệp báo cáo trước khi tải lên.")
            else:
                try:
                    validate_upload(uploaded)
                    task.report_file = uploaded
                    task.clean()
                    task.save()
                    messages.success(request, "Đã tải báo cáo hoạt động lên kho SharePoint.")
                except (ValidationError, OSError, ValueError, GraphAPIError) as exc:
                    _friendly_error(request, exc, "Không thể tải báo cáo lên kho tài liệu.")
        elif action == "add_comment" and employee:
            text = request.POST.get("comment_text", "").strip()
            if text:
                TaskComment.objects.create(task=task, employee=employee, comment_text=text, created_by=request.user)
                other_party = task.assigned_by_employee if is_owner else task.assigned_to_employee
                if other_party and other_party.pk != employee.pk:
                    notify(
                        other_party, "task", "Có bình luận mới trong công việc",
                        f"{task.task_name}: {text[:150]}",
                        link=f"/portal/tasks/{task.pk}/", created_by=request.user,
                    )
                messages.success(request, "Đã thêm bình luận.")
        return redirect("portal:task_detail", pk=pk)

    return render(request, "portal/task_detail.html", {
        "task": task, "is_owner": is_owner, "comments": task.comments.select_related("employee"),
    })


@login_required
def task_create(request):
    if not _can_assign_task(request):
        messages.error(request, "Anh/chị không có quyền giao việc.")
        return redirect("portal:task_list")

    employee = _get_employee(request)
    role_code = _role_code(request)
    if not request.user.is_superuser and role_code == "DEPT_MANAGER" and employee:
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        candidates = Employee.objects.filter(department_id__in=dept_ids)
    else:
        candidates = Employee.objects.all()

    if request.method == "POST":
        try:
            assignee = candidates.filter(
                pk=request.POST.get("assigned_to_employee")
            ).first()
            if assignee is None:
                raise ValidationError(
                    "Chỉ được giao việc cho nhân viên thuộc phạm vi phòng ban của anh/chị."
                )
            task = Task(
                task_code=request.POST.get("task_code"),
                task_name=request.POST.get("task_name"),
                description=request.POST.get("description", ""),
                assigned_by_employee=employee,
                assigned_to_employee=assignee,
                department_id=request.POST.get("department") or None,
                priority=request.POST.get("priority", "Normal"),
                start_date=request.POST.get("start_date") or timezone.localdate(),
                due_date=request.POST.get("due_date") or None,
                created_by=request.user,
            )
            task.full_clean()
            task.save()
            notify(
                task.assigned_to_employee, "task",
                "Bạn được giao công việc mới",
                f"{task.task_name} — hạn {task.due_date or 'chưa xác định'}.",
                link=f"/portal/tasks/{task.pk}/", created_by=request.user,
            )
            messages.success(request, "Đã giao việc thành công.")
            return redirect("portal:task_list")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo công việc.")

    return render(request, "portal/task_create.html", {
        "candidates": candidates,
        "today": timezone.localdate(),
    })

#=============================================
#KPI

from apps.kpi.models import KPIPeriod, KPITemplate, EmployeeKPI



def _can_score_kpi(request):
    return has_permission(request.user, "score_kpi")


@login_required
def kpi_my_list(request):
    employee = _get_employee(request)
    results = EmployeeKPI.objects.filter(employee=employee).select_related(
        "kpi_template", "kpi_period"
    ).order_by("-kpi_period__year", "-kpi_period__month") if employee else []
    return render(request, "portal/kpi_my_list.html", {
        "results": results, "can_score_kpi": _can_score_kpi(request),
    })


@login_required
def kpi_score_list(request):
    if not _can_score_kpi(request):
        messages.error(request, "Anh/chị không có quyền chấm điểm KPI.")
        return redirect("portal:dashboard")

    qs = EmployeeKPI.objects.select_related("employee", "kpi_template", "kpi_period")
    role_code = _role_code(request)
    if not request.user.is_superuser and role_code == "DEPT_MANAGER":
        employee = _get_employee(request)
        if not employee:
            qs = qs.none()
        else:
            dept_ids = employee.managed_departments.values_list("pk", flat=True)
            qs = qs.filter(employee__department_id__in=dept_ids)

    period_id = request.GET.get("period")
    if period_id:
        qs = qs.filter(kpi_period_id=period_id)

    return render(request, "portal/kpi_score_list.html", {
        "results": qs.order_by("employee__full_name"),
        "periods": KPIPeriod.objects.all().order_by("-year", "-month"),
        "selected_period": period_id,
    })


@login_required
@require_POST
def kpi_score_update(request, pk):
    if not _can_score_kpi(request):
        messages.error(request, "Anh/chị không có quyền chấm điểm KPI.")
        return redirect("portal:dashboard")

    result = get_object_or_404(EmployeeKPI, pk=pk)
    actual = request.POST.get("actual")
    try:
        result.actual = Decimal(actual) if actual not in (None, "") else None
        result.modified_by = request.user
        result.save()
        notify(
            result.employee, "kpi", "KPI của bạn vừa được chấm điểm",
            f"{result.kpi_template} - kỳ {result.kpi_period}: đạt {result.actual}.",
            link="/portal/kpi/", created_by=request.user,
        )
        messages.success(request, f"Đã cập nhật KPI cho {result.employee}.")
    except InvalidOperation:
        messages.error(request, "Giá trị 'Thực tế' không hợp lệ, vui lòng nhập số.")
    except Exception as e:
        _friendly_error(request, e, "Không thể cập nhật.")
    return redirect("portal:kpi_score_list")


@login_required
def kpi_assign_create(request):
    """Tạo chỉ tiêu KPI cho 1 nhân viên trong 1 kỳ — dành cho HR/Trưởng phòng."""
    if not _can_score_kpi(request):
        messages.error(request, "Anh/chị không có quyền tạo chỉ tiêu KPI.")
        return redirect("portal:dashboard")

    employee = _get_employee(request)
    role_code = _role_code(request)
    if not request.user.is_superuser and role_code == "DEPT_MANAGER" and employee:
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        candidates = Employee.objects.filter(department_id__in=dept_ids)
    else:
        candidates = Employee.objects.all()

    if request.method == "POST":
        try:
            kpi_result = EmployeeKPI(
                employee_id=request.POST.get("employee"),
                kpi_template_id=request.POST.get("kpi_template"),
                kpi_period_id=request.POST.get("kpi_period"),
                target=request.POST.get("target") or 0,
                created_by=request.user,
            )
            kpi_result.full_clean()
            kpi_result.save()
            notify(
                kpi_result.employee, "kpi", "Bạn có chỉ tiêu KPI mới",
                f"{kpi_result.kpi_template} - kỳ {kpi_result.kpi_period}: mục tiêu {kpi_result.target}.",
                link="/portal/kpi/", created_by=request.user,
            )
            messages.success(request, "Đã tạo chỉ tiêu KPI.")
            return redirect("portal:kpi_score_list")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo.")

    return render(request, "portal/kpi_assign_create.html", {
        "candidates": candidates,
        "templates": KPITemplate.objects.all(),
        "periods": KPIPeriod.objects.all().order_by("-year", "-month"),
    })


#=============================================
#Tài liệu
from apps.document.models import Document, DocumentApproval



def _can_approve_document(request):
    return has_permission(request.user, "approve_document")


@login_required
def document_list(request):
    documents = Document.objects.filter(status="Approved").select_related("category").order_by("-effective_date")
    return render(request, "portal/document_list.html", {
        "documents": documents,
    })


@login_required
def document_approval_list(request):
    if not _can_approve_document(request):
        messages.error(request, "Anh/chị không có quyền duyệt tài liệu.")
        return redirect("portal:dashboard")
    pending = Document.objects.filter(status="PendingApproval").select_related("category")
    return render(request, "portal/document_approval_list.html", {"documents": pending})


@login_required
@require_POST
def document_approve(request, pk):
    if not _can_approve_document(request):
        messages.error(request, "Anh/chị không có quyền duyệt tài liệu.")
        return redirect("portal:dashboard")

    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:document_approval_list")

    approver = _get_employee(request)
    if not approver:
        messages.error(request, "Tài khoản của anh/chị chưa gắn hồ sơ nhân viên, không thể duyệt.")
        return redirect("portal:document_approval_list")

    # transaction.atomic() + select_for_update(): tránh 2 người bấm duyệt cùng lúc
    # tạo 2 bản ghi DocumentApproval trùng cho cùng 1 tài liệu (mục 3.C).
    with transaction.atomic():
        document = get_object_or_404(
            Document.objects.select_for_update(), pk=pk, status="PendingApproval"
        )
        document.status = decision
        document.save()
        approval = DocumentApproval.objects.create(
            document=document, approver=approver, decision=decision, created_by=request.user
        )
        submitter = None
        if document.created_by_id and hasattr(document.created_by, "employee"):
            submitter = document.created_by.employee
        label = "được duyệt" if decision == "Approved" else "bị từ chối"
        if submitter:
            notify_with_email_draft(
                submitter,
                "document-decision",
                f"Tài liệu {label}",
                f"Tài liệu {document.document_no} đã {label}.",
                link="/portal/documents/",
                email_subject="IHRMS: Có cập nhật về tài liệu bạn gửi",
                email_body=(
                    "Có cập nhật mới về tài liệu anh/chị gửi trong IHRMS. "
                    "Vui lòng đăng nhập cổng nhân viên để xem trạng thái."
                ),
                created_by=request.user,
                idempotency_key=f"document-decision:{approval.pk}:{submitter.pk}",
            )

    if document.status == "Approved":
        for emp in Employee.objects.filter(employment_status="Active"):
            notify(emp, "document", "Tài liệu mới ban hành", document.document_name, link="/portal/documents/")

    label = "Đã ban hành" if decision == "Approved" else "Đã từ chối"
    messages.success(request, f"{label} tài liệu {document.document_no}.")
    return redirect("portal:document_approval_list")


def _visible_dispatches(request):
    """Các công văn mà tài khoản đang đăng nhập ĐƯỢC PHÉP xem.

    Dùng chung cho cả danh sách và trang chi tiết, nếu không trang chi tiết sẽ
    hở lỗ hổng: chỉ cần đổi số pk trên URL là đọc được công văn của phòng khác.
    """
    qs = OfficialDispatch.objects.select_related("department", "assigned_to", "dispatch_book")
    if _role_code(request) in ("PROJECT_MANAGER", "PROJECT_COORDINATOR"):
        return qs.filter(dispatch_type="Outgoing")
    if has_permission(request.user, "access_outgoing_document") or request.user.is_staff:
        return qs
    employee = _get_employee(request)
    if not employee:
        return qs.none()
    return qs.filter(models.Q(assigned_to=employee) | models.Q(department=employee.department))


@login_required
def dispatch_list(request):
    is_mgmt = _can_approve(request) or request.user.is_staff
    qs = _visible_dispatches(request)

    dispatch_type = request.GET.get("type")
    if dispatch_type in ("Incoming", "Outgoing", "InternalDirective"):
        qs = qs.filter(dispatch_type=dispatch_type)

    status = request.GET.get("status")
    if status:
        qs = qs.filter(status=status)

    search = request.GET.get("q")
    if search:
        qs = qs.filter(
            models.Q(dispatch_no__icontains=search) |
            models.Q(title__icontains=search) |
            models.Q(issuing_organization__icontains=search)
        )

    return render(request, "portal/dispatch_list.html", {
        "dispatches": qs,
        "selected_type": dispatch_type or "",
        "selected_status": status or "",
        "search_query": search or "",
        "is_management": is_mgmt,
    })


@login_required
def dispatch_detail(request, pk):
    dispatch = get_object_or_404(
        _visible_dispatches(request).select_related("signer", "reply_to"),
        pk=pk
    )
    trackings = dispatch.trackings.select_related("actor").order_by("-action_date")

    if request.method == "POST":
        action_type = request.POST.get("action_type", "Cập nhật tiến độ")
        comment = request.POST.get("comment", "").strip()
        new_status = request.POST.get("new_status")

        employee = _get_employee(request)
        if not employee:
            messages.error(request, "Tài khoản chưa liên kết hồ sơ nhân viên.")
            return redirect("portal:dispatch_detail", pk=pk)

        # Chỉ người được phân công xử lý hoặc cấp quản lý mới được cập nhật tiến độ.
        is_mgmt = _can_approve(request) or request.user.is_staff
        if not is_mgmt and dispatch.assigned_to_id != employee.pk:
            messages.error(request, "Anh/chị không có quyền cập nhật công văn này.")
            return redirect("portal:dispatch_detail", pk=pk)

        if comment:
            DispatchTracking.objects.create(
                dispatch=dispatch,
                actor=employee,
                action_type=action_type,
                comment=comment,
                created_by=request.user
            )

        if new_status in dict(OfficialDispatch.STATUS_CHOICES):
            dispatch.status = new_status
            dispatch.save(update_fields=["status"])

        # Báo cho "phía bên kia": người xử lý cập nhật thì báo người tạo/giao,
        # quản lý cập nhật thì báo người được giao xử lý.
        status_label = dict(OfficialDispatch.STATUS_CHOICES).get(dispatch.status, dispatch.status)
        recipient = None
        if dispatch.assigned_to_id == employee.pk:
            recipient = dispatch.created_by.employee if hasattr(dispatch.created_by, "employee") else None
        elif dispatch.assigned_to_id:
            recipient = dispatch.assigned_to
        if recipient and recipient.pk != employee.pk:
            notify(
                recipient, "dispatch", "Công văn có cập nhật mới",
                f"{dispatch.dispatch_no} - {dispatch.title}: {status_label}.",
                link=f"/portal/dispatches/{dispatch.pk}/", created_by=request.user,
            )

        messages.success(request, "Đã cập nhật tiến độ xử lý công văn.")
        return redirect("portal:dispatch_detail", pk=pk)

    return render(request, "portal/dispatch_detail.html", {
        "dispatch": dispatch,
        "trackings": trackings,
        "status_choices": OfficialDispatch.STATUS_CHOICES,
    })


from apps.asset.models import AssetAllocation

#=============================================
#Tài sản

@login_required
def asset_list(request):
    if not has_permission(request.user, "view_asset"):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    employee = _get_employee(request)
    allocations = AssetAllocation.objects.filter(
        employee=employee, return_date__isnull=True
    ).select_related("asset", "asset__asset_type") if employee else []
    return render(request, "portal/asset_list.html", {
        "allocations": allocations,
        "can_manage_asset": _can_manage_asset(request),
    })

#=============================================
#Tuyển dụng

from apps.recruitment.models import RecruitmentRequest
from apps.organization.models import Position



def _can_request_recruitment(request):
    return has_permission(request.user, "request_recruitment")


def _can_approve_recruitment(request):
    return has_permission(request.user, "approve_recruitment")


@login_required
def recruitment_list(request):
    employee = _get_employee(request)
    role_code = _role_code(request)

    if request.user.is_superuser or role_code == "HR_MANAGER":
        requests_qs = RecruitmentRequest.objects.all()
    elif role_code == "DEPT_MANAGER" and employee:
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        requests_qs = RecruitmentRequest.objects.filter(department_id__in=dept_ids)
    else:
        requests_qs = RecruitmentRequest.objects.none()

    return render(request, "portal/recruitment_list.html", {
        "requests": requests_qs.select_related("department", "position").order_by("-required_date"),
        "can_request": _can_request_recruitment(request),
        "can_approve": _can_approve_recruitment(request),
    })


@login_required
def recruitment_create(request):
    if not _can_request_recruitment(request):
        messages.error(request, "Anh/chị không có quyền đề xuất tuyển dụng.")
        return redirect("portal:recruitment_list")

    employee = _get_employee(request)
    role_code = _role_code(request)
    if not request.user.is_superuser and role_code == "DEPT_MANAGER" and employee:
        departments = employee.managed_departments.all()
    else:
        from apps.organization.models import Department
        departments = Department.objects.all()

    if request.method == "POST":
        try:
            recruitment_request = RecruitmentRequest(
                request_code=request.POST.get("request_code"),
                department_id=request.POST.get("department"),
                position_id=request.POST.get("position"),
                required_quantity=request.POST.get("required_quantity") or 1,
                request_reason=request.POST.get("request_reason", ""),
                required_date=request.POST.get("required_date"),
                created_by=request.user,
            )
            recruitment_request.full_clean()
            recruitment_request.save()
            messages.success(request, "Đã gửi đề xuất tuyển dụng, chờ Nhân sự phê duyệt.")
            return redirect("portal:recruitment_list")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo đề xuất.")

    return render(request, "portal/recruitment_create.html", {
        "departments": departments, "positions": Position.objects.all(),
    })


@login_required
@require_POST
def recruitment_approve(request, pk):
    if not _can_approve_recruitment(request):
        messages.error(request, "Anh/chị không có quyền duyệt đề xuất tuyển dụng.")
        return redirect("portal:recruitment_list")

    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:recruitment_list")

    with transaction.atomic():
        recruitment_request = get_object_or_404(
            RecruitmentRequest.objects.select_for_update(), pk=pk, status="Pending"
        )
        recruitment_request.status = decision
        recruitment_request.modified_by = request.user
        recruitment_request.save()

    label = "Đã duyệt" if decision == "Approved" else "Đã từ chối"

    notify(
        recruitment_request.created_by.employee if hasattr(recruitment_request.created_by, "employee") else None,
        "recruitment", f"Đề xuất tuyển dụng {label.lower()}",
        f"{recruitment_request.request_code} - {recruitment_request.position} {label.lower()}.",
        link="/portal/recruitment/",
    )
    messages.success(request, f"{label} đề xuất {recruitment_request.request_code}.")
    return redirect("portal:recruitment_list")

@login_required
def notification_list(request):
    employee = _get_employee(request)
    items = []
    if employee:
        items = Notification.objects.filter(employee=employee).prefetch_related(
            Prefetch(
                "email_outbox_entries",
                queryset=NotificationEmailOutbox.objects.filter(status__in=("Draft", "Sent")),
                to_attr="email_drafts",
            )
        )
        Notification.objects.filter(employee=employee, is_read=False).update(is_read=True)
    return render(request, "portal/notification_list.html", {"items": items})


@login_required
@require_POST
def notification_email_outbox_mark_sent(request, pk):
    employee = _get_employee(request)
    with transaction.atomic():
        outbox = get_object_or_404(NotificationEmailOutbox.objects.select_for_update(), pk=pk)
        if not has_permission(request.user, "manage_hr") and (
            not employee or outbox.notification.employee_id != employee.pk
        ):
            return HttpResponseForbidden("Không có quyền xác nhận gửi email này.")
        if outbox.status == "Sent":
            messages.info(request, "Email này đã được xác nhận gửi trước đó.")
            return redirect("portal:notification_list")
        if outbox.status != "Draft":
            messages.error(request, "Chỉ có thể xác nhận email đang ở trạng thái nháp Outlook.")
            return redirect("portal:notification_list")
        outbox.status = "Sent"
        outbox.sent_by = request.user
        outbox.sent_at = timezone.now()
        outbox.save(update_fields=("status", "sent_by", "sent_at", "modified_date"))
    messages.success(request, "Đã ghi nhận anh/chị gửi thư qua Outlook.")
    return redirect("portal:notification_list")

def _leave_days_in_month(employee, code, year, month):
    requests = LeaveRequest.objects.filter(
        employee=employee, status="Approved", leave_type__leave_type_code=code,
        from_date__lte=datetime.date(year, month, calendar_module.monthrange(year, month)[1]),
        to_date__gte=datetime.date(year, month, 1),
    )
    from apps.leave.models import leave_days_in_month
    return sum(
        (leave_days_in_month(request, year, month) for request in requests),
        Decimal("0"),
    )

def _weekday_count(year, month):
    """Số ngày Thứ 2 - Thứ 6 trong tháng (không tính Thứ 7, Chủ nhật)."""
    days_in_month = calendar_module.monthrange(year, month)[1]
    return sum(1 for d in range(1, days_in_month + 1) if calendar_module.weekday(year, month, d) < 5)

def _holiday_weekday_count(year, month):
    """Số ngày lễ rơi vào ngày trong tuần (Thứ 2-6) của tháng — dùng để trừ vào số ngày cần làm."""
    holidays = Holiday.objects.filter(holiday_date__year=year, holiday_date__month=month, is_active=True)
    return sum(1 for h in holidays if h.holiday_date.weekday() < 5)


@login_required
def timesheet_month(request):
    
    employee = _get_employee(request)
    today = timezone.now().date()
    year = parse_int(request.GET.get("year"), today.year, min_value=2000, max_value=2100)
    month = parse_int(request.GET.get("month"), today.month, min_value=1, max_value=12)

    days_in_month = calendar_module.monthrange(year, month)[1]
    holidays = {h.holiday_date.day: h for h in Holiday.objects.filter(
        holiday_date__year=year, holiday_date__month=month, is_active=True
    )}
    entries = TimesheetEntry.objects.filter(
        employee=employee, work_date__year=year, work_date__month=month
    ).select_related("project") if employee else []
    entries_by_day = {}
    for e in entries:
        entries_by_day.setdefault(e.work_date.day, []).append(e)

    # Dựng đủ các dòng trong tháng, phân loại: cuối tuần / ngày lễ / ngày làm việc
    day_rows = []
    for d in range(1, days_in_month + 1):
        weekday = calendar_module.weekday(year, month, d)
        is_weekend = weekday >= 5
        holiday = holidays.get(d)
        day_rows.append({
            "day": d,
            "date": f"{year}-{month:02d}-{d:02d}",
            "is_weekend": is_weekend,
            "holiday": holiday,
            "entries": entries_by_day.get(d, []),
            "day_total": sum((e.hours for e in entries_by_day.get(d, [])), Decimal("0")),
        })

    total_hours = sum((e.hours for e in entries), Decimal("0"))

    summary, _ = (
        MonthlyAttendanceSummary.objects.get_or_create(employee=employee, year=year, month=month)
        if employee else (None, False)
    )

    base_days = _weekday_count(year, month) - _holiday_weekday_count(year, month)
    working_calendar = WorkingCalendar.objects.filter(calendar_year=year).first()
    hours_per_day = working_calendar.working_hours_per_day if working_calendar else Decimal("8")
    leave_total = summary.total_leave_days if summary else Decimal("0")
    required_days = Decimal(base_days) - leave_total
    required_hours = required_days * hours_per_day

    from apps.attendance.models import ProjectTimeAllocation
    projects_used = {e.project for e in entries if e.project}
    allocations = list(ProjectTimeAllocation.objects.filter(
        employee=employee, year=year, month=month, project__in=projects_used
    ).select_related("project")) if employee else []
    # Nếu chỉ làm 1 dự án và chưa có cấu hình % -> mặc định 100%
    if len(projects_used) == 1 and not allocations:
        only_project = list(projects_used)[0]
        allocations = [ProjectTimeAllocation(employee=employee, project=only_project, year=year, month=month, allocation_percent=100)]
    diff_hours = total_hours - required_hours

    balance = LeaveBalance.objects.filter(employee=employee, year=year).first() if employee else None

    context = {
        "day_rows": day_rows, "year": year, "month": month,
        "total_hours": total_hours,
        "base_days": base_days, "hours_per_day": hours_per_day,
        "required_days": required_days, "required_hours": required_hours,
        "allocations": allocations, "diff_hours": diff_hours,
        "summary": summary, "balance": balance,
        "projects": Project.objects.filter(is_active=True),
        "months": range(1, 13), "years": range(today.year - 1, today.year + 2),
    }
    return render(request, "portal/timesheet_month.html", context)


@login_required
@require_POST
def timesheet_entry_create(request):
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:timesheet_month")

    work_date = request.POST.get("work_date", "")
    try:
        entry = TimesheetEntry(
            employee=employee,
            work_date=work_date,
            project_id=request.POST.get("project") or None,
            task_description=request.POST.get("task_description", ""),
            hours=request.POST.get("hours") or 0,
            created_by=request.user,
        )
        entry.full_clean()
        entry.save()
        messages.success(request, "Đã ghi nhận giờ làm việc.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")

    if work_date:
        y, m, _ = work_date.split("-")
        return redirect(f"/portal/timesheet/?year={y}&month={int(m)}")
    return redirect("portal:timesheet_month")


@login_required
@require_POST
def timesheet_entry_delete(request, pk):
    employee = _get_employee(request)
    entry = get_object_or_404(TimesheetEntry, pk=pk, employee=employee)
    year, month = entry.work_date.year, entry.work_date.month
    entry.delete()
    messages.success(request, "Đã xóa dòng chấm công.")
    return redirect(f"/portal/timesheet/?year={year}&month={month}")


@login_required
@require_POST
def timesheet_summary_update(request):

    employee = _get_employee(request)
    today = timezone.now().date()
    year = parse_int(request.POST.get("year"), today.year, min_value=2000, max_value=2100)
    month = parse_int(request.POST.get("month"), today.month, min_value=1, max_value=12)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:timesheet_month")

    summary, _ = MonthlyAttendanceSummary.objects.get_or_create(employee=employee, year=year, month=month)

    # I3 (đã sửa): số liệu này ảnh hưởng trực tiếp yêu cầu giờ công và lương (xem comment ở
    # model MonthlyAttendanceSummary) nên một khi đã được quản lý duyệt, nhân viên không được
    # tự sửa lại nữa — phải liên hệ quản lý để quản lý sửa/duyệt lại.
    if summary.status == MonthlyAttendanceSummary.STATUS_APPROVED:
        messages.error(
            request,
            "Bảng tổng hợp chấm công tháng này đã được duyệt, không thể tự sửa. "
            "Liên hệ quản lý nếu cần điều chỉnh.",
        )
        return redirect(f"/portal/timesheet/?year={year}&month={month}")

    try:
        summary.sick_days = request.POST.get("sick_days") or 0
        summary.maternity_days = request.POST.get("maternity_days") or 0
        summary.compensatory_days = request.POST.get("compensatory_days") or 0
        summary.annual_leave_days = request.POST.get("annual_leave_days") or 0
        summary.unpaid_leave_days = request.POST.get("unpaid_leave_days") or 0
        summary.status = MonthlyAttendanceSummary.STATUS_PENDING
        summary.approved_by = None
        summary.approved_at = None
        summary.modified_by = request.user
        summary.full_clean()
        summary.save()
        messages.success(request, "Đã cập nhật số ngày nghỉ trong tháng — đang chờ quản lý duyệt.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")

    return redirect(f"/portal/timesheet/?year={year}&month={month}")


def _can_approve_timesheet(request):
    return has_permission(request.user, "approve_timesheet")


def _approvable_timesheet_summaries(request):
    """Các bảng tổng hợp chấm công tháng (trạng thái Chờ duyệt) mà tài khoản đang đăng nhập
    thực sự có quyền duyệt — cùng cách phạm vi với `_approvable_leave_requests`."""
    qs = MonthlyAttendanceSummary.objects.filter(
        status=MonthlyAttendanceSummary.STATUS_PENDING
    ).select_related("employee", "employee__department")
    if has_permission(request.user, "approve_timesheet_all_departments"):
        return qs
    if _role_code(request) == "DEPT_MANAGER":
        employee = _get_employee(request)
        if not employee:
            return qs.none()
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        return qs.filter(employee__department_id__in=dept_ids)
    return qs.none()


@login_required
def timesheet_approval_list(request):
    if not _can_approve_timesheet(request):
        messages.error(request, "Anh/chị không có quyền duyệt bảng chấm công.")
        return redirect("portal:dashboard")
    qs = _approvable_timesheet_summaries(request)
    return render(request, "portal/timesheet_approval_list.html", {
        "summaries": qs.order_by("-year", "-month", "employee__full_name"),
    })


@login_required
@require_POST
def timesheet_summary_approve(request, pk):
    if not _can_approve_timesheet(request):
        messages.error(request, "Anh/chị không có quyền duyệt bảng chấm công.")
        return redirect("portal:dashboard")

    with transaction.atomic():
        summary = get_object_or_404(
            _approvable_timesheet_summaries(request).select_for_update(), pk=pk,
        )
        summary.status = MonthlyAttendanceSummary.STATUS_APPROVED
        summary.approved_by = request.user
        summary.approved_at = timezone.now()
        summary.save()
    notify(
        summary.employee, "attendance", "Bảng chấm công tháng đã được duyệt",
        f"Tổng hợp chấm công {summary.month}/{summary.year} của bạn đã được duyệt.",
        link="/portal/timesheet/", created_by=request.user,
    )
    messages.success(request, "Đã duyệt bảng tổng hợp chấm công.")
    return redirect("portal:timesheet_approval_list")



def _can_manage_calendar(request):
    return has_permission(request.user, "manage_calendar")


@login_required
def holiday_list(request):
    if not _can_manage_calendar(request):
        messages.error(request, "Anh/chị không có quyền quản lý ngày lễ.")
        return redirect("portal:dashboard")

    if request.method == "POST":
        try:
            holiday = Holiday(
                holiday_code=request.POST.get("holiday_code"),
                holiday_name=request.POST.get("holiday_name"),
                holiday_date=request.POST.get("holiday_date"),
                created_by=request.user,
            )
            holiday.full_clean()
            holiday.save()
            messages.success(request, "Đã thêm ngày lễ.")
        except Exception as e:
            _friendly_error(request, e, "Không thể thêm.")
        return redirect("portal:holiday_list")

    return render(request, "portal/holiday_list.html", {
        "holidays": Holiday.objects.filter(is_active=True).order_by("holiday_date"),
    })


@login_required
@require_POST
def holiday_delete(request, pk):
    if not _can_manage_calendar(request):
        messages.error(request, "Anh/chị không có quyền quản lý ngày lễ.")
        return redirect("portal:dashboard")
    holiday = get_object_or_404(Holiday, pk=pk)
    holiday.is_active = False
    holiday.save(update_fields=["is_active"])
    messages.success(request, "Đã xóa ngày lễ.")
    return redirect("portal:holiday_list")



def _can_view_reports(request):
    return has_permission(request.user, "view_reports")


@login_required
def report_hub(request):
    if not _can_view_reports(request):
        messages.error(request, "Anh/chị không có quyền xem báo cáo.")
        return redirect("portal:dashboard")
    return render(request, "portal/report_hub.html")


@login_required
def report_timesheet(request):
    if not _can_view_reports(request):
        messages.error(request, "Anh/chị không có quyền xem báo cáo.")
        return redirect("portal:dashboard")

    today = timezone.now().date()
    year = parse_int(request.GET.get("year"), today.year, min_value=2000, max_value=2100)
    month = parse_int(request.GET.get("month"), today.month, min_value=1, max_value=12)

    base_days = _weekday_count(year, month) - _holiday_weekday_count(year, month)
    working_calendar = WorkingCalendar.objects.filter(calendar_year=year).first()
    hours_per_day = working_calendar.working_hours_per_day if working_calendar else Decimal("8")

    rows = []
    for emp in Employee.objects.filter(employment_status="Active").select_related("department").order_by("department__department_name", "full_name"):
        total_hours = TimesheetEntry.objects.filter(
            employee=emp, work_date__year=year, work_date__month=month
        ).aggregate(total=Sum("hours"))["total"] or Decimal("0")
        
        # I3 (đã sửa): báo cáo BCC cho quản lý chỉ tính theo số liệu đã được duyệt
        # (status="Approved"), không dùng thẳng số nhân viên tự khai chưa ai xác nhận.
        summary = MonthlyAttendanceSummary.objects.filter(
            employee=emp, year=year, month=month, status=MonthlyAttendanceSummary.STATUS_APPROVED,
        ).first()
        leave_total = summary.total_leave_days if summary else Decimal("0")
        required_hours = (Decimal(base_days) - leave_total) * hours_per_day
        balance = LeaveBalance.objects.filter(employee=emp, year=year).first()
        has_pending_summary = MonthlyAttendanceSummary.objects.filter(
            employee=emp, year=year, month=month, status=MonthlyAttendanceSummary.STATUS_PENDING,
        ).exists()
        rows.append({
            "employee": emp, "total_hours": total_hours, "required_hours": required_hours,
            "diff_hours": total_hours - required_hours,
            "leave_remaining": balance.remaining_days if balance else None,
            "has_pending_summary": has_pending_summary,
        })

    return render(request, "portal/report_timesheet.html", {
        "rows": rows, "year": year, "month": month,
        "months": range(1, 13), "years": range(today.year - 1, today.year + 2),
    })


@login_required
def report_personnel_changes(request):
    if not _can_view_reports(request):
        messages.error(request, "Anh/chị không có quyền xem báo cáo.")
        return redirect("portal:dashboard")

    today = timezone.now().date()
    default_from = today.replace(day=1)
    from_date = request.GET.get("from_date", default_from.isoformat())
    to_date = request.GET.get("to_date", today.isoformat())

    new_hires = Employee.objects.filter(join_date__range=[from_date, to_date]).select_related("department")
    departures = Employee.objects.filter(
        leave_date__range=[from_date, to_date], employment_status="Resigned"
    ).select_related("department")
    transfers = EmployeeTransfer.objects.filter(
        effective_date__range=[from_date, to_date]
    ).select_related("employee", "from_department", "to_department")
    promotions = EmployeePromotion.objects.filter(
        effective_date__range=[from_date, to_date]
    ).select_related("employee", "old_position", "new_position")

    return render(request, "portal/report_personnel_changes.html", {
        "new_hires": new_hires, "departures": departures,
        "transfers": transfers, "promotions": promotions,
        "from_date": from_date, "to_date": to_date,
    })

def _can_approve_office_request(request):
    return has_permission(request.user, "approve_office_request")


def _office_request_approvers():
    return _employees_with_permission("approve_office_request")


# --- Đặt phòng họp ---
@login_required
def room_booking_list(request):
    employee = _get_employee(request)
    upcoming = RoomBooking.objects.filter(
        status__in=["Pending", "Approved"], end_time__gte=timezone.now(),
    ).select_related("room", "employee").order_by("start_time")
    return render(request, "portal/room_booking_list.html", {
        "upcoming": upcoming, "rooms": MeetingRoom.objects.filter(is_active=True),
    })


@login_required
@require_POST
def room_booking_create(request):
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:room_booking_list")
    try:
        booking = RoomBooking(
            room_id=request.POST.get("room"), employee=employee,
            purpose=request.POST.get("purpose"),
            start_time=request.POST.get("start_time"), end_time=request.POST.get("end_time"),
            status="Pending", created_by=request.user,
        )
        booking.full_clean()
        booking.save()
        for approver in _office_request_approvers():
            notify(
                approver, "office", "Có yêu cầu đặt phòng họp chờ duyệt",
                f"{employee.full_name} đề nghị đặt {booking.room} "
                f"({booking.start_time:%d/%m %H:%M} - {booking.end_time:%H:%M}).",
                link="/portal/rooms/approvals/", created_by=request.user,
            )
        messages.success(request, "Đã gửi yêu cầu đặt phòng, chờ duyệt.")
    except Exception as e:
        _friendly_error(request, e, "Không thể đặt phòng.")
    return redirect("portal:room_booking_list")


@login_required
@require_POST
def room_booking_cancel(request, pk):
    employee = _get_employee(request)
    booking = get_object_or_404(RoomBooking, pk=pk, employee=employee)
    booking.status = "Cancelled"
    booking.save(update_fields=["status"])
    messages.success(request, "Đã hủy đặt phòng.")
    return redirect("portal:room_booking_list")


@login_required
def room_booking_approval_list(request):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đặt phòng họp.")
        return redirect("portal:dashboard")
    pending = RoomBooking.objects.filter(status="Pending").select_related("room", "employee").order_by("start_time")
    return render(request, "portal/room_booking_approval_list.html", {"bookings": pending})


@login_required
@require_POST
def room_booking_approve(request, pk):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đặt phòng họp.")
        return redirect("portal:dashboard")
    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:room_booking_approval_list")
    try:
        # transaction.atomic() + select_for_update(): tránh 2 người bấm duyệt
        # cùng lúc dẫn tới 2 phòng chồng giờ đều được duyệt (mục 3.C).
        with transaction.atomic():
            booking = get_object_or_404(
                RoomBooking.objects.select_for_update(), pk=pk, status="Pending"
            )
            booking.status = decision
            booking.modified_by = request.user
            booking.full_clean()
            booking.save()
        label = "duyệt" if decision == "Approved" else "từ chối"
        notify(
            booking.employee, "office", f"Yêu cầu đặt phòng họp đã được {label}",
            f"Đặt {booking.room} ({booking.start_time:%d/%m %H:%M} - {booking.end_time:%H:%M}) đã {label}.",
            link="/portal/rooms/", created_by=request.user,
        )
        messages.success(request, f"Đã {label} yêu cầu đặt phòng.")
    except Exception as e:
        _friendly_error(request, e, "Không thể xử lý yêu cầu đặt phòng.")
    return redirect("portal:room_booking_approval_list")


# --- Đặt xe ---
@login_required
def vehicle_booking_list(request):
    upcoming = VehicleBooking.objects.filter(
        status__in=["Pending", "Approved"], end_time__gte=timezone.now(),
    ).select_related("vehicle", "employee").order_by("start_time")
    return render(request, "portal/vehicle_booking_list.html", {
        "upcoming": upcoming, "vehicles": Vehicle.objects.filter(is_active=True),
    })


@login_required
@require_POST
def vehicle_booking_create(request):
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:vehicle_booking_list")
    try:
        booking = VehicleBooking(
            vehicle_id=request.POST.get("vehicle"), employee=employee,
            destination=request.POST.get("destination"), purpose=request.POST.get("purpose", ""),
            start_time=request.POST.get("start_time"), end_time=request.POST.get("end_time"),
            status="Pending", created_by=request.user,
        )
        booking.full_clean()
        booking.save()
        for approver in _office_request_approvers():
            notify(
                approver, "office", "Có yêu cầu đặt xe chờ duyệt",
                f"{employee.full_name} đề nghị đặt xe đi {booking.destination} "
                f"({booking.start_time:%d/%m %H:%M} - {booking.end_time:%H:%M}).",
                link="/portal/vehicles/approvals/", created_by=request.user,
            )
        messages.success(request, "Đã gửi yêu cầu đặt xe, chờ duyệt.")
    except Exception as e:
        _friendly_error(request, e, "Không thể đặt xe.")
    return redirect("portal:vehicle_booking_list")


@login_required
@require_POST
def vehicle_booking_cancel(request, pk):
    employee = _get_employee(request)
    booking = get_object_or_404(VehicleBooking, pk=pk, employee=employee)
    booking.status = "Cancelled"
    booking.save(update_fields=["status"])
    messages.success(request, "Đã hủy đặt xe.")
    return redirect("portal:vehicle_booking_list")


@login_required
def vehicle_booking_approval_list(request):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đặt xe.")
        return redirect("portal:dashboard")
    pending = VehicleBooking.objects.filter(status="Pending").select_related("vehicle", "employee").order_by("start_time")
    return render(request, "portal/vehicle_booking_approval_list.html", {"bookings": pending})


@login_required
@require_POST
def vehicle_booking_approve(request, pk):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đặt xe.")
        return redirect("portal:dashboard")
    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:vehicle_booking_approval_list")
    try:
        with transaction.atomic():
            booking = get_object_or_404(
                VehicleBooking.objects.select_for_update(), pk=pk, status="Pending"
            )
            booking.status = decision
            booking.modified_by = request.user
            booking.full_clean()
            booking.save()
        label = "duyệt" if decision == "Approved" else "từ chối"
        notify(
            booking.employee, "office", f"Yêu cầu đặt xe đã được {label}",
            f"Đặt xe đi {booking.destination} ({booking.start_time:%d/%m %H:%M} - {booking.end_time:%H:%M}) đã {label}.",
            link="/portal/vehicles/", created_by=request.user,
        )
        messages.success(request, f"Đã {label} yêu cầu đặt xe.")
    except Exception as e:
        _friendly_error(request, e, "Không thể xử lý yêu cầu đặt xe.")
    return redirect("portal:vehicle_booking_approval_list")


# --- Đề nghị VPP ---
@login_required
def supply_request_list(request):
    employee = _get_employee(request)
    requests_qs = SupplyRequest.objects.filter(employee=employee).prefetch_related("items__item").order_by("-request_date") if employee else []
    return render(request, "portal/supply_request_list.html", {
        "requests": requests_qs, "items": SupplyItem.objects.all(),
    })


@login_required
@require_POST
def supply_request_create(request):
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:supply_request_list")

    item_ids = request.POST.getlist("item")
    quantities = request.POST.getlist("quantity")
    valid_lines = [(i, q) for i, q in zip(item_ids, quantities) if i and q and int(q) > 0]

    if not valid_lines:
        messages.error(request, "Vui lòng chọn ít nhất 1 vật tư với số lượng > 0.")
        return redirect("portal:supply_request_list")

    try:
        with transaction.atomic():
            supply_request = SupplyRequest.objects.create(
                request_no=f"VPP-{employee.employee_code}-{timezone.now():%Y%m%d%H%M%S}",
                employee=employee, request_date=timezone.now().date(), created_by=request.user,
            )
            for item_id, qty in valid_lines:
                SupplyRequestItem.objects.create(
                    supply_request=supply_request, item_id=item_id, quantity_requested=qty, created_by=request.user,
                )
        for approver in _office_request_approvers():
            notify(
                approver, "office", "Có đề nghị cấp VPP chờ duyệt",
                f"{employee.full_name} gửi đề nghị {supply_request.request_no}.",
                link="/portal/supplies/approvals/", created_by=request.user,
            )
        messages.success(request, "Đã gửi đề nghị cấp vật tư, chờ Hành chính duyệt.")
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo đề nghị.")
    return redirect("portal:supply_request_list")


@login_required
def supply_request_approval_list(request):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đề nghị VPP.")
        return redirect("portal:dashboard")
    pending = SupplyRequest.objects.filter(status="Pending").select_related("employee").prefetch_related("items__item").order_by("request_date")
    approved = SupplyRequest.objects.filter(status="Approved").select_related("employee").prefetch_related("items__item").order_by("request_date")
    return render(request, "portal/supply_request_approval_list.html", {
        "pending": pending, "approved": approved,
    })


@login_required
@require_POST
def supply_request_approve(request, pk):
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền duyệt đề nghị VPP.")
        return redirect("portal:dashboard")
    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:supply_request_approval_list")
    try:
        with transaction.atomic():
            supply_request = get_object_or_404(
                SupplyRequest.objects.select_for_update(), pk=pk, status="Pending"
            )
            supply_request.status = decision
            supply_request.modified_by = request.user
            supply_request.save(update_fields=["status", "modified_by"])
        label = "duyệt" if decision == "Approved" else "từ chối"
        notify(
            supply_request.employee, "office", f"Đề nghị cấp VPP đã được {label}",
            f"Đề nghị {supply_request.request_no} đã {label}.",
            link="/portal/supplies/", created_by=request.user,
        )
        messages.success(request, f"Đã {label} đề nghị {supply_request.request_no}.")
    except Exception as e:
        _friendly_error(request, e, "Không thể xử lý đề nghị.")
    return redirect("portal:supply_request_approval_list")


@login_required
@require_POST
def supply_request_issue(request, pk):
    """Xuất kho cho đề nghị đã duyệt — trừ tồn kho theo từng dòng vật tư.

    Bọc transaction.atomic() + select_for_update() trên cả SupplyRequest lẫn
    SupplyItem để 2 người bấm "Xuất kho" cùng lúc không trừ tồn kho 2 lần
    hoặc xuất vượt tồn (mục 3.C).
    """
    if not _can_approve_office_request(request):
        messages.error(request, "Anh/chị không có quyền cấp phát VPP.")
        return redirect("portal:dashboard")
    try:
        with transaction.atomic():
            supply_request = get_object_or_404(
                SupplyRequest.objects.select_for_update(), pk=pk, status="Approved"
            )
            lines = list(supply_request.items.select_related("item").select_for_update())
            items_locked = {
                item.pk: item
                for item in SupplyItem.objects.select_for_update().filter(
                    pk__in=[line.item_id for line in lines]
                )
            }
            for line in lines:
                item = items_locked[line.item_id]
                to_issue = line.quantity_requested - line.quantity_issued
                if to_issue > 0 and item.quantity_in_stock < to_issue:
                    raise ValidationError(
                        f"Không đủ tồn kho cho {item.item_name} (còn {item.quantity_in_stock}, cần {to_issue})."
                    )
            for line in lines:
                item = items_locked[line.item_id]
                to_issue = line.quantity_requested - line.quantity_issued
                if to_issue <= 0:
                    continue
                item.quantity_in_stock -= to_issue
                item.save(update_fields=["quantity_in_stock"])
                line.quantity_issued = line.quantity_requested
                line.save(update_fields=["quantity_issued"])
            supply_request.status = "Issued"
            supply_request.modified_by = request.user
            supply_request.save(update_fields=["status", "modified_by"])
        notify(
            supply_request.employee, "office", "Đã cấp phát vật tư",
            f"Đề nghị {supply_request.request_no} đã được cấp phát.",
            link="/portal/supplies/", created_by=request.user,
        )
        messages.success(request, f"Đã cấp phát vật tư cho đề nghị {supply_request.request_no}.")
    except Exception as e:
        _friendly_error(request, e, "Không thể cấp phát vật tư.")
    return redirect("portal:supply_request_approval_list")


def _can_request_procurement(request):
    return has_permission(request.user, "request_procurement")


def _can_approve_procurement(request):
    return has_permission(request.user, "approve_procurement")


# --- Đề nghị mua sắm ---
@login_required
def procurement_request_list(request):
    employee = _get_employee(request)
    if _can_approve_procurement(request):
        requests_qs = ProcurementRequest.objects.select_related(
            "requested_by", "department"
        ).prefetch_related("items").order_by("-request_date")
    elif employee:
        requests_qs = ProcurementRequest.objects.filter(requested_by=employee).select_related(
            "requested_by", "department"
        ).prefetch_related("items").order_by("-request_date")
    else:
        requests_qs = ProcurementRequest.objects.none()
    return render(request, "portal/procurement_request_list.html", {
        "requests": requests_qs,
        "can_create": _can_request_procurement(request) and employee is not None,
        "can_approve": _can_approve_procurement(request),
    })


@login_required
@require_POST
def procurement_request_create(request):
    if not _can_request_procurement(request):
        messages.error(request, "Anh/chị không có quyền đề nghị mua sắm.")
        return redirect("portal:procurement_request_list")
    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:procurement_request_list")
    if not employee.department_id:
        messages.error(request, "Hồ sơ nhân viên chưa gắn phòng ban.")
        return redirect("portal:procurement_request_list")

    item_names = request.POST.getlist("item_name")
    quantities = request.POST.getlist("quantity")
    costs = request.POST.getlist("estimated_cost")
    valid_lines = [
        (name.strip(), qty, cost) for name, qty, cost in zip(item_names, quantities, costs)
        if name.strip() and qty and qty.isdigit() and int(qty) > 0
    ]
    if not valid_lines:
        messages.error(request, "Vui lòng nhập ít nhất 1 hạng mục với số lượng > 0.")
        return redirect("portal:procurement_request_list")

    try:
        with transaction.atomic():
            proc_request = ProcurementRequest.objects.create(
                request_no=f"MS-{employee.employee_code}-{timezone.now():%Y%m%d%H%M%S}",
                requested_by=employee, department=employee.department,
                request_date=timezone.now().date(), reason=request.POST.get("reason", ""),
                created_by=request.user,
            )
            for name, qty, cost in valid_lines:
                ProcurementItem.objects.create(
                    procurement_request=proc_request, item_name=name, quantity=qty,
                    estimated_cost=cost or None, created_by=request.user,
                )
        for approver in _employees_with_permission("approve_procurement"):
            notify(
                approver, "procurement", "Có đề nghị mua sắm chờ duyệt",
                f"{employee.full_name} ({proc_request.department}) gửi đề nghị {proc_request.request_no}.",
                link="/portal/procurement/", created_by=request.user,
            )
        messages.success(request, f"Đã gửi đề nghị mua sắm {proc_request.request_no}, chờ duyệt.")
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo đề nghị mua sắm.")
    return redirect("portal:procurement_request_list")


@login_required
@require_POST
def procurement_request_approve(request, pk):
    if not _can_approve_procurement(request):
        messages.error(request, "Anh/chị không có quyền duyệt đề nghị mua sắm.")
        return redirect("portal:procurement_request_list")
    decision = request.POST.get("decision")
    if decision not in ("Approved", "Rejected"):
        messages.error(request, "Yêu cầu không hợp lệ.")
        return redirect("portal:procurement_request_list")
    approver = _get_employee(request)
    if not approver:
        messages.error(request, "Tài khoản của anh/chị chưa gắn hồ sơ nhân viên, không thể duyệt.")
        return redirect("portal:procurement_request_list")

    try:
        with transaction.atomic():
            proc_request = get_object_or_404(
                ProcurementRequest.objects.select_for_update(), pk=pk, status="Pending"
            )
            proc_request.status = decision
            proc_request.modified_by = request.user
            proc_request.save(update_fields=["status", "modified_by"])
            ProcurementApproval.objects.create(
                procurement_request=proc_request, approver=approver, decision=decision, created_by=request.user,
            )
        label = "duyệt" if decision == "Approved" else "từ chối"
        notify(
            proc_request.requested_by, "procurement", f"Đề nghị mua sắm đã được {label}",
            f"Đề nghị {proc_request.request_no} đã {label}.",
            link="/portal/procurement/", created_by=request.user,
        )
        messages.success(request, f"Đã {label} đề nghị {proc_request.request_no}.")
    except Exception as e:
        _friendly_error(request, e, "Không thể xử lý đề nghị mua sắm.")
    return redirect("portal:procurement_request_list")


@login_required
@require_POST
def procurement_request_mark_purchased(request, pk):
    if not _can_approve_procurement(request):
        messages.error(request, "Anh/chị không có quyền cập nhật đề nghị mua sắm.")
        return redirect("portal:procurement_request_list")
    try:
        with transaction.atomic():
            proc_request = get_object_or_404(
                ProcurementRequest.objects.select_for_update(), pk=pk, status="Approved"
            )
            proc_request.status = "Purchased"
            proc_request.modified_by = request.user
            proc_request.save(update_fields=["status", "modified_by"])
        messages.success(request, f"Đã đánh dấu {proc_request.request_no} là đã mua.")
    except Exception as e:
        _friendly_error(request, e, "Không thể cập nhật đề nghị mua sắm.")
    return redirect("portal:procurement_request_list")



def _can_create_event(request):
    return has_permission(request.user, "create_event")


@login_required
def event_list(request):
    employee = _get_employee(request)
    today = timezone.now()
    if employee:
        dept_ids = list(employee.department_id and [employee.department_id] or [])
        events = Event.objects.filter(
            end_time__gte=today
        ).filter(models.Q(department__isnull=True) | models.Q(department_id__in=dept_ids)).order_by("start_time")
    else:
        events = Event.objects.filter(end_time__gte=today, department__isnull=True).order_by("start_time")

    if request.method == "POST":
        if not _can_create_event(request):
            messages.error(request, "Anh/chị không có quyền tạo sự kiện.")
            return redirect("portal:event_list")
        try:
            event = Event(
                title=request.POST.get("title"),
                event_type=request.POST.get("event_type", "Meeting"),
                department_id=request.POST.get("department") or None,
                start_time=request.POST.get("start_time"),
                end_time=request.POST.get("end_time"),
                location=request.POST.get("location", ""),
                description=request.POST.get("description", ""),
                created_by_employee=employee,
                created_by=request.user,
            )
            event.full_clean()
            event.save()
            if event.department_id:
                recipients = Employee.objects.filter(department_id=event.department_id, employment_status="Active")
            else:
                recipients = Employee.objects.filter(employment_status="Active")
            for emp in recipients:
                notify(
                    emp, "event", "Sự kiện mới",
                    f"{event.title} — {event.start_time:%d/%m/%Y %H:%M} tại {event.location or 'chưa rõ địa điểm'}.",
                    link="/portal/events/", created_by=request.user,
                )
            messages.success(request, "Đã tạo sự kiện.")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo.")
        return redirect("portal:event_list")

    my_departments = employee.managed_departments.all() if employee and _role_code(request) == "DEPT_MANAGER" else Department.objects.all()
    return render(request, "portal/event_list.html", {
        "events": events, "can_create_event": _can_create_event(request), "departments": my_departments,
    })


@login_required
def training_my_list(request):
    employee = _get_employee(request)
    enrollments = TrainingEnrollment.objects.filter(employee=employee).select_related("course") if employee else []
    return render(request, "portal/training_my_list.html", {"enrollments": enrollments})


@login_required
def report_headcount(request):
    if not _can_view_reports(request):
        messages.error(request, "Anh/chị không có quyền xem báo cáo.")
        return redirect("portal:dashboard")
    today = timezone.now().date()
    year = parse_int(request.GET.get("year"), today.year, min_value=2000, max_value=2100)
    plans = HeadcountPlan.objects.filter(year=year).select_related("department")
    return render(request, "portal/report_headcount.html", {
        "plans": plans, "year": year, "years": range(today.year - 1, today.year + 2),
    })


@login_required
def report_budget(request):
    if not _can_view_reports(request):
        messages.error(request, "Anh/chị không có quyền xem báo cáo.")
        return redirect("portal:dashboard")
    today = timezone.now().date()
    year = parse_int(request.GET.get("year"), today.year, min_value=2000, max_value=2100)
    budgets = AnnualBudget.objects.filter(year=year).select_related("department")

    rows = []
    for b in budgets:
        actual = Decimal("0")
        if b.category == "Procurement":
            from apps.procurement.models import ProcurementRequest
            qs = ProcurementRequest.objects.filter(request_date__year=year, status__in=["Approved", "Purchased"])
            if b.department:
                qs = qs.filter(department=b.department)
            actual = sum((pr.total_estimated_cost for pr in qs), Decimal("0"))
        elif b.category == "Training":
            from apps.planning.models import TrainingCourse
            qs = TrainingCourse.objects.filter(start_date__year=year)
            actual = sum((c.cost for c in qs), Decimal("0"))
        rows.append({"budget": b, "actual": actual, "remaining": b.allocated_amount - actual})

    return render(request, "portal/report_budget.html", {"rows": rows, "year": year, "years": range(today.year - 1, today.year + 2)})

#====================================================================================================================================
#Portal — Quản lý Hồ sơ Nhân viên



def _can_manage_hr(request):
    return has_permission(request.user, "manage_hr")


def _employee_scope(request):
    """Trả về queryset Employee theo đúng quyền xem của tài khoản đang đăng nhập."""
    if has_permission(request.user, "view_all_employees"):
        return Employee.objects.all()
    if _role_code(request) == "DEPT_MANAGER":
        employee = _get_employee(request)
        if not employee:
            return Employee.objects.none()
        dept_ids = employee.managed_departments.values_list("pk", flat=True)
        return Employee.objects.filter(department_id__in=dept_ids)
    return Employee.objects.none()


@login_required
def hr_employee_list(request):
    q = request.GET.get("q", "").strip()
    employees = _employee_scope(request).select_related("department", "position").order_by("full_name")
    if q:
        employees = employees.filter(models.Q(full_name__icontains=q) | models.Q(employee_code__icontains=q))
    return render(request, "portal/hr_employee_list.html", {
        "employees": employees, "q": q, "can_manage_hr": _can_manage_hr(request),
    })


@login_required
def hr_employee_detail(request, pk):
    employee = get_object_or_404(_employee_scope(request), pk=pk)
    return render(request, "portal/hr_employee_detail.html", {
        "employee": employee, "can_manage_hr": _can_manage_hr(request),
        "contracts": employee.contracts.select_related("contract_type").order_by("-start_date"),
        "current_contract": employee.contracts.order_by("-start_date").first(),
        "transfers": employee.transfers.select_related("from_department", "to_department"),
        "promotions": employee.promotions.select_related("old_position", "new_position"),
        "rewards": employee.rewards.all(),
        "disciplines": employee.disciplines.all(),
        "salaries": employee.salaries.order_by("-effective_date"),
        "departments": Department.objects.filter(is_active=True),
        "positions": Position.objects.all(),
    })


def _save_employee_form(request, employee):
    employee.employee_code = request.POST.get("employee_code")
    employee.first_name = request.POST.get("first_name")
    employee.last_name = request.POST.get("last_name")
    employee.full_name = request.POST.get("full_name")
    employee.date_of_birth = request.POST.get("date_of_birth")
    employee.gender = request.POST.get("gender")
    employee.national_id = request.POST.get("national_id")
    employee.tax_code = request.POST.get("tax_code", "")
    employee.social_insurance_no = request.POST.get("social_insurance_no", "")
    employee.personal_email = request.POST.get("personal_email", "")
    employee.company_email = request.POST.get("company_email")
    employee.mobile_phone = request.POST.get("mobile_phone")
    employee.address = request.POST.get("address", "")
    employee.join_date = request.POST.get("join_date")
    employee.leave_date = request.POST.get("leave_date") or None
    employee.department_id = request.POST.get("department")
    employee.position_id = request.POST.get("position")
    employee.job_title_id = request.POST.get("job_title") or None
    employee.manager_id = request.POST.get("manager") or None
    employee.employment_status = request.POST.get("employment_status", "Active")
    employee.full_clean()
    employee.save()


@login_required
def hr_employee_create(request):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền quản lý hồ sơ nhân viên.")
        return redirect("portal:hr_employee_list")

    if request.method == "POST":
        try:
            employee = Employee(created_by=request.user)
            _save_employee_form(request, employee)
            messages.success(request, "Đã tạo hồ sơ nhân viên mới.")
            return redirect("portal:hr_employee_detail", pk=employee.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo hồ sơ.")

    initial_fields = (
        "employee_code", "first_name", "last_name", "full_name",
        "personal_email", "mobile_phone", "address", "join_date",
    )
    initial = {field: request.GET.get(field, "") for field in initial_fields}
    if not initial.get("employee_code"):
        initial["employee_code"] = _generate_next_employee_code()
    if "employee_code" not in initial or not initial.get("employee_code"):
        initial["employee_code"] = _generate_next_employee_code()
    return render(request, "portal/hr_employee_form.html", {
        "employee": None, "initial": initial,
        "departments": Department.objects.filter(is_active=True), "positions": Position.objects.all(),
        "job_titles": JobTitle.objects.all(), "managers": Employee.objects.filter(employment_status="Active"),
    })


@login_required
def hr_employee_edit(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền quản lý hồ sơ nhân viên.")
        return redirect("portal:hr_employee_list")

    employee = get_object_or_404(Employee, pk=pk)
    if request.method == "POST":
        try:
            employee.modified_by = request.user
            _save_employee_form(request, employee)
            messages.success(request, "Đã cập nhật hồ sơ nhân viên.")
            return redirect("portal:hr_employee_detail", pk=employee.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể cập nhật.")

    initial = {
        field: "" for field in (
            "employee_code", "first_name", "last_name", "full_name",
            "personal_email", "mobile_phone", "address", "join_date",
        )
    }
    return render(request, "portal/hr_employee_form.html", {
        "employee": employee, "initial": initial,
        "departments": Department.objects.filter(
            Q(is_active=True) | Q(pk=employee.department_id)
        ),
        "current_department_inactive": not employee.department.is_active,
        "positions": Position.objects.all(),
        "job_titles": JobTitle.objects.all(),
        "managers": Employee.objects.filter(employment_status="Active").exclude(pk=employee.pk),
    })


@login_required
@require_POST
def hr_employee_history_create(request, pk, history_type):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền thực hiện thao tác này.")
        return redirect("portal:hr_employee_list")

    employee = get_object_or_404(Employee, pk=pk)
    try:
        if history_type == "transfer":
            transfer = EmployeeTransfer.objects.create(
                employee=employee, from_department=employee.department,
                to_department_id=request.POST.get("to_department"),
                effective_date=request.POST.get("effective_date"),
                reason=request.POST.get("reason", ""), created_by=request.user,
            )
            notify(
                employee, "hr", "Bạn có quyết định điều chuyển",
                f"Điều chuyển từ {transfer.from_department} sang {transfer.to_department}, hiệu lực {transfer.effective_date}.",
                link=f"/portal/hr/employees/{pk}/", created_by=request.user,
            )
            messages.success(request, "Đã ghi nhận điều chuyển. Phòng ban của nhân viên đã được cập nhật.")
        elif history_type == "promotion":
            promotion = EmployeePromotion.objects.create(
                employee=employee, old_position=employee.position,
                new_position_id=request.POST.get("new_position"),
                effective_date=request.POST.get("effective_date"),
                reason=request.POST.get("reason", ""), created_by=request.user,
            )
            notify(
                employee, "hr", "Bạn có quyết định bổ nhiệm/thăng chức",
                f"Từ {promotion.old_position} lên {promotion.new_position}, hiệu lực {promotion.effective_date}.",
                link=f"/portal/hr/employees/{pk}/", created_by=request.user,
            )
            messages.success(request, "Đã ghi nhận bổ nhiệm/thăng chức. Vị trí của nhân viên đã được cập nhật.")
        elif history_type == "reward":
            reward = EmployeeReward.objects.create(
                employee=employee, reward_type=request.POST.get("reward_type"),
                reward_date=request.POST.get("reward_date"),
                amount=request.POST.get("amount") or None,
                description=request.POST.get("description", ""), created_by=request.user,
            )
            notify(
                employee, "hr", "Bạn nhận được khen thưởng",
                f"{reward.reward_type} ngày {reward.reward_date}.",
                link=f"/portal/hr/employees/{pk}/", created_by=request.user,
            )
            messages.success(request, "Đã ghi nhận khen thưởng.")
        elif history_type == "discipline":
            discipline = EmployeeDiscipline.objects.create(
                employee=employee, discipline_type=request.POST.get("discipline_type"),
                discipline_date=request.POST.get("discipline_date"),
                description=request.POST.get("description", ""), created_by=request.user,
            )
            notify(
                employee, "hr", "Bạn có quyết định kỷ luật",
                f"{discipline.discipline_type} ngày {discipline.discipline_date}.",
                link=f"/portal/hr/employees/{pk}/", created_by=request.user,
            )
            messages.success(request, "Đã ghi nhận kỷ luật.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect("portal:hr_employee_detail", pk=pk)

#===============================================================================================
#Portal — Hợp đồng lao động

@login_required
def hr_contract_list(request):
    if not _can_manage_hr(request):
        return HttpResponseForbidden("Bạn không có quyền xem danh sách hợp đồng nhân sự.")
    contracts = Contract.objects.select_related("employee", "contract_type").order_by("-start_date")
    return render(request, "portal/hr_contract_list.html", {"contracts": contracts, "can_manage_hr": _can_manage_hr(request)})


@login_required
def hr_contract_create(request):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền quản lý hợp đồng.")
        return redirect("portal:hr_contract_list")

    if request.method == "POST":
        try:
            contract = Contract(
                employee_id=request.POST.get("employee"),
                contract_type_id=request.POST.get("contract_type"),
                contract_no=request.POST.get("contract_no"),
                start_date=request.POST.get("start_date"),
                end_date=request.POST.get("end_date") or None,
                basic_salary=request.POST.get("basic_salary"),
                status=request.POST.get("status", "Draft"),
                created_by=request.user,
            )
            contract.full_clean()
            contract.save()
            notify(
                contract.employee, "contract", "Bạn có hợp đồng lao động mới",
                f"Hợp đồng {contract.contract_no}, hiệu lực từ {contract.start_date}.",
                link=f"/portal/hr/employees/{contract.employee_id}/#hopdong", created_by=request.user,
            )
            messages.success(request, "Đã tạo hợp đồng.")
            return redirect("portal:hr_contract_list")
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo hợp đồng.")

    return render(request, "portal/hr_contract_form.html", {
        "employees": Employee.objects.filter(employment_status__in=["Active", "Probation"]),
        "contract_types": ContractType.objects.all(),
    })


@login_required
@require_POST
def hr_contract_appendix_create(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền quản lý hợp đồng.")
        return redirect("portal:hr_contract_list")
    contract = get_object_or_404(Contract, pk=pk)
    try:
        appendix = ContractAppendix.objects.create(
            contract=contract, appendix_no=request.POST.get("appendix_no"),
            effective_date=request.POST.get("effective_date"),
            description=request.POST.get("description", ""), created_by=request.user,
        )
        notify(
            contract.employee, "contract", "Hợp đồng của bạn có phụ lục mới",
            f"Phụ lục {appendix.appendix_no} - {appendix.description or contract.contract_no}.",
            link=f"/portal/hr/employees/{contract.employee_id}/#hopdong", created_by=request.user,
        )
        messages.success(request, "Đã thêm phụ lục hợp đồng.")
    except Exception as e:
        _friendly_error(request, e, "Không thể thêm phụ lục.")
    return redirect("portal:hr_employee_detail", pk=contract.employee_id)

#=====================================================================
#Portal — Tuyển dụng đầy đủ (Ứng viên → Phỏng vấn → Offer → Onboarding)

@login_required
def hr_candidate_list(request):
    candidates = Candidate.objects.all().order_by("-created_date")
    return render(request, "portal/hr_candidate_list.html", {"candidates": candidates, "can_manage_hr": _can_manage_hr(request)})


@login_required
def hr_candidate_create(request):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền quản lý ứng viên.")
        return redirect("portal:hr_candidate_list")
    if request.method == "POST":
        try:
            candidate = Candidate(
                candidate_code=request.POST.get("candidate_code"),
                full_name=request.POST.get("full_name"),
                phone_number=request.POST.get("phone_number"),
                email=request.POST.get("email"),
                current_address=request.POST.get("current_address", ""),
                source=request.POST.get("source", ""),
                created_by=request.user,
            )
            candidate.full_clean()
            candidate.save()
            messages.success(request, "Đã thêm ứng viên.")
            return redirect("portal:hr_candidate_detail", pk=candidate.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể thêm ứng viên.")
    return render(request, "portal/hr_candidate_form.html")


@login_required
def hr_candidate_detail(request, pk):
    candidate = get_object_or_404(Candidate, pk=pk)
    return render(request, "portal/hr_candidate_detail.html", {
        "candidate": candidate, "can_manage_hr": _can_manage_hr(request),
        "interviews": candidate.interviews.select_related("interviewer"),
        "offers": candidate.offers.all(),
        "interviewers": Employee.objects.filter(employment_status="Active"),
    })


@login_required
@require_POST
def hr_interview_create(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_candidate_detail", pk=pk)
    candidate = get_object_or_404(Candidate, pk=pk)
    try:
        interview = Interview.objects.create(
            candidate=candidate, interview_date=request.POST.get("interview_date"),
            interview_type=request.POST.get("interview_type"),
            interviewer_id=request.POST.get("interviewer"),
            created_by=request.user,
        )
        candidate.status = "Interviewing"
        candidate.save(update_fields=["status"])
        if interview.interviewer_id:
            notify(
                interview.interviewer, "recruitment", "Bạn được phân công phỏng vấn",
                f"Ứng viên {candidate.full_name} - lúc {interview.interview_date}.",
                link=f"/portal/hr/candidates/{candidate.pk}/", created_by=request.user,
            )
        messages.success(request, "Đã lên lịch phỏng vấn.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect("portal:hr_candidate_detail", pk=pk)


@login_required
@require_POST
def hr_interview_result(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_candidate_list")
    interview = get_object_or_404(Interview, pk=pk)
    interview.result = request.POST.get("result")
    interview.comments = request.POST.get("comments", "")
    interview.modified_by = request.user
    interview.save()
    messages.success(request, "Đã cập nhật kết quả phỏng vấn.")
    return redirect("portal:hr_candidate_detail", pk=interview.candidate_id)


@login_required
@require_POST
def hr_offer_create(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_candidate_detail", pk=pk)
    candidate = get_object_or_404(Candidate, pk=pk)
    try:
        Offer.objects.create(
            candidate=candidate, offer_date=request.POST.get("offer_date"),
            expected_join_date=request.POST.get("expected_join_date"),
            salary_offer=request.POST.get("salary_offer"),
            created_by=request.user,
        )
        candidate.status = "Offered"
        candidate.save(update_fields=["status"])
        messages.success(request, "Đã tạo thư mời làm việc.")
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo offer.")
    return redirect("portal:hr_candidate_detail", pk=pk)


@login_required
@require_POST
def hr_offer_update_status(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_candidate_list")
    offer = get_object_or_404(Offer, pk=pk)
    offer.status = request.POST.get("status")
    offer.modified_by = request.user
    offer.save()
    messages.success(request, "Đã cập nhật trạng thái Offer.")
    return redirect("portal:hr_candidate_detail", pk=offer.candidate_id)


@login_required
@require_POST
def hr_candidate_onboard(request, pk):
    """Tạo hồ sơ nhân viên từ ứng viên đã có Offer 'Đã chấp nhận' — điền sẵn thông tin."""
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_candidate_detail", pk=pk)

    candidate = get_object_or_404(Candidate, pk=pk)
    accepted_offer = candidate.offers.filter(status="Accepted").order_by("-offer_date").first()
    if not accepted_offer:
        messages.error(request, "Ứng viên chưa có Offer ở trạng thái 'Đã chấp nhận'.")
        return redirect("portal:hr_candidate_detail", pk=pk)

    candidate.status = "Hired"
    candidate.save(update_fields=["status"])

    name_parts = candidate.full_name.strip().split(" ")
    last_name = name_parts[-1] if name_parts else ""
    first_name = " ".join(name_parts[:-1]) if len(name_parts) > 1 else candidate.full_name

    params = {
        "first_name": first_name, "last_name": last_name, "full_name": candidate.full_name,
        "personal_email": candidate.email, "mobile_phone": candidate.phone_number,
        "address": candidate.current_address or "", "join_date": accepted_offer.expected_join_date.isoformat(),
    }
    import urllib.parse
    query = urllib.parse.urlencode(params)
    messages.success(request, "Đã đánh dấu Trúng tuyển. Vui lòng hoàn tất hồ sơ nhân viên bên dưới.")
    return redirect(f"/portal/hr/employees/create/?{query}")

#========================
#Tự sinh mã nhân viên
def _generate_next_employee_code():
    max_num = 0
    for code in Employee.objects.values_list("employee_code", flat=True):
        m = re.match(r"^VH(\d+)$", code or "")
        if m:
            max_num = max(max_num, int(m.group(1)))
    return f"VH{max_num + 1:04d}"

#====================== 
#Thêm cột "Thao tác" ở danh sách Hồ sơ Nhân viên

@login_required
@require_POST
def hr_salary_create(request, pk):
    if not (_can_manage_hr(request) or _can_manage_finance(request)):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_employee_detail", pk=pk)
    employee = get_object_or_404(Employee, pk=pk)
    try:
        effective_date = request.POST.get("effective_date")
        # transaction.atomic() + select_for_update(): tránh 2 người cùng bấm "nâng lương"
        # cho 1 nhân viên cùng lúc tạo ra 2 bản ghi lương "Active" song song (mục 3.C).
        with transaction.atomic():
            employee = Employee.objects.select_for_update().get(pk=employee.pk)
            EmployeeSalary.objects.filter(employee=employee, status="Active").update(
                status="Ended", end_date=_day_before(effective_date)
            )
            new_salary = EmployeeSalary(
                employee=employee, current_basic_salary=request.POST.get("current_basic_salary"),
                effective_date=effective_date, status="Active", created_by=request.user,
            )
            new_salary.full_clean()
            new_salary.save()
        notify(
            employee, "salary", "Mức lương của bạn vừa được cập nhật",
            f"Lương cơ bản mới {new_salary.current_basic_salary}, áp dụng từ {effective_date}.",
            link=f"/portal/hr/employees/{pk}/#luong", created_by=request.user,
        )
        messages.success(request, "Đã ghi nhận nâng lương.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect(f"/portal/hr/employees/{pk}/#luong")


@login_required
@require_POST
def hr_contract_renew(request, pk):
    if not _can_manage_hr(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:hr_employee_detail", pk=pk)
    employee = get_object_or_404(Employee, pk=pk)
    contract = employee.contracts.order_by("-start_date").first()
    if not contract:
        messages.error(request, "Nhân viên chưa có hợp đồng nào để gia hạn.")
        return redirect("portal:hr_employee_detail", pk=pk)
    try:
        new_end_date = request.POST.get("new_end_date")
        # transaction.atomic() + select_for_update(): tránh 2 người cùng gia hạn 1 hợp
        # đồng cùng lúc tạo 2 phụ lục trùng nhau (mục 3.C).
        with transaction.atomic():
            contract = Contract.objects.select_for_update().get(pk=contract.pk)
            appendix = ContractAppendix(
                contract=contract, appendix_no=request.POST.get("appendix_no") or f"GH-{timezone.now():%Y%m%d%H%M}",
                effective_date=timezone.now().date(),
                description=f"Gia hạn hợp đồng đến {new_end_date}", created_by=request.user,
            )
            appendix.full_clean()
            appendix.save()
            contract.end_date = new_end_date
            contract.modified_by = request.user
            contract.full_clean(exclude=["created_by", "modified_by"])
            contract.save(update_fields=["end_date", "modified_by"])
        notify(
            employee, "contract", "Hợp đồng của bạn đã được gia hạn",
            f"Hợp đồng {contract.contract_no} gia hạn đến {new_end_date}.",
            link=f"/portal/hr/employees/{pk}/#hopdong", created_by=request.user,
        )
        messages.success(request, "Đã gia hạn hợp đồng.")
    except Exception as e:
        _friendly_error(request, e, "Không thể gia hạn.")
    return redirect(f"/portal/hr/employees/{pk}/#hopdong")

#====================================
#Chạy lương hàng tháng (Payroll Run)

def _can_manage_finance(request):
    return has_permission(request.user, "manage_finance")


@login_required
def finance_period_list(request):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền truy cập.")
        return redirect("portal:dashboard")
    periods = PayrollPeriod.objects.order_by("-payroll_year", "-payroll_month")
    return render(request, "portal/finance_period_list.html", {"periods": periods})


@login_required
@require_POST
def finance_period_create(request):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")
    month = parse_int(request.POST.get("payroll_month"), None, min_value=1, max_value=12)
    year = parse_int(request.POST.get("payroll_year"), None, min_value=2000, max_value=2100)
    if month is None or year is None:
        messages.error(request, "Tháng/năm không hợp lệ.")
        return redirect("portal:finance_period_list")
    try:
        first_day = f"{year}-{month:02d}-01"
        last_day = calendar_module.monthrange(year, month)[1]
        period = PayrollPeriod.objects.create(
            payroll_month=month, payroll_year=year,
            start_date=first_day, end_date=f"{year}-{month:02d}-{last_day:02d}",
            created_by=request.user,
        )
        messages.success(request, f"Đã tạo kỳ lương {month}/{year}.")
        return redirect("portal:finance_payroll_run", pk=period.pk)
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo kỳ lương.")
        return redirect("portal:finance_period_list")


def _employees_payable_for_period(period):
    """Return active employees and inactive employees whose last day falls in this period."""
    return Employee.objects.filter(
        models.Q(employment_status="Active")
        | models.Q(leave_date__gte=period.start_date, leave_date__lte=period.end_date)
    )


@login_required
def finance_payroll_run(request, pk):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")

    period = get_object_or_404(PayrollPeriod, pk=pk)
    existing_emp_ids = set(Payroll.objects.filter(payroll_period=period).values_list("employee_id", flat=True))

    payrolls = Payroll.objects.filter(payroll_period=period).select_related("employee")
    missing_employees = _employees_payable_for_period(period).exclude(id__in=existing_emp_ids)
    unissued_payslips = payrolls.filter(payslip__isnull=True).count() if period.status == "Locked" else 0

    return render(request, "portal/finance_payroll_run.html", {
        "period": period, "payrolls": payrolls, "missing_employees": missing_employees,
        "unissued_payslips": unissued_payslips,
    })


@login_required
@require_POST
def finance_payroll_generate(request, pk):
    """Tạo dòng Bảng lương cho các nhân viên chưa có trong kỳ này, lấy lương cơ bản từ EmployeeSalary."""
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")

    period = get_object_or_404(PayrollPeriod, pk=pk)
    if period.status == "Locked":
        messages.error(request, "Kỳ lương đã khóa — không thể tạo thêm dòng bảng lương.")
        return redirect("portal:finance_payroll_run", pk=pk)

    with transaction.atomic():
        # Khóa bản ghi kỳ lương để 2 người bấm cùng lúc không tạo trùng
        period = PayrollPeriod.objects.select_for_update().get(pk=pk)
        existing_emp_ids = set(
            Payroll.objects.filter(payroll_period=period).values_list("employee_id", flat=True)
        )
        employees = list(
            _employees_payable_for_period(period).exclude(id__in=existing_emp_ids)
        )
        # Lấy lương cơ bản của tất cả nhân viên trong 1 lần truy vấn thay vì lặp từng người
        salary_map = {}
        for sal in (
            EmployeeSalary.objects.filter(employee__in=employees, effective_date__lte=period.start_date)
            .filter(models.Q(end_date__isnull=True) | models.Q(end_date__gte=period.start_date))
            .order_by("effective_date")
        ):
            salary_map[sal.employee_id] = sal.current_basic_salary

        new_rows = [
            Payroll(
                payroll_period=period, employee=employee,
                basic_salary=salary_map.get(employee.id, 0),
                insurance_salary=salary_map.get(employee.id, 0),
                created_by=request.user,
            )
            for employee in employees
        ]
        Payroll.objects.bulk_create(new_rows)
        count = len(new_rows)

    messages.success(request, f"Đã tạo {count} dòng bảng lương mới cho kỳ {period}.")
    return redirect("portal:finance_payroll_run", pk=pk)


@login_required
@require_POST
def finance_payroll_calculate_all(request, pk):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")

    period = get_object_or_404(PayrollPeriod, pk=pk)
    if period.status == "Locked":
        messages.error(request, "Kỳ lương đã khóa — không thể tính lại lương.")
        return redirect("portal:finance_payroll_run", pk=pk)

    # I13: mỗi nhân viên dùng 1 savepoint atomic() riêng (nested), KHÔNG bọc chung
    # 1 atomic() cho cả vòng lặp. Nếu bọc chung, lỗi của 1 nhân viên (kể cả
    # ValidationError thuần Python) sẽ khiến Django đánh dấu toàn bộ transaction
    # ngoài là "needs_rollback" — mọi câu lệnh DB của các nhân viên tính SAU đó
    # trong cùng vòng lặp sẽ ăn theo lỗi TransactionManagementError, dù bản thân
    # họ tính hoàn toàn đúng. Dùng savepoint riêng từng người để cô lập lỗi.
    count, errors = 0, []
    for payroll in Payroll.objects.filter(payroll_period=period).select_related("employee"):
        try:
            with transaction.atomic():
                payroll.calculate()
                payroll.modified_by = request.user
                payroll.save()
            count += 1
        except ValidationError as e:
            detail = "; ".join(e.messages) if hasattr(e, "messages") else str(e)
            errors.append(f"{payroll.employee}: {detail}")
        except Exception as e:
            logger.exception("Lỗi tính lương nhân viên %s (kỳ %s)", payroll.employee, period)
            errors.append(f"{payroll.employee}: Có lỗi hệ thống khi tính lương, đã ghi log.")

    messages.success(request, f"Đã tính lương cho {count} nhân viên.")
    for err in errors:
        messages.error(request, err)
    return redirect("portal:finance_payroll_run", pk=pk)


@login_required
@require_POST
def finance_payroll_lock(request, pk):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")
    with transaction.atomic():
        period = get_object_or_404(PayrollPeriod.objects.select_for_update(), pk=pk)
        period.status = "Locked"
        period.modified_by = request.user
        period.save()
        payrolls = list(Payroll.objects.filter(payroll_period=period).select_related("employee"))
        existing_ids = set(Payslip.objects.filter(payroll__payroll_period=period).values_list("payroll_id", flat=True))
        to_issue = [payroll for payroll in payrolls if payroll.pk not in existing_ids]
        Payslip.objects.bulk_create([
            Payslip(payroll=payroll, created_by=request.user) for payroll in to_issue
        ])
        period_managers = Employee.objects.filter(
            employment_status="Active",
            user_account__role__role_code__in=("DIRECTOR", "FINANCE_OFFICER"),
        ).select_related("user_account", "user_account__role").distinct()
        for recipient in period_managers:
            is_finance = recipient.user_account.role.role_code == "FINANCE_OFFICER"
            notify_with_email_draft(
                recipient,
                "payroll-period-locked",
                "Kỳ lương đã khóa",
                f"Kỳ lương {period} đã được khóa trong IHRMS.",
                link="/portal/finance/periods/" if is_finance else "/portal/reports/",
                email_subject="IHRMS: Kỳ lương đã khóa",
                email_body=(
                    f"Kỳ lương {period} đã được khóa. "
                    "Vui lòng đăng nhập IHRMS để xem thông tin theo quyền được cấp."
                ),
                created_by=request.user,
                idempotency_key=f"payroll-period-locked:{period.pk}:{recipient.pk}",
            )
    for payroll in to_issue:
        notify(
            payroll.employee, "payroll", "Phiếu lương của bạn đã sẵn sàng",
            f"Bảng lương kỳ {period} đã được chốt; tải phiếu lương PDF trong mục Phiếu lương.",
            link="/portal/payslip/", created_by=request.user,
        )
    messages.success(request, f"Đã khóa kỳ lương {period}. Không thể sửa bảng lương của kỳ này nữa.")
    return redirect("portal:finance_period_list")


@login_required
@require_POST
def finance_payslip_issue(request, pk):
    """Backfill payslip records for periods locked before PDF release was enabled."""
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")
    with transaction.atomic():
        period = get_object_or_404(PayrollPeriod.objects.select_for_update(), pk=pk)
        if period.status != "Locked":
            messages.error(request, "Chỉ có thể phát hành phiếu lương sau khi kỳ đã khóa.")
            return redirect("portal:finance_payroll_run", pk=period.pk)
        existing_ids = set(Payslip.objects.filter(payroll__payroll_period=period).values_list("payroll_id", flat=True))
        to_issue = list(
            Payroll.objects.filter(payroll_period=period).exclude(pk__in=existing_ids).select_related("employee")
        )
        Payslip.objects.bulk_create([
            Payslip(payroll=payroll, created_by=request.user) for payroll in to_issue
        ])
    for payroll in to_issue:
        notify(
            payroll.employee, "payroll", "Phiếu lương của bạn đã sẵn sàng",
            f"Phiếu lương PDF kỳ {period} đã được phát hành trong mục Phiếu lương.",
            link="/portal/payslip/", created_by=request.user,
        )
    messages.success(request, f"Đã phát hành {len(to_issue)} phiếu lương PDF cho kỳ {period}.")
    return redirect("portal:finance_payroll_run", pk=period.pk)


@login_required
@require_POST
def finance_payroll_update_line(request, pk):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:finance_period_list")
    payroll = get_object_or_404(Payroll, pk=pk)
    if payroll.payroll_period.status == "Locked":
        messages.error(request, "Kỳ lương đã khóa, không thể sửa.")
        return redirect("portal:finance_payroll_run", pk=payroll.payroll_period_id)
    try:
        payroll.allowance = request.POST.get("allowance") or 0
        payroll.overtime_amount = request.POST.get("overtime_amount") or 0
        payroll.overtime_tax_exempt_amount = request.POST.get("overtime_tax_exempt_amount") or 0
        payroll.insurance_salary = request.POST.get("insurance_salary") or 0
        payroll.deduction_amount = request.POST.get("deduction_amount") or 0
        payroll.modified_by = request.user
        payroll.full_clean(exclude=["created_by", "modified_by"])
        payroll.save()
        messages.success(request, "Đã cập nhật, cần bấm \"Tính lương lại\" để cập nhật số liệu.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect("portal:finance_payroll_run", pk=payroll.payroll_period_id)

#Lương cơ bản nhân viên (Kế toán xem/sửa toàn công ty)
@login_required
def finance_salary_list(request):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    salaries = EmployeeSalary.objects.filter(status="Active").select_related("employee").order_by("employee__full_name")
    return render(request, "portal/finance_salary_list.html", {"salaries": salaries})

#Cấu hình Bảo hiểm & Thuế TNCN
@login_required
def finance_tax_config(request):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")

    if request.method == "POST":
        form_type = request.POST.get("form_type")
        try:
            if form_type == "insurance":
                effective_from = request.POST.get("effective_from")
                with transaction.atomic():
                    # Đóng effective_to của các thế hệ cũ còn bỏ ngỏ (C1: nếu không đóng,
                    # calculate() có thể vẫn chỉ lấy .first() theo effective_from mới nhất nên ít rủi ro
                    # hơn bracket, nhưng đóng lại cho rõ ràng, tránh nhầm lẫn khi tra cứu).
                    InsuranceRateConfig.objects.filter(effective_to__isnull=True).exclude(
                        effective_from__gte=effective_from
                    ).update(effective_to=_day_before(effective_from))
                    InsuranceRateConfig.objects.create(
                        effective_from=effective_from,
                        employee_social_insurance_rate=request.POST.get("employee_social_insurance_rate"),
                        employee_health_insurance_rate=request.POST.get("employee_health_insurance_rate"),
                        employee_unemployment_insurance_rate=request.POST.get("employee_unemployment_insurance_rate"),
                        employer_social_insurance_rate=request.POST.get("employer_social_insurance_rate"),
                        employer_health_insurance_rate=request.POST.get("employer_health_insurance_rate"),
                        employer_unemployment_insurance_rate=request.POST.get("employer_unemployment_insurance_rate"),
                        legal_reference=request.POST.get("legal_reference", ""),
                        created_by=request.user,
                    )
                messages.success(request, "Đã thêm cấu hình tỷ lệ Bảo hiểm.")
            elif form_type == "deduction":
                effective_from = request.POST.get("effective_from")
                with transaction.atomic():
                    TaxDeductionConfig.objects.filter(effective_to__isnull=True).exclude(
                        effective_from__gte=effective_from
                    ).update(effective_to=_day_before(effective_from))
                    TaxDeductionConfig.objects.create(
                        effective_from=effective_from,
                        self_deduction_monthly=request.POST.get("self_deduction_monthly"),
                        dependent_deduction_monthly=request.POST.get("dependent_deduction_monthly"),
                        created_by=request.user,
                    )
                messages.success(request, "Đã thêm cấu hình giảm trừ thuế.")
            elif form_type == "bracket":
                # C1 (sửa): tự đóng effective_to của TOÀN BỘ thế hệ bậc thuế cũ (mọi bracket_no
                # cùng effective_from cũ, chưa có effective_to) trước khi thêm bậc của thế hệ mới —
                # để calculate() (và bất kỳ báo cáo nào lọc theo effective_to) không bao giờ thấy
                # 2 thế hệ cùng "đang hiệu lực" một lúc, kể cả khi người dùng không khai đủ tất cả
                # các bậc trong cùng 1 lần submit.
                effective_from = request.POST.get("effective_from")
                with transaction.atomic():
                    PersonalIncomeTaxBracket.objects.filter(effective_to__isnull=True).exclude(
                        effective_from__gte=effective_from
                    ).update(effective_to=_day_before(effective_from))
                    PersonalIncomeTaxBracket.objects.create(
                        effective_from=effective_from,
                        bracket_no=request.POST.get("bracket_no"),
                        income_from_monthly=request.POST.get("income_from_monthly"),
                        income_to_monthly=request.POST.get("income_to_monthly") or None,
                        tax_rate=request.POST.get("tax_rate"),
                        created_by=request.user,
                    )
                messages.success(request, "Đã thêm bậc thuế TNCN.")
        except Exception as e:
            _friendly_error(request, e, "Không thể lưu.")
        return redirect("portal:finance_tax_config")

    return render(request, "portal/finance_tax_config.html", {
        "insurance_configs": InsuranceRateConfig.objects.order_by("-effective_from"),
        "deduction_configs": TaxDeductionConfig.objects.order_by("-effective_from"),
        "tax_brackets": PersonalIncomeTaxBracket.objects.order_by("-effective_from", "bracket_no"),
    })

#Người phụ thuộc
@login_required
def finance_dependent_list(request):
    if not _can_manage_finance(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    if request.method == "POST":
        try:
            dependent = EmployeeDependent(
                employee_id=request.POST.get("employee"),
                full_name=request.POST.get("full_name"),
                relationship=request.POST.get("relationship"),
                registered_from=request.POST.get("registered_from"),
                created_by=request.user,
            )
            dependent.full_clean()
            dependent.save()
            messages.success(request, "Đã thêm người phụ thuộc.")
        except Exception as e:
            _friendly_error(request, e, "Không thể lưu.")
        return redirect("portal:finance_dependent_list")
    return render(request, "portal/finance_dependent_list.html", {
        "dependents": EmployeeDependent.objects.filter(is_active=True).select_related("employee"),
        "employees": Employee.objects.filter(employment_status="Active"),
    })

#====================================
#Portal — Kiểm kê tài sản


def _can_manage_asset(request):
    return has_permission(request.user, "manage_asset")


def _can_view_asset(request):
    return has_permission(request.user, "view_asset")


@login_required
def asset_admin_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    can_manage_asset = _can_manage_asset(request)
    assets = Asset.objects.select_related("asset_type").order_by("asset_type__asset_group", "asset_code")
    return render(request, "portal/asset_admin_list.html", {
        "assets": assets,
        "asset_types": AssetType.objects.all() if can_manage_asset else AssetType.objects.none(),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def asset_admin_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:asset_admin_list")
    try:
        asset = Asset(
            asset_code=request.POST.get("asset_code"), asset_name=request.POST.get("asset_name"),
            asset_type_id=request.POST.get("asset_type"), location=request.POST.get("location", ""),
            brand=request.POST.get("brand", ""), serial_number=request.POST.get("serial_number", ""),
            funding_source=request.POST.get("funding_source", ""),
            purchase_date=request.POST.get("purchase_date") or None,
            quantity=request.POST.get("quantity") or 1, unit=request.POST.get("unit", "chiếc"),
            purchase_value=request.POST.get("purchase_value") or None,
            note=request.POST.get("note", ""), created_by=request.user,
        )
        asset.full_clean()
        asset.save()
        messages.success(request, "Đã thêm tài sản.")
    except Exception as e:
        _friendly_error(request, e, "Không thể thêm.")
    return redirect("portal:asset_admin_list")


@login_required
def inventory_session_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    sessions = InventorySession.objects.order_by("-inventory_datetime")
    return render(request, "portal/inventory_session_list.html", {
        "sessions": sessions, "can_manage_asset": _can_manage_asset(request),
    })


@login_required
@require_POST
def inventory_session_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:inventory_session_list")
    try:
        session = InventorySession.objects.create(
            session_name=request.POST.get("session_name"),
            office_name=request.POST.get("office_name"),
            inventory_datetime=request.POST.get("inventory_datetime"),
            created_by=request.user,
        )
        # Tự nạp toàn bộ tài sản đang hoạt động (chưa thanh lý) vào đợt kiểm kê
        for asset in Asset.objects.exclude(status="Disposed"):
            InventoryLine.objects.create(session=session, asset=asset, actual_condition=asset.condition)
        messages.success(request, "Đã tạo đợt kiểm kê và nạp toàn bộ tài sản.")
        return redirect("portal:inventory_session_detail", pk=session.pk)
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo.")
        return redirect("portal:inventory_session_list")


@login_required
def inventory_session_detail(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)

    if request.method == "POST" and not _can_manage_asset(request):
        messages.error(request, "Anh/chị chỉ có quyền xem thông tin tài sản.")
        return redirect("portal:inventory_session_detail", pk=pk)

    if request.method == "POST" and request.POST.get("action") == "add_member":
        try:
            InventoryMember.objects.create(
                session=session, employee_id=request.POST.get("employee"),
                role=request.POST.get("role"), created_by=request.user,
            )
            messages.success(request, "Đã thêm thành viên Ban kiểm kê.")
        except Exception as e:
            _friendly_error(request, e, "Không thể thêm.")
        return redirect("portal:inventory_session_detail", pk=pk)

    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")
    can_manage_asset = _can_manage_asset(request)
    return render(request, "portal/inventory_session_detail.html", {
        "session": session, "lines": lines,
        "uncounted_lines": lines.filter(is_counted=False).count(),
        "line_count": lines.count(),
        "members": session.members.select_related("employee"),
        "employees": Employee.objects.filter(employment_status="Active") if can_manage_asset else Employee.objects.none(),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def inventory_line_update(request, pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    line = get_object_or_404(InventoryLine, pk=pk)
    if line.session.status == "Completed":
        messages.error(request, "Đợt kiểm kê đã hoàn tất, không thể thay đổi dòng kiểm kê.")
        return redirect("portal:inventory_session_detail", pk=line.session_id)
    line.is_counted = True
    line.actual_condition = request.POST.get("actual_condition", line.actual_condition)
    line.note = request.POST.get("note", "")
    line.modified_by = request.user
    line.save()
    messages.success(request, f"Đã kiểm đếm: {line.asset}.")
    return redirect("portal:inventory_session_detail", pk=line.session_id)


@login_required
@require_POST
def inventory_session_complete(request, pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    line_count = session.lines.count()
    uncounted_count = session.lines.filter(is_counted=False).count()
    if line_count == 0:
        messages.error(request, "Chưa thể hoàn tất đợt kiểm kê: chưa có tài sản nào trong danh sách.")
        return redirect("portal:inventory_session_detail", pk=pk)
    if uncounted_count:
        messages.error(request, f"Chưa thể hoàn tất đợt kiểm kê: còn {uncounted_count} tài sản chưa kiểm đếm.")
        return redirect("portal:inventory_session_detail", pk=pk)
    session.status = "Completed"
    session.recommendation = request.POST.get("recommendation", "Không.")
    session.modified_by = request.user
    session.save()
    messages.success(request, "Đã hoàn tất đợt kiểm kê.")
    return redirect("portal:inventory_session_print", pk=pk)


@login_required
def inventory_session_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")

    machinery = [l for l in lines if l.asset.asset_type.asset_group == "Machinery"]
    other = [l for l in lines if l.asset.asset_type.asset_group == "Other"]

    return render(request, "portal/inventory_session_print.html", {
        "session": session, "machinery": machinery, "other": other,
        "chairman": session.members.filter(role="Chairman").first(),
        "recorder": session.members.filter(role="Recorder").first(),
        "member_list": session.members.filter(role="Member"),
        "all_members": session.members.select_related("employee"),
    })


@login_required
def inventory_session_export_excel(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")

    import openpyxl
    from openpyxl.styles import Font, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Kiểm kê"
    bold = Font(bold=True)
    center = Alignment(horizontal="center")

    ws["A1"] = "BIÊN BẢN KIỂM KÊ TRANG THIẾT BỊ"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = session.office_name
    ws["A3"] = f"Thời điểm kiểm kê: {session.inventory_datetime:%H giờ %M, ngày %d tháng %m năm %Y}"

    row = 5
    for m in session.members.select_related("employee"):
        ws.cell(row=row, column=1, value=f"{m.employee.full_name} — {m.get_role_display()}")
        row += 1

    row += 1
    headers = ["STT", "Tên thiết bị", "Mã số", "Loại", "Nơi SD", "Nhãn hiệu", "Số Seri", "Nguồn mua sắm", "Ngày mua", "SL", "ĐVT", "Giá trị", "Tình trạng", "Ghi chú"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = bold
    header_row = row
    row += 1

    def write_group(group_name, lines_group, stt_group):
        nonlocal row
        ws.cell(row=row, column=1, value=stt_group).font = bold
        ws.cell(row=row, column=2, value=group_name).font = bold
        row += 1
        for i, l in enumerate(lines_group, start=1):
            a = l.asset
            values = [i, a.asset_name, a.asset_code, a.asset_type.asset_type_name, a.location, a.brand,
                      a.serial_number, a.funding_source, a.purchase_date.isoformat() if a.purchase_date else "",
                      a.quantity, a.unit, float(a.purchase_value) if a.purchase_value else "", l.get_actual_condition_display(), l.note or ""]
            for c, v in enumerate(values, start=1):
                ws.cell(row=row, column=c, value=v)
            row += 1

    machinery = [l for l in lines if l.asset.asset_type.asset_group == "Machinery"]
    other = [l for l in lines if l.asset.asset_type.asset_group == "Other"]
    write_group("I. Máy móc thiết bị", machinery, "I")
    write_group("II. Thiết bị khác", other, "II")

    row += 1
    ws.cell(row=row, column=2, value=f"* Kết quả kiểm kê: có {len(machinery)} máy móc thiết bị và {len(other)} thiết bị khác.")
    row += 1
    ws.cell(row=row, column=2, value=f"* Kiến nghị: {session.recommendation}")

    for c in range(1, 15):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 18

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="BBKK_{session.session_name}.xlsx"'
    wb.save(response)
    return response

#Portal — Đề nghị & Biên bản Thanh lý
@login_required
def disposal_request_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    requests_qs = AssetDisposalRequest.objects.order_by("-request_date")
    return render(request, "portal/disposal_request_list.html", {
        "requests": requests_qs, "can_manage_asset": _can_manage_asset(request),
    })


@login_required
def disposal_request_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:disposal_request_list")
    employee = _get_employee(request)
    if request.method == "POST":
        try:
            disposal_request = AssetDisposalRequest.objects.create(
                request_no=request.POST.get("request_no"),
                request_date=request.POST.get("request_date"),
                location_name=request.POST.get("location_name"),
                disposal_type=request.POST.get("disposal_type", "Liquidate"),
                requested_by=employee, created_by=request.user,
            )
            messages.success(request, "Đã tạo đề nghị thanh lý.")
            return redirect("portal:disposal_request_detail", pk=disposal_request.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo.")
    return render(request, "portal/disposal_request_form.html", {})


def _can_approve_disposal(request):
    """C3 (đã sửa): quyền duyệt thanh lý CHỈ dành cho Giám đốc/superuser — tách hẳn khỏi
    `_can_manage_asset` (HR_OFFICER trở lên, chỉ nên được thêm thành viên/tài sản vào biên bản,
    không được tự duyệt). Dùng chung logic với biến `can_approve` đã hiển thị ở template."""
    return _director_employees().filter(user_account=request.user).exists() or request.user.is_superuser


@login_required
def disposal_request_detail(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)

    if request.method == "POST":
        action = request.POST.get("action")
        if action != "approve" and not _can_manage_asset(request):
            messages.error(request, "Anh/chị chỉ có quyền xem thông tin tài sản.")
            return redirect("portal:disposal_request_detail", pk=pk)
        if action == "add_member":
            try:
                member = AssetDisposalRequestMember(
                    request=disposal_request, employee_id=request.POST.get("employee"),
                    participation=request.POST.get("participation", "Direct"),
                    role=request.POST.get("role", "Member"), created_by=request.user,
                )
                member.full_clean()
                member.save()
                messages.success(request, "Đã thêm thành viên.")
            except Exception as e:
                _friendly_error(request, e, "Không thể thêm thành viên.")
        elif action == "add_item":
            try:
                item = AssetDisposalItem(
                    request=disposal_request, asset_id=request.POST.get("asset"),
                    damage_description=request.POST.get("damage_description"), created_by=request.user,
                )
                item.full_clean()
                item.save()
                messages.success(request, "Đã thêm tài sản vào đề nghị thanh lý.")
            except Exception as e:
                _friendly_error(request, e, "Không thể thêm tài sản.")
        elif action == "approve":
            # C3 (đã sửa): trước đây chỉ check `_can_manage_asset` (HR_OFFICER trở lên) rồi set
            # thẳng status="Approved" — biến `can_approve` (Director-only) chỉ dùng ẩn/hiện nút ở
            # template, ai POST tay với `action=approve` đều tự duyệt được. Giờ kiểm tra lại
            # `_can_approve_disposal` ngay server-side, và chỉ duyệt đơn đang "Pending".
            if not _can_approve_disposal(request):
                messages.error(request, "Chỉ Giám đốc mới có quyền duyệt đề nghị thanh lý.")
                return redirect("portal:disposal_request_detail", pk=pk)
            # transaction.atomic() + select_for_update(): tránh 2 người cùng duyệt 1 đề
            # nghị thanh lý cùng lúc (mục 3.C).
            with transaction.atomic():
                disposal_request = get_object_or_404(
                    AssetDisposalRequest.objects.select_for_update(), pk=pk, status="Pending",
                )
                disposal_request.status = "Approved"
                disposal_request.modified_by = request.user
                disposal_request.save()
            submitter = disposal_request.created_by.employee if hasattr(disposal_request.created_by, "employee") else None
            notify(
                submitter, "asset", "Đề nghị thanh lý tài sản đã được duyệt",
                f"Đề nghị {disposal_request.request_no} đã được duyệt.",
                link=f"/portal/disposal/{disposal_request.pk}/", created_by=request.user,
            )
            messages.success(request, "Đã duyệt đề nghị thanh lý.")
        return redirect("portal:disposal_request_detail", pk=pk)

    can_manage_asset = _can_manage_asset(request)
    return render(request, "portal/disposal_request_detail.html", {
        "req": disposal_request, "members": disposal_request.members.select_related("employee"),
        "items": disposal_request.items.select_related("asset"),
        "employees": Employee.objects.filter(employment_status="Active") if can_manage_asset else Employee.objects.none(),
        "assets": Asset.objects.exclude(status="Disposed") if can_manage_asset else Asset.objects.none(),
        "directors": _director_employees() if can_manage_asset else Employee.objects.none(),
        "can_approve": _can_approve_disposal(request),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def disposal_minutes_create(request, item_pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    item = get_object_or_404(AssetDisposalItem, pk=item_pk)
    try:
        AssetDisposalMinutes.objects.create(
            item=item,
            decision_date=request.POST.get("decision_date"),
            origin_country=request.POST.get("origin_country", ""),
            manufacture_year=request.POST.get("manufacture_year") or None,
            put_into_use_year=request.POST.get("put_into_use_year") or None,
            original_value=request.POST.get("original_value") or 0,
            depreciated_value=request.POST.get("depreciated_value") or 0,
            conclusion=request.POST.get("conclusion"),
            disposal_cost=request.POST.get("disposal_cost") or 0,
            recovery_value=request.POST.get("recovery_value") or 0,
            director_signer_id=request.POST.get("director_signer") or None,
            accountant_signer_id=request.POST.get("accountant_signer") or None,
            created_by=request.user,
        )
        messages.success(request, "Đã lập biên bản thanh lý. Tài sản đã chuyển trạng thái 'Đã thanh lý'.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect("portal:disposal_request_detail", pk=item.request_id)


@login_required
def disposal_minutes_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    minutes = get_object_or_404(AssetDisposalMinutes, pk=pk)
    return render(request, "portal/disposal_minutes_print.html", {
        "minutes": minutes, "item": minutes.item, "req": minutes.item.request,
        "committee": minutes.item.request.members.select_related("employee"),
    })


@login_required
def disposal_request_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_request_print.html", {
        "req": disposal_request, "members": disposal_request.members.select_related("employee"),
        "items": disposal_request.items.select_related("asset"),
    })

@login_required
def disposal_decision_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_decision_print.html", {
        "req": disposal_request, "items": disposal_request.items.select_related("asset"),
    })


@login_required
def disposal_execution_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_execution_print.html", {
        "req": disposal_request, "items": disposal_request.items.select_related("asset"),
        "chairman": disposal_request.members.filter(role="Chairman").first(),
        "member_list": disposal_request.members.filter(role="Member"),
    })

#===============================================================
@login_required
def dispatch_create(request):
    if not has_permission(request.user, "manage_office"):
        messages.error(request, "Anh/chị không có quyền vào sổ công văn.")
        return redirect("portal:dispatch_list")

    employee = _get_employee(request)
    if not employee:
        messages.error(request, "Tài khoản chưa gắn hồ sơ nhân viên.")
        return redirect("portal:dispatch_list")

    if request.method == "POST":
        try:
            dispatch_type = request.POST.get("dispatch_type")
            issue_date = request.POST.get("issue_date")
            year = int(issue_date[:4])

            book_names = {
                "Incoming": "Sổ công văn đến", "Outgoing": "Sổ công văn đi", "InternalDirective": "Sổ văn bản nội bộ",
            }
            book, _ = DispatchBook.objects.get_or_create(
                book_type=dispatch_type, year=year,
                defaults={
                    "book_code": f"{dispatch_type}-{year}",
                    "book_name": f"{book_names.get(dispatch_type)} {year}",
                    "created_by": request.user,
                },
            )
            book.current_number += 1
            book.save(update_fields=["current_number"])

            dispatch = OfficialDispatch(
                dispatch_type=dispatch_type, dispatch_book=book,
                dispatch_no=request.POST.get("dispatch_no"),
                original_doc_no=request.POST.get("original_doc_no", ""),
                title=request.POST.get("title"),
                issuing_organization=request.POST.get("issuing_organization"),
                issue_date=issue_date,
                received_or_sent_date=request.POST.get("received_or_sent_date"),
                urgency_level=request.POST.get("urgency_level", "Normal"),
                security_level=request.POST.get("security_level", "Normal"),
                department_id=request.POST.get("department") or None,
                assigned_to_id=request.POST.get("assigned_to") or None,
                signer_id=request.POST.get("signer") or None,
                due_date=request.POST.get("due_date") or None,
                attachment=request.FILES.get("attachment"),
                created_by=request.user,
            )
            dispatch.full_clean()
            dispatch.save()
            DispatchTracking.objects.create(
                dispatch=dispatch, actor=employee, action_type="Vào sổ",
                comment="Tiếp nhận/Phát hành công văn.", created_by=request.user,
            )
            if dispatch.assigned_to_id:
                notify(
                    dispatch.assigned_to, "dispatch", "Bạn được giao xử lý công văn",
                    f"{dispatch.dispatch_no} - {dispatch.title}.",
                    link=f"/portal/dispatches/{dispatch.pk}/", created_by=request.user,
                )
            messages.success(request, f"Đã vào sổ công văn {dispatch.dispatch_no}.")
            return redirect("portal:dispatch_detail", pk=dispatch.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể lưu.")

    return render(request, "portal/dispatch_create.html", {
        "departments": Department.objects.filter(is_active=True),
        "employees": Employee.objects.filter(employment_status="Active"),
    })
