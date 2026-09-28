from django.db import transaction
from django.views.decorators.http import require_POST
from apps.security.permissions import has_permission
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.shortcuts import get_object_or_404, redirect, render
from apps.notification.services import notify
from django.http import HttpResponse
from apps.asset.models import (
    Asset,
    AssetDisposalItem,
    AssetDisposalMinutes,
    AssetDisposalRequest,
    AssetDisposalRequestMember,
    AssetType,
    InventoryLine,
    InventoryMember,
    InventorySession,
)
from apps.hrm.models import Employee
from apps.asset.models import AssetAllocation
from .common import _director_employees, _friendly_error, _get_employee
@login_required


def asset_list(request):
    if not has_permission(request.user, "view_asset"):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    employee = _get_employee(request)
    allocations = AssetAllocation.objects.filter(
        employee=employee, return_date__isnull=True
    ).select_related("asset", "asset__asset_type") if employee else []
    return render(request, "portal/asset_list.html", {
        "allocations": allocations,
        "can_manage_asset": _can_manage_asset(request),
    })


def _can_manage_asset(request):
    return has_permission(request.user, "manage_asset")


def _can_view_asset(request):
    return has_permission(request.user, "view_asset")


@login_required
def asset_admin_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    can_manage_asset = _can_manage_asset(request)
    assets = Asset.objects.select_related("asset_type").order_by("asset_type__asset_group", "asset_code")
    return render(request, "portal/asset_admin_list.html", {
        "assets": assets,
        "asset_types": AssetType.objects.all() if can_manage_asset else AssetType.objects.none(),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def asset_admin_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:asset_admin_list")
    try:
        asset = Asset(
            asset_code=request.POST.get("asset_code"), asset_name=request.POST.get("asset_name"),
            asset_type_id=request.POST.get("asset_type"), location=request.POST.get("location", ""),
            brand=request.POST.get("brand", ""), serial_number=request.POST.get("serial_number", ""),
            funding_source=request.POST.get("funding_source", ""),
            purchase_date=request.POST.get("purchase_date") or None,
            quantity=request.POST.get("quantity") or 1, unit=request.POST.get("unit", "chiếc"),
            purchase_value=request.POST.get("purchase_value") or None,
            note=request.POST.get("note", ""), created_by=request.user,
        )
        asset.full_clean()
        asset.save()
        messages.success(request, "Đã thêm tài sản.")
    except Exception as e:
        _friendly_error(request, e, "Không thể thêm.")
    return redirect("portal:asset_admin_list")


@login_required
def inventory_session_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    sessions = InventorySession.objects.order_by("-inventory_datetime")
    return render(request, "portal/inventory_session_list.html", {
        "sessions": sessions, "can_manage_asset": _can_manage_asset(request),
    })


@login_required
@require_POST
def inventory_session_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:inventory_session_list")
    try:
        session = InventorySession.objects.create(
            session_name=request.POST.get("session_name"),
            office_name=request.POST.get("office_name"),
            inventory_datetime=request.POST.get("inventory_datetime"),
            created_by=request.user,
        )
        # Tự nạp toàn bộ tài sản đang hoạt động (chưa thanh lý) vào đợt kiểm kê
        for asset in Asset.objects.exclude(status="Disposed"):
            InventoryLine.objects.create(session=session, asset=asset, actual_condition=asset.condition)
        messages.success(request, "Đã tạo đợt kiểm kê và nạp toàn bộ tài sản.")
        return redirect("portal:inventory_session_detail", pk=session.pk)
    except Exception as e:
        _friendly_error(request, e, "Không thể tạo.")
        return redirect("portal:inventory_session_list")


@login_required
def inventory_session_detail(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)

    if request.method == "POST" and not _can_manage_asset(request):
        messages.error(request, "Anh/chị chỉ có quyền xem thông tin tài sản.")
        return redirect("portal:inventory_session_detail", pk=pk)

    if request.method == "POST" and request.POST.get("action") == "add_member":
        try:
            InventoryMember.objects.create(
                session=session, employee_id=request.POST.get("employee"),
                role=request.POST.get("role"), created_by=request.user,
            )
            messages.success(request, "Đã thêm thành viên Ban kiểm kê.")
        except Exception as e:
            _friendly_error(request, e, "Không thể thêm.")
        return redirect("portal:inventory_session_detail", pk=pk)

    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")
    can_manage_asset = _can_manage_asset(request)
    return render(request, "portal/inventory_session_detail.html", {
        "session": session, "lines": lines,
        "uncounted_lines": lines.filter(is_counted=False).count(),
        "line_count": lines.count(),
        "members": session.members.select_related("employee"),
        "employees": Employee.objects.filter(employment_status="Active") if can_manage_asset else Employee.objects.none(),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def inventory_line_update(request, pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    line = get_object_or_404(InventoryLine, pk=pk)
    if line.session.status == "Completed":
        messages.error(request, "Đợt kiểm kê đã hoàn tất, không thể thay đổi dòng kiểm kê.")
        return redirect("portal:inventory_session_detail", pk=line.session_id)
    line.is_counted = True
    line.actual_condition = request.POST.get("actual_condition", line.actual_condition)
    line.note = request.POST.get("note", "")
    line.modified_by = request.user
    line.save()
    messages.success(request, f"Đã kiểm đếm: {line.asset}.")
    return redirect("portal:inventory_session_detail", pk=line.session_id)


@login_required
@require_POST
def inventory_session_complete(request, pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    line_count = session.lines.count()
    uncounted_count = session.lines.filter(is_counted=False).count()
    if line_count == 0:
        messages.error(request, "Chưa thể hoàn tất đợt kiểm kê: chưa có tài sản nào trong danh sách.")
        return redirect("portal:inventory_session_detail", pk=pk)
    if uncounted_count:
        messages.error(request, f"Chưa thể hoàn tất đợt kiểm kê: còn {uncounted_count} tài sản chưa kiểm đếm.")
        return redirect("portal:inventory_session_detail", pk=pk)
    session.status = "Completed"
    session.recommendation = request.POST.get("recommendation", "Không.")
    session.modified_by = request.user
    session.save()
    messages.success(request, "Đã hoàn tất đợt kiểm kê.")
    return redirect("portal:inventory_session_print", pk=pk)


@login_required
def inventory_session_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")

    machinery = [l for l in lines if l.asset.asset_type.asset_group == "Machinery"]
    other = [l for l in lines if l.asset.asset_type.asset_group == "Other"]

    return render(request, "portal/inventory_session_print.html", {
        "session": session, "machinery": machinery, "other": other,
        "chairman": session.members.filter(role="Chairman").first(),
        "recorder": session.members.filter(role="Recorder").first(),
        "member_list": session.members.filter(role="Member"),
        "all_members": session.members.select_related("employee"),
    })


@login_required
def inventory_session_export_excel(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    session = get_object_or_404(InventorySession, pk=pk)
    lines = session.lines.select_related("asset", "asset__asset_type").order_by("asset__asset_type__asset_group", "asset__asset_code")

    import openpyxl
    from openpyxl.styles import Font, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Kiểm kê"
    bold = Font(bold=True)
    center = Alignment(horizontal="center")

    ws["A1"] = "BIÊN BẢN KIỂM KÊ TRANG THIẾT BỊ"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = session.office_name
    ws["A3"] = f"Thời điểm kiểm kê: {session.inventory_datetime:%H giờ %M, ngày %d tháng %m năm %Y}"

    row = 5
    for m in session.members.select_related("employee"):
        ws.cell(row=row, column=1, value=f"{m.employee.full_name} — {m.get_role_display()}")
        row += 1

    row += 1
    headers = ["STT", "Tên thiết bị", "Mã số", "Loại", "Nơi SD", "Nhãn hiệu", "Số Seri", "Nguồn mua sắm", "Ngày mua", "SL", "ĐVT", "Giá trị", "Tình trạng", "Ghi chú"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = bold
    header_row = row
    row += 1

    def write_group(group_name, lines_group, stt_group):
        nonlocal row
        ws.cell(row=row, column=1, value=stt_group).font = bold
        ws.cell(row=row, column=2, value=group_name).font = bold
        row += 1
        for i, l in enumerate(lines_group, start=1):
            a = l.asset
            values = [i, a.asset_name, a.asset_code, a.asset_type.asset_type_name, a.location, a.brand,
                      a.serial_number, a.funding_source, a.purchase_date.isoformat() if a.purchase_date else "",
                      a.quantity, a.unit, float(a.purchase_value) if a.purchase_value else "", l.get_actual_condition_display(), l.note or ""]
            for c, v in enumerate(values, start=1):
                ws.cell(row=row, column=c, value=v)
            row += 1

    machinery = [l for l in lines if l.asset.asset_type.asset_group == "Machinery"]
    other = [l for l in lines if l.asset.asset_type.asset_group == "Other"]
    write_group("I. Máy móc thiết bị", machinery, "I")
    write_group("II. Thiết bị khác", other, "II")

    row += 1
    ws.cell(row=row, column=2, value=f"* Kết quả kiểm kê: có {len(machinery)} máy móc thiết bị và {len(other)} thiết bị khác.")
    row += 1
    ws.cell(row=row, column=2, value=f"* Kiến nghị: {session.recommendation}")

    for c in range(1, 15):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 18

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="BBKK_{session.session_name}.xlsx"'
    wb.save(response)
    return response


@login_required
def disposal_request_list(request):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    requests_qs = AssetDisposalRequest.objects.order_by("-request_date")
    return render(request, "portal/disposal_request_list.html", {
        "requests": requests_qs, "can_manage_asset": _can_manage_asset(request),
    })


@login_required
def disposal_request_create(request):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:disposal_request_list")
    employee = _get_employee(request)
    if request.method == "POST":
        try:
            disposal_request = AssetDisposalRequest.objects.create(
                request_no=request.POST.get("request_no"),
                request_date=request.POST.get("request_date"),
                location_name=request.POST.get("location_name"),
                disposal_type=request.POST.get("disposal_type", "Liquidate"),
                requested_by=employee, created_by=request.user,
            )
            messages.success(request, "Đã tạo đề nghị thanh lý.")
            return redirect("portal:disposal_request_detail", pk=disposal_request.pk)
        except Exception as e:
            _friendly_error(request, e, "Không thể tạo.")
    return render(request, "portal/disposal_request_form.html", {})


def _can_approve_disposal(request):
    """C3 (đã sửa): quyền duyệt thanh lý CHỈ dành cho Giám đốc/superuser — tách hẳn khỏi
    `_can_manage_asset` (HR_OFFICER trở lên, chỉ nên được thêm thành viên/tài sản vào biên bản,
    không được tự duyệt). Dùng chung logic với biến `can_approve` đã hiển thị ở template."""
    return _director_employees().filter(user_account=request.user).exists() or request.user.is_superuser


@login_required
def disposal_request_detail(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)

    if request.method == "POST":
        action = request.POST.get("action")
        if action != "approve" and not _can_manage_asset(request):
            messages.error(request, "Anh/chị chỉ có quyền xem thông tin tài sản.")
            return redirect("portal:disposal_request_detail", pk=pk)
        if action == "add_member":
            try:
                member = AssetDisposalRequestMember(
                    request=disposal_request, employee_id=request.POST.get("employee"),
                    participation=request.POST.get("participation", "Direct"),
                    role=request.POST.get("role", "Member"), created_by=request.user,
                )
                member.full_clean()
                member.save()
                messages.success(request, "Đã thêm thành viên.")
            except Exception as e:
                _friendly_error(request, e, "Không thể thêm thành viên.")
        elif action == "add_item":
            try:
                item = AssetDisposalItem(
                    request=disposal_request, asset_id=request.POST.get("asset"),
                    damage_description=request.POST.get("damage_description"), created_by=request.user,
                )
                item.full_clean()
                item.save()
                messages.success(request, "Đã thêm tài sản vào đề nghị thanh lý.")
            except Exception as e:
                _friendly_error(request, e, "Không thể thêm tài sản.")
        elif action == "approve":
            # C3 (đã sửa): trước đây chỉ check `_can_manage_asset` (HR_OFFICER trở lên) rồi set
            # thẳng status="Approved" — biến `can_approve` (Director-only) chỉ dùng ẩn/hiện nút ở
            # template, ai POST tay với `action=approve` đều tự duyệt được. Giờ kiểm tra lại
            # `_can_approve_disposal` ngay server-side, và chỉ duyệt đơn đang "Pending".
            if not _can_approve_disposal(request):
                messages.error(request, "Chỉ Giám đốc mới có quyền duyệt đề nghị thanh lý.")
                return redirect("portal:disposal_request_detail", pk=pk)
            # transaction.atomic() + select_for_update(): tránh 2 người cùng duyệt 1 đề
            # nghị thanh lý cùng lúc (mục 3.C).
            with transaction.atomic():
                disposal_request = get_object_or_404(
                    AssetDisposalRequest.objects.select_for_update(), pk=pk, status="Pending",
                )
                disposal_request.status = "Approved"
                disposal_request.modified_by = request.user
                disposal_request.save()
            submitter = disposal_request.created_by.employee if hasattr(disposal_request.created_by, "employee") else None
            notify(
                submitter, "asset", "Đề nghị thanh lý tài sản đã được duyệt",
                f"Đề nghị {disposal_request.request_no} đã được duyệt.",
                link=f"/portal/disposal/{disposal_request.pk}/", created_by=request.user,
            )
            messages.success(request, "Đã duyệt đề nghị thanh lý.")
        return redirect("portal:disposal_request_detail", pk=pk)

    can_manage_asset = _can_manage_asset(request)
    return render(request, "portal/disposal_request_detail.html", {
        "req": disposal_request, "members": disposal_request.members.select_related("employee"),
        "items": disposal_request.items.select_related("asset"),
        "employees": Employee.objects.filter(employment_status="Active") if can_manage_asset else Employee.objects.none(),
        "assets": Asset.objects.exclude(status="Disposed") if can_manage_asset else Asset.objects.none(),
        "directors": _director_employees() if can_manage_asset else Employee.objects.none(),
        "can_approve": _can_approve_disposal(request),
        "can_manage_asset": can_manage_asset,
    })


@login_required
@require_POST
def disposal_minutes_create(request, item_pk):
    if not _can_manage_asset(request):
        messages.error(request, "Anh/chị không có quyền.")
        return redirect("portal:dashboard")
    item = get_object_or_404(AssetDisposalItem, pk=item_pk)
    try:
        AssetDisposalMinutes.objects.create(
            item=item,
            decision_date=request.POST.get("decision_date"),
            origin_country=request.POST.get("origin_country", ""),
            manufacture_year=request.POST.get("manufacture_year") or None,
            put_into_use_year=request.POST.get("put_into_use_year") or None,
            original_value=request.POST.get("original_value") or 0,
            depreciated_value=request.POST.get("depreciated_value") or 0,
            conclusion=request.POST.get("conclusion"),
            disposal_cost=request.POST.get("disposal_cost") or 0,
            recovery_value=request.POST.get("recovery_value") or 0,
            director_signer_id=request.POST.get("director_signer") or None,
            accountant_signer_id=request.POST.get("accountant_signer") or None,
            created_by=request.user,
        )
        messages.success(request, "Đã lập biên bản thanh lý. Tài sản đã chuyển trạng thái 'Đã thanh lý'.")
    except Exception as e:
        _friendly_error(request, e, "Không thể lưu.")
    return redirect("portal:disposal_request_detail", pk=item.request_id)


@login_required
def disposal_minutes_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    minutes = get_object_or_404(AssetDisposalMinutes, pk=pk)
    return render(request, "portal/disposal_minutes_print.html", {
        "minutes": minutes, "item": minutes.item, "req": minutes.item.request,
        "committee": minutes.item.request.members.select_related("employee"),
    })


@login_required
def disposal_request_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_request_print.html", {
        "req": disposal_request, "members": disposal_request.members.select_related("employee"),
        "items": disposal_request.items.select_related("asset"),
    })


@login_required
def disposal_decision_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_decision_print.html", {
        "req": disposal_request, "items": disposal_request.items.select_related("asset"),
    })


@login_required
def disposal_execution_print(request, pk):
    if not _can_view_asset(request):
        messages.error(request, "Anh/chị không có quyền xem thông tin tài sản.")
        return redirect("portal:dashboard")
    disposal_request = get_object_or_404(AssetDisposalRequest, pk=pk)
    return render(request, "portal/disposal_execution_print.html", {
        "req": disposal_request, "items": disposal_request.items.select_related("asset"),
        "chairman": disposal_request.members.filter(role="Chairman").first(),
        "member_list": disposal_request.members.filter(role="Member"),
    })
