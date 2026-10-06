from datetime import date, timedelta
from openpyxl import load_workbook
from django.test import SimpleTestCase, TestCase
from apps.portal import tests as portal_tests
from apps.contract.models import Contract, ContractType

from apps.reporting.services import (
    ComplianceReportNotConfigured,
    generate_insurance_change_report_excel,
    generate_payroll_cost_by_department_excel,
    generate_pit_finalization_data,
    generate_expiring_contracts_excel,
)


class ComplianceReportScaffoldTests(SimpleTestCase):
    def test_bhxh_export_fails_closed_until_template_is_confirmed(self):
        with self.assertRaises(ComplianceReportNotConfigured):
            generate_insurance_change_report_excel("2026-09", [])

    def test_pit_mapping_fails_closed_until_year_template_is_confirmed(self):
        with self.assertRaises(ComplianceReportNotConfigured):
            generate_pit_finalization_data(2026)

    def test_department_payroll_export_fails_closed_until_allocation_is_confirmed(self):
        with self.assertRaises(ComplianceReportNotConfigured):
            generate_payroll_cost_by_department_excel("2026-09")


class ContractExportBoundaryTests(TestCase):
    setUpTestData = classmethod(portal_tests.PortalSensitiveWorkflowTests.setUpTestData.__func__)
    make_employee = portal_tests.PortalSensitiveWorkflowTests.make_employee

    def test_upcoming_export_excludes_past_and_beyond_cutoff(self):
        employee = self.make_employee('EXPORT-BOUNDARY')
        contract_type = ContractType.objects.create(contract_type_code='EXPORT-TEST', contract_type_name='Synthetic')
        today = date.today()
        for code, offset in [('PAST', -1), ('TODAY', 0), ('CUTOFF', 60), ('LATER', 61)]:
            Contract.objects.create(employee=employee, contract_type=contract_type, contract_no=code,
                start_date=today - timedelta(days=100), end_date=today + timedelta(days=offset),
                basic_salary=100, status='Active')
        workbook = load_workbook(generate_expiring_contracts_excel(days_threshold=60))
        codes = [row[1].value for row in workbook.active.iter_rows(min_row=4)]
        self.assertEqual(codes, ['TODAY', 'CUTOFF'])
