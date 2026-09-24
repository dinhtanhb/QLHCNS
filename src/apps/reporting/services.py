import io
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.hrm.models import Employee
from apps.contract.models import Contract
from apps.document.models import OfficialDispatch


def _style_header_cell(cell, text):
    cell.value = text
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_table_borders(ws, start_row, end_row, start_col, end_col):
    thin = Side(border_style="thin", color="D9D9D9")
    table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(row=r, column=c).border = table_border


def _auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)


def generate_employee_list_excel():
    """Xuất danh sách nhân sự toàn công ty sang Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách nhân viên"
    ws.views.sheetView[0].showGridLines = True

    # Tiêu đề báo cáo
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "DANH SÁCH NHÂN SỰ DOANH NGHIỆP (IHRMS)"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["A2"] = f"Ngày xuất báo cáo: {date.today().strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(name="Arial", size=10, italic=True)

    headers = [
        "STT", "Mã nhân viên", "Họ và tên", "Giới tính", "Ngày sinh",
        "Số CCCD", "Phòng ban", "Vị trí", "Ngày vào làm", "Trạng thái"
    ]
    header_row = 4
    ws.row_dimensions[header_row].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        _style_header_cell(cell, header)

    employees = Employee.objects.select_related("department", "position").order_by("department__department_name", "employee_code")
    current_row = 5
    for idx, emp in enumerate(employees, 1):
        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=emp.employee_code).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=emp.full_name)
        ws.cell(row=current_row, column=4, value=emp.get_gender_display()).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=5, value=emp.date_of_birth.strftime("%d/%m/%Y") if emp.date_of_birth else "").alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=6, value=emp.national_id).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=7, value=emp.department.department_name if emp.department else "")
        ws.cell(row=current_row, column=8, value=emp.position.position_name if emp.position else "")
        ws.cell(row=current_row, column=9, value=emp.join_date.strftime("%d/%m/%Y") if emp.join_date else "").alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=10, value=emp.get_employment_status_display()).alignment = Alignment(horizontal="center")
        current_row += 1

    _apply_table_borders(ws, 4, current_row - 1, 1, len(headers))
    _auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_expiring_contracts_excel(days_threshold=60):
    """Xuất danh sách hợp đồng lao động sắp hết hạn."""
    wb = Workbook()
    ws = wb.active
    ws.title = "HĐ sắp hết hạn"
    ws.views.sheetView[0].showGridLines = True

    limit_date = date.today() + timedelta(days=days_threshold)
    contracts = Contract.objects.filter(
        status="Active", end_date__isnull=False, end_date__lte=limit_date
    ).select_related("employee", "employee__department", "contract_type").order_by("end_date")

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"DANH SÁCH HỢP ĐỒNG SẮP HẾT HẠN (TRONG VÒNG {days_threshold} NGÀY)"
    title_cell.font = Font(name="Arial", size=13, bold=True, color="C00000")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = [
        "STT", "Số hợp đồng", "Mã NV", "Họ tên nhân viên",
        "Phòng ban", "Loại hợp đồng", "Ngày kết thúc", "Số ngày còn lại"
    ]
    header_row = 3
    ws.row_dimensions[header_row].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        _style_header_cell(cell, header)

    current_row = 4
    for idx, c in enumerate(contracts, 1):
        days_left = (c.end_date - date.today()).days if c.end_date else 0
        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=c.contract_no).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=c.employee.employee_code).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=4, value=c.employee.full_name)
        ws.cell(row=current_row, column=5, value=c.employee.department.department_name if c.employee.department else "")
        ws.cell(row=current_row, column=6, value=c.contract_type.contract_type_name if c.contract_type else "")
        ws.cell(row=current_row, column=7, value=c.end_date.strftime("%d/%m/%Y") if c.end_date else "").alignment = Alignment(horizontal="center")
        
        days_cell = ws.cell(row=current_row, column=8, value=days_left)
        days_cell.alignment = Alignment(horizontal="center")
        if days_left <= 15:
            days_cell.font = Font(color="FF0000", bold=True)
        current_row += 1

    if current_row > 4:
        _apply_table_borders(ws, 3, current_row - 1, 1, len(headers))
    _auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_dispatch_book_excel(dispatch_type="Incoming", year=None):
    """Xuất Sổ công văn đến hoặc Sổ công văn đi sang file Excel chuẩn lưu trữ."""
    wb = Workbook()
    ws = wb.active
    type_name = "CÔNG VĂN ĐẾN" if dispatch_type == "Incoming" else "CÔNG VĂN ĐI"
    ws.title = f"Sổ {type_name}"
    ws.views.sheetView[0].showGridLines = True

    qs = OfficialDispatch.objects.filter(dispatch_type=dispatch_type)
    if year:
        qs = qs.filter(received_or_sent_date__year=year)
    dispatches = qs.select_related("department", "assigned_to").order_by("dispatch_no")

    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"SỔ THEO DÕI {type_name} NĂM {year or date.today().year}"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    if dispatch_type == "Incoming":
        headers = [
            "Số đến", "Ngày đến", "Tác giả / Cơ quan gửi", "Số ký hiệu gốc",
            "Ngày ban hành gốc", "Trích yếu nội dung", "Người / Đơn vị nhận", "Hạn xử lý", "Trạng thái"
        ]
    else:
        headers = [
            "Số đi", "Ngày gửi", "Đơn vị soạn thảo", "Người ký",
            "Nơi nhận văn bản", "Trích yếu nội dung", "Độ khẩn", "Độ mật", "Trạng thái"
        ]

    header_row = 3
    ws.row_dimensions[header_row].height = 25
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        _style_header_cell(cell, header)

    current_row = 4
    for d in dispatches:
        ws.cell(row=current_row, column=1, value=d.dispatch_no).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=d.received_or_sent_date.strftime("%d/%m/%Y") if d.received_or_sent_date else "").alignment = Alignment(horizontal="center")
        if dispatch_type == "Incoming":
            ws.cell(row=current_row, column=3, value=d.issuing_organization)
            ws.cell(row=current_row, column=4, value=d.original_doc_no or "").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=5, value=d.issue_date.strftime("%d/%m/%Y") if d.issue_date else "").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=6, value=d.title)
            dest = f"{d.department.department_name if d.department else ''} ({d.assigned_to.full_name if d.assigned_to else ''})"
            ws.cell(row=current_row, column=7, value=dest)
            ws.cell(row=current_row, column=8, value=d.due_date.strftime("%d/%m/%Y") if d.due_date else "").alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=9, value=d.get_status_display()).alignment = Alignment(horizontal="center")
        else:
            ws.cell(row=current_row, column=3, value=d.department.department_name if d.department else "")
            ws.cell(row=current_row, column=4, value=d.signer.full_name if d.signer else "")
            ws.cell(row=current_row, column=5, value=d.issuing_organization)
            ws.cell(row=current_row, column=6, value=d.title)
            ws.cell(row=current_row, column=7, value=d.get_urgency_level_display()).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=8, value=d.get_security_level_display()).alignment = Alignment(horizontal="center")
            ws.cell(row=current_row, column=9, value=d.get_status_display()).alignment = Alignment(horizontal="center")
        current_row += 1

    if current_row > 4:
        _apply_table_borders(ws, 3, current_row - 1, 1, len(headers))
    _auto_fit_columns(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
