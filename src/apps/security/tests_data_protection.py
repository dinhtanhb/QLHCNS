"""Regression tests for sensitive data access and administrative bypasses."""
import tempfile
from datetime import date
from decimal import Decimal
from django.test import TestCase, override_settings
from django.urls import reverse
from django.contrib import admin
from django.core.files.base import ContentFile
from apps.portal import tests as fixtures
from apps.security.models import User
from apps.hrm.models import Employee
from apps.payroll.models import EmployeeSalary, PayrollPeriod
from apps.leave.models import LeaveRequest, LeaveType


class SensitiveDataTests(TestCase):
    make_employee = fixtures.PortalSensitiveWorkflowTests.make_employee
    make_user = fixtures.PortalSensitiveWorkflowTests.make_user

    @classmethod
    def setUpTestData(cls):
        fixtures.PortalSensitiveWorkflowTests.setUpTestData.__func__(cls)

    def setUp(self):
        self.owner = self.make_employee('SEC-OWNER')
        self.other = self.make_employee('SEC-OTHER')
        self.owner_user = self.make_user('EMPLOYEE', self.owner, 'sec-owner')
        self.other_user = self.make_user('EMPLOYEE', self.other, 'sec-other')
        self.sys_user = self.make_user('SYS_ADMIN', self.make_employee('SEC-SYS'), 'sec-sys')
        self.sys_user.is_staff = True
        self.sys_user.save(update_fields=['is_staff'])
        self.media = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name, M365_STORAGE_ENABLED=False)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.addCleanup(self.media.cleanup)

    def test_system_admin_cannot_assign_superuser_permissions(self):
        self.client.force_login(self.sys_user)
        response = self.client.get(reverse('admin:security_user_change', args=(self.owner_user.pk,)))
        self.assertEqual(response.status_code, 200)
        fields = response.context['adminform'].form.fields
        for field in ['is_superuser', 'groups', 'user_permissions', 'role', 'is_staff', 'employee']:
            self.assertNotIn(field, fields)

    def test_view_only_admin_cannot_soft_delete_employee_using_action(self):
        director = self.make_user('DIRECTOR', self.make_employee('SEC-DIRECTOR'), 'sec-director')
        director.is_staff = True
        director.save(update_fields=['is_staff'])
        self.client.force_login(director)
        self.client.post(reverse('admin:hrm_employee_changelist'), {'action':'soft_delete_selected', '_selected_action':str(self.owner.pk), 'select_across':'0'})
        self.assertTrue(Employee.objects.filter(pk=self.owner.pk).exists())

    def test_locked_payroll_period_cannot_be_deleted_by_bulk_action(self):
        finance = self.make_user('FINANCE_OFFICER', self.make_employee('SEC-FINANCE'), 'sec-finance')
        finance.is_staff = True
        finance.save(update_fields=['is_staff'])
        period = PayrollPeriod.objects.create(payroll_month=10, payroll_year=2026, start_date=date(2026,10,1), end_date=date(2026,10,31), status='Locked')
        self.client.force_login(finance)
        self.client.post(reverse('admin:payroll_payrollperiod_changelist'), {'action':'soft_delete_selected', '_selected_action':str(period.pk), 'select_across':'0'})
        self.assertTrue(PayrollPeriod.objects.filter(pk=period.pk).exists())

    def test_department_manager_does_not_receive_salary_objects_in_context(self):
        manager = self.make_user('DEPT_MANAGER', self.make_employee('SEC-MANAGER'), 'sec-manager')
        self.department.manager = manager.employee
        self.department.save(update_fields=['manager'])
        EmployeeSalary.objects.create(employee=self.owner, current_basic_salary=Decimal('12345678'), effective_date=date(2026,1,1))
        self.client.force_login(manager)
        response = self.client.get(reverse('portal:hr_employee_detail', args=(self.owner.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['salaries'])
        self.assertFalse(response.context['contracts'])

    def test_system_admin_cannot_export_personnel_identifiers(self):
        self.client.force_login(self.sys_user)
        response = self.client.get(reverse('reporting:export_employees'))
        self.assertEqual(response.status_code, 403)

    def test_profile_response_is_not_cacheable(self):
        self.client.force_login(self.owner_user)
        response = self.client.get(reverse('portal:profile'))
        self.assertIn('no-store', response.get('Cache-Control', ''))

    def test_attachment_url_requires_owner_or_hr_and_hides_deleted_rows(self):
        kind = LeaveType.objects.create(leave_type_code='SEC-TEST', leave_type_name='Test')
        leave = LeaveRequest.objects.create(employee=self.owner, leave_type=kind, from_date=date(2026,11,2), to_date=date(2026,11,2))
        leave.attachment.save('secret.pdf', ContentFile(b'%PDF-1.4 fake private file'))
        url = leave.attachment.url
        self.assertTrue(url.startswith('/private-files/'))
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.other_user)
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.force_login(self.owner_user)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('no-store', response['Cache-Control'])
        response.close()
        leave.delete()
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_raw_media_path_is_never_public_even_in_debug(self):
        with override_settings(DEBUG=True):
            self.assertEqual(self.client.get('/media/secret.pdf').status_code, 404)
        self.client.force_login(self.owner_user)
        for path in ['/private-files/not-recorded.pdf', '/private-files/../secret.pdf']:
            self.assertEqual(self.client.get(path).status_code, 404)

    def test_system_admin_cannot_reset_privileged_password(self):
        hr = self.make_user('HR_MANAGER', self.make_employee('SEC-HR'), 'sec-hr')
        self.client.force_login(self.sys_user)
        response = self.client.post(reverse('admin:auth_user_password_change', args=(hr.pk,)),
                                    {'password1':'new-password-test-123!', 'password2':'new-password-test-123!'})
        self.assertIn(response.status_code, (302, 403, 404))
        hr.refresh_from_db()
        self.assertFalse(hr.check_password('new-password-test-123!'))

    def test_system_admin_employee_view_omits_identifiers(self):
        self.client.force_login(self.sys_user)
        response = self.client.get(reverse('admin:hrm_employee_change', args=(self.owner.pk,)))
        fields = [name for _, options in response.context['adminform'].fieldsets for name in options['fields']]
        for name in ('national_id', 'tax_code', 'social_insurance_no'):
            self.assertNotIn(name, fields)

    def test_idle_session_expires(self):
        import time
        self.client.force_login(self.owner_user)
        session = self.client.session
        session['security_last_activity'] = int(time.time()) - 1801
        session.save()
        self.assertEqual(self.client.get(reverse('portal:profile')).status_code, 302)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_five_failures_lock_username_even_after_ip_changes(self):
        self.owner_user.set_password('secure-test-password-123!')
        self.owner_user.save()
        for number in range(5):
            self.client.post(reverse('portal:login'),
                {'username':self.owner_user.username, 'password':'wrong-password'},
                REMOTE_ADDR='192.0.2.' + str(number + 1))
        response = self.client.post(reverse('portal:login'),
            {'username':self.owner_user.username, 'password':'secure-test-password-123!'},
            REMOTE_ADDR='192.0.2.99')
        self.assertEqual(response.status_code, 429)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_export_treats_user_formula_as_literal_text(self):
        from openpyxl import load_workbook
        from apps.reporting.services import generate_employee_list_excel
        self.owner.full_name = '=HYPERLINK("https://example.invalid","open")'
        self.owner.save(update_fields=['full_name'])
        workbook = load_workbook(generate_employee_list_excel())
        cell = next(cell for row in workbook.active.iter_rows() for cell in row if cell.value == self.owner.full_name)
        self.assertEqual(cell.data_type, 's')

    def test_production_exception_logs_do_not_contain_payloads(self):
        import logging
        import sys
        from apps.security.logging import SafeExceptionFilter
        try:
            raise RuntimeError('password=never-persist-this-sensitive-value')
        except RuntimeError:
            record = logging.LogRecord('test', logging.ERROR, __file__, 1, 'unsafe %s', ('private',), sys.exc_info())
        with override_settings(DEBUG=False):
            SafeExceptionFilter().filter(record)
        rendered = logging.Formatter().format(record)
        self.assertIn('RuntimeError', rendered)
        self.assertNotIn('never-persist', rendered)
        self.assertNotIn('private', rendered)
        self.assertIsNone(record.exc_info)

    def test_authorized_hr_export_creates_metadata_audit_event(self):
        from apps.security.models import SecurityEvent
        hr = self.make_user('HR_MANAGER', self.make_employee('SEC-HR-EXPORT'), 'sec-hr-export')
        self.client.force_login(hr)
        response = self.client.get(reverse('reporting:export_employees'))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(SecurityEvent.objects.filter(user=hr, action='data_export', resource='export_employees').exists())

    def test_throttle_expires_after_cooloff(self):
        from axes.models import AccessAttempt
        from django.utils import timezone
        from datetime import timedelta
        self.owner_user.set_password('secure-test-password-123!')
        self.owner_user.save()
        for number in range(5):
            self.client.post(reverse('portal:login'), {'username':self.owner_user.username, 'password':'wrong'})
        AccessAttempt.objects.update(attempt_time=timezone.now() - timedelta(minutes=31))
        response = self.client.post(reverse('portal:login'), {'username':self.owner_user.username, 'password':'secure-test-password-123!'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(int(self.client.session['_auth_user_id']), self.owner_user.pk)

    def test_confidential_dispatch_scope_ignores_staff_flag(self):
        from apps.document.models import OfficialDispatch, DispatchBook
        from apps.security.data_access import visible_dispatches_for
        book = DispatchBook.objects.create(book_code='SEC-BOOK', book_name='Test', book_type='Incoming', year=2026)
        dispatch = OfficialDispatch.objects.create(dispatch_type='Incoming', dispatch_book=book, dispatch_no='SEC-DOC',
            title='Synthetic confidential', issuing_organization='Test', issue_date=date(2026,10,1),
            received_or_sent_date=date(2026,10,1), security_level='Confidential', department=self.department)
        self.other_user.is_staff = True
        self.other_user.save()
        self.assertFalse(visible_dispatches_for(self.other_user).filter(pk=dispatch.pk).exists())
        dispatch.assigned_to = self.other
        dispatch.save(update_fields=['assigned_to'])
        self.assertTrue(visible_dispatches_for(self.other_user).filter(pk=dispatch.pk).exists())

    def test_sharepoint_url_is_private_without_graph_request(self):
        from unittest.mock import patch
        from apps.m365.storage import ConfiguredStorage, SharePointStorage
        with override_settings(M365_STORAGE_ENABLED=True), patch('apps.m365.storage.get_graph_client') as graph:
            self.assertEqual(ConfiguredStorage('dispatches').url('dispatches/private.pdf'), '/private-files/dispatches/private.pdf')
            remote = SharePointStorage('Documents')
            self.assertEqual(remote.url('dispatches/private.pdf'), '/private-files/dispatches/private.pdf')
            graph.return_value.get_web_url.assert_not_called()

    def test_document_link_rejects_active_content_and_raw_media(self):
        from apps.document.models import Document
        for url in ['javascript:alert(1)', 'data:text/html,test', '//external.invalid/path', '/media/secret.pdf']:
            self.assertEqual(Document(document_url=url).safe_document_url, '')
        self.assertEqual(Document(document_url='https://example.invalid/policy.pdf').safe_document_url, 'https://example.invalid/policy.pdf')

    def test_system_admin_cannot_reset_staff_employee_password(self):
        self.other_user.is_staff = True
        self.other_user.save()
        self.client.force_login(self.sys_user)
        self.client.post(reverse('admin:auth_user_password_change', args=(self.other_user.pk,)),
            {'password1':'new-password-test-123!', 'password2':'new-password-test-123!'})
        self.other_user.refresh_from_db()
        self.assertFalse(self.other_user.check_password('new-password-test-123!'))
