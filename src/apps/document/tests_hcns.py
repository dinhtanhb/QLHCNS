import io
from datetime import date
from decimal import Decimal
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from django.test import TestCase
from django.core.exceptions import ValidationError
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.asset.models import Asset, AssetType, InventoryLine, InventorySession
from apps.hrm.models import Employee
from apps.leave.models import LeaveRequest, LeaveType
from apps.organization.models import Department, Organization, Position
from apps.security.models import Role, SecurityEvent, User

from .hcns.catalog import FORMS, fields_for, tables_for
from .hcns.exports import export_xlsx
from .hcns.forms import initial_document_data
from .hcns.rendering import document_context, scored_rows
from .hcns.money import amount_in_words


class HCNSFormTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        org = Organization.objects.create(organization_code="HCNS-TEST", organization_name="Tổ chức giả")
        cls.department = Department.objects.create(department_code="HCNS-TEST", department_name="Phòng giả", organization=org)
        cls.position = Position.objects.create(position_code="HCNS-TEST", position_name="Chuyên viên giả")
        cls.employees = []
        for i in range(3):
            cls.employees.append(Employee.objects.create(employee_code=f"HCNS-{i}", first_name="Giả", last_name="Kiểm thử", full_name=f"Nhân viên giả {i}",
                date_of_birth=date(1990, 1, 1), gender="Other", national_id=f"SYNTHETIC-{i}", company_email=f"hcns-{i}@example.test",
                mobile_phone="0000000000", join_date=date(2020, 1, 1), department=cls.department, position=cls.position))
        cls.hr = cls.user("hcns-hr", "HR_OFFICER", cls.employees[0])
        cls.staff = cls.user("hcns-staff", "EMPLOYEE", cls.employees[1])
        cls.admin = cls.user("hcns-admin", "SYS_ADMIN", cls.employees[2])
        cls.superuser = User.objects.create_superuser(username="hcns-superuser", password="synthetic-test", email="root@example.test")

    @classmethod
    def user(cls, username, role_code, employee):
        role, _ = Role.objects.get_or_create(role_code=role_code, defaults={"role_name": role_code})
        return User.objects.create_user(username=username, password="synthetic-test", employee=employee, role=role)

    def url(self, slug):
        return reverse("portal:hcns_draft", args=[slug])

    def data(self, slug, **overrides):
        spec = FORMS[slug]
        result = {field["key"]: field.get("initial", "") for field in fields_for(spec)}
        result["action"] = "preview"
        for table in tables_for(spec):
            key = table["key"]
            result.update({f"{key}-TOTAL_FORMS": str(len(table["rows"])), f"{key}-INITIAL_FORMS": str(len(table["rows"])),
                f"{key}-MIN_NUM_FORMS": "0", f"{key}-MAX_NUM_FORMS": str(table["max_rows"])})
            for i, row in enumerate(table["rows"]):
                for col in table["columns"]:
                    if not col.get("locked"):
                        result[f"{key}-{i}-{col['key']}"] = row.get(col["key"], "")
        result.update(overrides)
        return result

    def test_all_layouts_and_exports_open_without_originals_or_real_records(self):
        self.client.force_login(self.superuser)
        for slug, spec in FORMS.items():
            with self.subTest(slug=slug):
                response = self.client.get(self.url(slug))
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "BẢN SOẠN THẢO")
                response = self.client.post(self.url(slug), self.data(slug, action="export"))
                self.assertEqual(response.status_code, 200)
                self.assertIn("no-store", response["Cache-Control"])
                if spec["kind"] == "docx":
                    with ZipFile(io.BytesIO(response.content)) as archive:
                        root = ET.fromstring(archive.read("word/document.xml"))
                        self.assertIn("BẢN SOẠN THẢO", "".join(root.itertext()))
                else:
                    wb = load_workbook(io.BytesIO(response.content))
                    self.assertGreater(wb.active.max_row, 5)
                    self.assertFalse(wb.vba_archive)
        self.assertEqual(SecurityEvent.objects.filter(action="data_export").count(), len(FORMS))

    def test_login_own_personal_export_and_foreign_employee_denial(self):
        self.assertEqual(self.client.get(self.url("employee-information")).status_code, 302)
        self.client.force_login(self.staff)
        own = self.client.get(self.url("employee-information"), {"employee_id": self.employees[1].pk})
        self.assertContains(own, "SYNTHETIC-1")
        foreign = self.client.get(self.url("employee-information"), {"employee_id": self.employees[0].pk, "export": "1"})
        self.assertEqual(foreign.status_code, 403)
        self.assertNotContains(foreign, "SYNTHETIC-0", status_code=403)
        self.assertEqual(self.client.get(self.url("salary-decision")).status_code, 403)
        self.assertEqual(self.client.get(self.url("asset-inventory")).status_code, 403)

    def test_sysadmin_has_no_hr_or_compensation_export_access(self):
        self.client.force_login(self.admin)
        for slug in ("employee-evaluation", "candidate-interview", "salary-scale", "termination-decision"):
            self.assertEqual(self.client.get(self.url(slug)).status_code, 403)
        self.assertEqual(self.client.get(self.url("employee-information"), {"employee_id": self.employees[0].pk}).status_code, 403)

    def test_leave_prefill_scoped_and_type_checked_without_changing_status(self):
        leave_type, _ = LeaveType.objects.get_or_create(leave_type_code="AL", defaults={"leave_type_name": "Phép năm", "is_paid": True})
        obj = LeaveRequest.objects.create(employee=self.employees[1], leave_type=leave_type, from_date=date(2026, 11, 2), to_date=date(2026, 11, 3), reason="Lý do giả")
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(self.url("annual-leave"), {"leave_id": obj.pk}), "Lý do giả")
        self.assertEqual(self.client.get(self.url("unpaid-leave"), {"leave_id": obj.pk}).status_code, 404)
        self.client.force_login(self.hr)
        self.assertEqual(self.client.get(self.url("maternity-decision"), {"leave_id": obj.pk}).status_code, 404)
        maternity, _ = LeaveType.objects.get_or_create(leave_type_code="ML", defaults={"leave_type_name": "Thai sản", "is_paid": True})
        obj.leave_type = maternity
        obj.save(update_fields=["leave_type"])
        response = self.client.get(self.url("maternity-decision"), {"leave_id": obj.pk})
        self.assertContains(response, self.employees[1].full_name)
        self.assertContains(response, "02/11/2026")
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(self.url("annual-leave"), {"leave_id": obj.pk}).status_code, 403)
        obj.refresh_from_db()
        self.assertEqual(obj.status, "Pending")

    def test_export_validation_xss_formula_injection_and_audit_metadata(self):
        self.client.force_login(self.hr)
        response = self.client.post(self.url("annual-leave"), self.data("annual-leave", from_date="2026-11-03", to_date="2026-11-02", action="export"))
        self.assertEqual(response.status_code, 422)
        response = self.client.post(self.url("resignation"), self.data("resignation", reason='<script>alert("x")</script>'))
        self.assertContains(response, "&lt;script&gt;")
        self.assertNotContains(response, '<script>alert("x")</script>')
        response = self.client.post(self.url("resignation"), self.data("resignation", reason="Tên giả\ufffe\uffff", action="export"))
        with ZipFile(io.BytesIO(response.content)) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
            self.assertIn("Tên giả", "".join(root.itertext()))
        response = self.client.post(self.url("offboarding-handover"), self.data("offboarding-handover", address="=HYPERLINK(\"https://example.test\")", action="export"))
        wb = load_workbook(io.BytesIO(response.content))
        injected = next(c for row in wb.active for c in row if isinstance(c.value, str) and c.value.startswith("=HYPERLINK"))
        self.assertEqual(injected.data_type, "s")
        event = SecurityEvent.objects.latest("pk")
        self.assertEqual(event.resource, "hcns:offboarding-handover")
        self.assertEqual(event.record_id, "")

    def test_score_average_both_types_zero_and_missing_values(self):
        rows = []
        for interviewer, score in (("A", Decimal(10)), ("B", Decimal(0))):
            row = {"candidate_code": "TEST-1", "candidate": "Ứng viên giả", "interviewer": interviewer}
            row.update({key: score for key in ("expertise", "experience", "communication", "english", "interest", "other", "plan_test", "budget_test")})
            rows.append(row)
        result = scored_rows(rows)
        self.assertEqual(result[0]["average_interview"], Decimal(30))
        self.assertEqual(result[0]["average_test"], Decimal(10))
        self.assertEqual(result[0]["combined"], Decimal(40))
        rows[1]["plan_test"] = None
        self.assertIsNone(scored_rows(rows)[0]["average_test"])
        self.assertIsNone(scored_rows(rows)[0]["combined"])
        spec = FORMS["candidate-score-summary"]
        values, _ = initial_document_data(spec, {}, {})
        wb = load_workbook(io.BytesIO(export_xlsx(document_context(spec, values, {"results": result}))))
        formulas = [c.value for row in wb.active for c in row if c.data_type == "f"]
        self.assertEqual(sum("AVERAGE(" in f for f in formulas), 4)
        self.assertTrue(all(f.startswith("=IF(COUNT(") for f in formulas))

    def test_score_ceiling_and_fixed_criteria_cannot_be_tampered(self):
        self.client.force_login(self.hr)
        response = self.client.post(self.url("candidate-interview"), self.data("candidate-interview", **{"scores-0-score": "21", "scores-0-maximum": "999", "action": "export"}))
        self.assertEqual(response.status_code, 422)
        response = self.client.post(self.url("candidate-interview"), self.data("candidate-interview", **{"scores-TOTAL_FORMS": "999999", "action": "export"}))
        self.assertEqual(response.status_code, 422)

    def test_inventory_prefill_and_large_source_exports_all_rows_without_macro(self):
        obj = InventorySession.objects.create(session_name="Đợt giả", office_name="Văn phòng giả", inventory_datetime=timezone.now())
        kind = AssetType.objects.create(asset_type_code="HCNS-TEST", asset_type_name="Thiết bị giả")
        for i in range(51):
            asset = Asset.objects.create(asset_code=f"HCNS-ASSET-{i}", asset_name=f"Tài sản giả {i}", asset_type=kind)
            InventoryLine.objects.create(session=obj, asset=asset)
        self.client.force_login(self.hr)
        response = self.client.get(self.url("asset-inventory"), {"inventory_id": obj.pk})
        self.assertContains(response, "Tài sản giả 50")
        response = self.client.get(self.url("asset-inventory"), {"inventory_id": obj.pk, "export": "1"})
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb["Dán nhãn"].max_row, 51)
        self.assertEqual(wb["Dán nhãn"].cell(51, 1).value, "HCNS-ASSET-50")
        self.assertFalse(wb.vba_archive)
        obj.refresh_from_db()
        self.assertEqual(obj.status, "Draft")

    def test_bad_source_values_show_a_safe_error_instead_of_crashing(self):
        obj = InventorySession.objects.create(session_name="Đợt dữ liệu lỗi giả", office_name="Văn phòng giả", inventory_datetime=timezone.now())
        kind = AssetType.objects.create(asset_type_code="HCNS-BAD", asset_type_name="Thiết bị giả")
        asset = Asset.objects.create(asset_code="HCNS-BAD", asset_name="Tài sản dữ liệu lỗi giả", asset_type=kind, quantity=-1)
        InventoryLine.objects.create(session=obj, asset=asset)
        self.client.force_login(self.hr)
        response = self.client.get(self.url("asset-inventory"), {"inventory_id": obj.pk, "export": "1"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dữ liệu hồ sơ nguồn chưa phù hợp")
        self.assertNotIn("Content-Disposition", response)

    def test_zero_only_dynamic_rows_are_retained_and_require_score_identity(self):
        self.client.force_login(self.hr)
        response = self.client.post(self.url("asset-inventory"), self.data("asset-inventory", **{"assets-0-quantity": "0"}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["document"]["sections"][-1]["table"]["raw_rows"][0]["quantity"], Decimal(0))
        response = self.client.post(self.url("candidate-score-summary"), self.data("candidate-score-summary", **{"results-0-expertise": "0", "action": "export"}))
        self.assertEqual(response.status_code, 422)
        self.assertNotIn("Content-Disposition", response)

    def test_source_defaults_validate_date_and_time_ranges(self):
        with self.assertRaises(ValidationError):
            initial_document_data(FORMS["annual-leave"], {"from_date": date(2026, 11, 3), "to_date": date(2026, 11, 2)}, {})
        with self.assertRaises(ValidationError):
            initial_document_data(FORMS["asset-destruction-minutes"], {"start_time": "15:00", "end_time": "14:00"}, {})

    def test_appointment_reference_choices_autofill_and_word_layout(self):
        self.client.force_login(self.hr)
        response = self.client.get(self.url("appointment-decision"), {"employee_id": self.employees[1].pk})
        form = response.context["form"]
        self.assertEqual(form["full_name"].value(), self.employees[1].pk)
        self.assertContains(response, 'select name="full_name"')
        self.assertContains(response, 'id="hcns-reference-data"')
        self.assertContains(response, 'hcns-decision-head')
        self.assertNotContains(response, "Nội dung căn cứ</h3>")
        data = self.data("appointment-decision", full_name=self.employees[1].pk, signer=self.employees[0].pk,
                         salary="20000000", salary_words="Sai do người gửi", effective_date="2026-10-01", action="export")
        response = self.client.post(self.url("appointment-decision"), data)
        self.assertEqual(response.status_code, 200)
        with ZipFile(io.BytesIO(response.content)) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        text = "".join(root.itertext())
        self.assertIn(self.employees[1].full_name, text)
        self.assertIn(self.position.position_name, text)
        self.assertIn(self.department.department_name, text)
        self.assertIn("Hai mươi triệu đồng chẵn", text)
        self.assertNotIn("Sai do người gửi", text)
        self.assertEqual(len(root.findall('.//{*}tbl')), 2)
        self.assertNotIn("Ngày lập:", text)
        self.assertIn(self.employees[0].full_name, text)
        data.update(decision_action="Miễn nhiệm", action="preview")
        response = self.client.post(self.url("appointment-decision"), data)
        self.assertContains(response, "Miễn nhiệm chức vụ")
        self.assertContains(response, "thôi giữ chức vụ")

    def test_appointment_rejects_forged_or_deleted_reference_choices(self):
        self.client.force_login(self.hr)
        self.employees[2].delete()
        for key, value in (("full_name", self.employees[2].pk), ("signer", self.employees[2].pk),
                           ("department", "999999"), ("position", "Nhập tuỳ ý"), ("salary", "1.5")):
            with self.subTest(key=key):
                response = self.client.post(self.url("appointment-decision"), self.data("appointment-decision", **{key: value, "action": "export"}))
                self.assertEqual(response.status_code, 422)
                self.assertNotIn("Content-Disposition", response)
        self.department.delete()
        self.position.delete()
        response = self.client.post(self.url("appointment-decision"), self.data("appointment-decision", full_name=self.employees[1].pk, action="export"))
        self.assertEqual(response.status_code, 422)
        self.assertNotIn("Content-Disposition", response)

    def test_whole_dong_words_zero_groups_and_boundaries(self):
        for amount, expected in ((0, "Không đồng chẵn"), (5, "Năm đồng chẵn"), (15, "Mười lăm đồng chẵn"),
                                 (21, "Hai mươi mốt đồng chẵn"), (105, "Một trăm linh năm đồng chẵn"),
                                 (1005, "Một nghìn không trăm linh năm đồng chẵn"), (1000001, "Một triệu không trăm linh một đồng chẵn"),
                                 (20000000, "Hai mươi triệu đồng chẵn"), (1000000000, "Một tỷ đồng chẵn")):
            self.assertEqual(amount_in_words(Decimal(amount)), expected)
        for invalid in (-1, Decimal("0.5"), Decimal("NaN"), 10000000000000):
            with self.assertRaises(ValueError):
                amount_in_words(invalid)

    def test_catalog_selected_employee_reaches_every_template_and_matching_fields(self):
        employee = self.employees[1]
        employee.personal_email = "prefill@example.test"
        employee.address = "Địa chỉ kiểm thử"
        employee.tax_code = "SYN-TAX"
        employee.social_insurance_no = "SYN-INS"
        employee.manager = self.employees[0]
        employee.save()
        self.client.force_login(self.hr)
        query = {"employee_id": employee.pk}
        response = self.client.get(reverse("portal:hcns_catalog"), query)
        self.assertEqual(response.status_code, 200)
        for slug, spec in FORMS.items():
            with self.subTest(slug=slug):
                if slug == "salary-scale":
                    self.assertNotContains(response, self.url(slug))
                    self.assertEqual(self.client.get(self.url(slug), query).status_code, 403)
                    continue
                self.assertContains(response, self.url(slug) + "?employee_id=" + str(employee.pk))
                draft = self.client.get(self.url(slug), query)
                self.assertEqual(draft.status_code, 200)
                form = draft.context["form"]
                expected = {"full_name": employee.full_name, "position": str(self.position), "department": str(self.department),
                    "new_position": str(self.position),
                    "email": employee.personal_email, "address": employee.address, "phone": employee.mobile_phone,
                    "date_of_birth": employee.date_of_birth, "gender": employee.get_gender_display(), "birth_year": employee.date_of_birth.year,
                    "national_id": employee.national_id, "tax_code": employee.tax_code, "insurance_no": employee.social_insurance_no,
                    "join_date": employee.join_date, "manager": self.employees[0].full_name, "manager_position": str(self.position)}
                if slug == "candidate-interview":
                    expected = {"position": self.position.pk}
                else:
                    expected.update(full_name=employee.pk, position=self.position.pk, department=self.department.pk, new_position=self.position.pk)
                    expected.update(manager=self.employees[0].pk, manager_position=self.position.pk)
                for key, value in expected.items():
                    if key in form.fields:
                        self.assertEqual(form[key].value(), value)
                self.assertContains(draft, reverse("portal:hcns_catalog") + "?employee_id=" + str(employee.pk))
                exported = self.client.get(self.url(slug), {**query, "export": "1"})
                self.assertEqual(exported.status_code, 200)
                self.assertIn("Content-Disposition", exported)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(self.url("candidate-interview"), query).status_code, 403)
        self.assertEqual(self.client.get(self.url("employee-information"), {"employee_id": self.employees[0].pk}).status_code, 403)

    def test_employee_prefill_current_salary_contract_and_unavailable_data(self):
        from apps.contract.models import Contract, ContractType
        from apps.payroll.models import EmployeeSalary
        employee = self.employees[1]
        today = timezone.localdate()
        EmployeeSalary.objects.create(employee=employee, current_basic_salary=Decimal("20000000.00"), effective_date=date(2020, 1, 1))
        EmployeeSalary.objects.create(employee=employee, current_basic_salary=Decimal("90000000"), effective_date=date(today.year + 1, 1, 1))
        kind = ContractType.objects.create(contract_type_code="SYN", contract_type_name="Hợp đồng giả")
        Contract.objects.create(employee=employee, contract_type=kind, contract_no="SYN-HĐ", start_date=date(2020, 1, 1), basic_salary=0, status="Active")
        self.client.force_login(self.hr)
        for slug in ("appointment-decision", "salary-decision"):
            response = self.client.get(self.url(slug), {"employee_id": employee.pk})
            self.assertEqual(response.context["form"]["salary"].value(), Decimal("20000000"))
            self.assertEqual(response.context["form"]["salary_words"].value(), "Hai mươi triệu đồng chẵn")
            self.assertEqual(response.context["form"]["effective_date"].value(), "")
        response = self.client.get(self.url("termination-decision"), {"employee_id": employee.pk})
        self.assertEqual(response.context["form"]["contract_no"].value(), "SYN-HĐ")
        self.assertEqual(response.context["form"]["contract_date"].value(), "")
        response = self.client.get(self.url("employee-information"), {"employee_id": employee.pk})
        self.assertEqual(response.context["form"]["account_number"].value(), "")
        EmployeeSalary.objects.create(employee=employee, current_basic_salary=Decimal("25000000"), effective_date=date(2021, 1, 1))
        response = self.client.get(self.url("salary-decision"), {"employee_id": employee.pk})
        self.assertEqual(response.context["form"]["salary"].value(), "")
