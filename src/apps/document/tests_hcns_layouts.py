import io
from datetime import date
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.template.loader import render_to_string
from django.test import SimpleTestCase
from openpyxl import load_workbook

from .hcns.catalog import FORMS
from .hcns.exports import W, export_docx, export_xlsx
from .hcns.forms import initial_document_data
from .hcns.rendering import document_context


class PrintLayoutTests(SimpleTestCase):
    def context(self, slug, **overrides):
        spec = FORMS[slug]
        values, rows = initial_document_data(spec, {}, {})
        values.update(overrides)
        return document_context(spec, values, rows, "SOURCE-METADATA-ONLY")

    def test_all_word_page_sizes_and_margins_match_reference_profiles(self):
        for slug, spec in FORMS.items():
            if spec["kind"] != "docx":
                continue
            with self.subTest(slug=slug):
                context = self.context(slug)
                with ZipFile(io.BytesIO(export_docx(context))) as archive:
                    root = ET.fromstring(archive.read("word/document.xml"))
                size = root.find(f".//{{{W}}}pgSz")
                margins = root.find(f".//{{{W}}}pgMar")
                for key in ("top", "bottom", "left", "right", "header", "footer"):
                    self.assertEqual(int(margins.get(f"{{{W}}}{key}")), context["profile"][key])
                self.assertEqual(int(size.get(f"{{{W}}}w")), context["profile"]["width"])
                self.assertNotIn("SOURCE-METADATA-ONLY", "".join(root.itertext()))

    def test_all_hr_decisions_have_two_column_header_and_no_editor_headings(self):
        for slug, spec in FORMS.items():
            if spec["group"] != "decision":
                continue
            with self.subTest(slug=slug):
                context = self.context(slug, full_name="Nhân viên giả", decision_no="TEST-DECISION", document_date=date(2026, 10, 6))
                html = render_to_string("document/hcns/document.html", {"document": context})
                self.assertIn('class="hcns-decision-head"', html)
                self.assertIn("TEST-DECISION", html)
                self.assertIn("ngày 06 tháng 10 năm 2026", html)
                for label in ("Nội dung căn cứ", "Danh sách nơi nhận", "Hồ sơ nguồn", "Số quyết định:"):
                    self.assertNotIn(label, html)

    def test_two_column_checklist_keeps_tracking_data_and_leave_uses_date_line(self):
        spec = FORMS["onboarding-documents"]
        values, rows = initial_document_data(spec, {}, {})
        rows["documents"][0].update(received="Có", received_date=date(2026, 10, 6), note="Đã kiểm tra")
        context = document_context(spec, values, rows)
        table = next(b["table"] for b in context["layout"]["blocks"] if b.get("table"))
        self.assertEqual(len(table["columns"]), 2)
        self.assertIn("Đã kiểm tra", table["rows"][0][1])
        self.assertIn("06/10/2026", table["rows"][0][1])
        html = render_to_string("document/hcns/document.html", {"document": self.context("annual-leave")})
        self.assertIn("Trong thời gian xin nghỉ tôi đã bàn giao", html)
        self.assertNotIn("Nội dung đề nghị", html)
        self.assertNotIn("Ngày lập:", html)

    def test_workbook_group_headers_and_print_ranges(self):
        for slug in ("salary-scale", "candidate-score-summary"):
            with self.subTest(slug=slug):
                ws = load_workbook(io.BytesIO(export_xlsx(self.context(slug)))).active
                self.assertTrue(ws.print_area)
                self.assertTrue(ws.print_title_rows)
                expected = "Bậc" if slug == "salary-scale" else "Tiêu chí và kết quả phỏng vấn"
                self.assertIn(expected, [c.value for row in ws for c in row])
                self.assertNotIn("SOURCE-METADATA-ONLY", [c.value for row in ws for c in row])
        ws = load_workbook(io.BytesIO(export_xlsx(self.context("offboarding-handover")))).active
        self.assertGreaterEqual(ws.column_dimensions["H"].width, 14)
        self.assertEqual(ws["H11"].value, "Ghi chú")
        inventory = load_workbook(io.BytesIO(export_xlsx(self.context("asset-inventory")))).active
        self.assertIn("A9:N9", [str(r) for r in inventory.merged_cells.ranges])

    def test_rating_grid_uses_grouped_repeatable_headers_and_keeps_selected_levels(self):
        spec = FORMS["employee-evaluation"]
        values, rows = initial_document_data(spec, {}, {})
        rows["criteria"][0].update(frequency="2", quality="3")
        context = document_context(spec, values, rows)
        table = next(b["table"] for b in context["layout"]["blocks"] if b.get("table") and b["table"].get("key") == "criteria")
        self.assertEqual(table["rows"][1][2:], ["", "✓", "", "", "", "✓"])
        html = render_to_string("document/hcns/document.html", {"document": context})
        self.assertIn('width:52.0000%', html)
        with ZipFile(io.BytesIO(export_docx(context))) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        self.assertEqual(len(root.findall(f".//{{{W}}}gridSpan")), 2)
        self.assertIn("Tần suất", "".join(root.itertext()))
        self.assertIn("Chất lượng", "".join(root.itertext()))

    def test_asset_decision_keeps_inventory_between_articles(self):
        context = self.context("asset-disposal-decision")
        with ZipFile(io.BytesIO(export_docx(context))) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        body = root.find(f"{{{W}}}body")
        blocks = ["".join(el.itertext()) for el in body]
        index = next(i for i, text in enumerate(blocks) if text.startswith("Điều 1:"))
        self.assertIn("Tên thiết bị", blocks[index+1])
        self.assertTrue(blocks[index+2].startswith("Điều 2:"))
        proposal = self.context("asset-disposal-proposal")["layout"]
        minutes = self.context("asset-destruction-minutes")["layout"]
        self.assertEqual(len(proposal["signatures"]), 5)
        self.assertEqual(len(minutes["signatures"]), 4)
        self.assertEqual(minutes["signature_columns"], 2)
