"""Portable DOCX and macro-free XLSX exports from the validated layout data."""
import io
from datetime import date, time
from decimal import Decimal
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .rendering import INTERVIEW_KEYS, TEST_KEYS, display


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ET.register_namespace("w", W)


def _w(parent, tag, **attrs):
    return ET.SubElement(parent, "{" + W + "}" + tag, {"{" + W + "}" + key: str(value) for key, value in attrs.items()})


def _paragraph(parent, text, *, bold=False, centered=False, size=24, keep=False, italic=False, justified=False):
    p = _w(parent, "p")
    props = _w(p, "pPr")
    _w(props, "spacing", after=70, line=240, lineRule="auto")
    if centered:
        _w(props, "jc", val="center")
    elif justified:
        _w(props, "jc", val="both")
    if keep:
        _w(props, "keepNext")
    for i, line in enumerate(str(text).split("\n")):
        run = _w(p, "r")
        rpr = _w(run, "rPr")
        _w(rpr, "rFonts", ascii="Times New Roman", hAnsi="Times New Roman", cs="Times New Roman")
        _w(rpr, "sz", val=size)
        _w(rpr, "color", val="000000")
        if bold:
            _w(rpr, "b")
        if italic:
            _w(rpr, "i")
        if i:
            _w(run, "br")
        node = _w(run, "t")
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        node.text = display(line, placeholder=False)
    return p


def _word_table(parent, headers, rows, width, groups=None, font_size=22):
    table = _w(parent, "tbl")
    props = _w(table, "tblPr")
    _w(props, "tblW", w=width, type="dxa")
    borders = _w(props, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _w(borders, edge, val="single", sz=4, color="999999")
    _w(props, "tblLayout", type="fixed")
    margins = _w(props, "tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        _w(margins, edge, w=60, type="dxa")
    grid = _w(table, "tblGrid")
    weights = [max(7, min(70, max([len(label)] + [len(str(row[i])) for row in rows[:20]]))) for i, label in enumerate(headers)]
    if len(headers) == 8 and headers[0] == "TT":
        weights = [6, 52, 7, 7, 7, 7, 7, 7]
    elif headers == ["STT", "Loại hồ sơ"]:
        weights = [8, 92]
    widths = [int(width * weight / sum(weights)) for weight in weights]
    for cell_width in widths:
        _w(grid, "gridCol", w=cell_width)
    if groups:
        # Two repeatable header rows preserve source column groups without
        # forcing their labels into individual narrow rating cells.
        row = _w(table, "tr")
        _w(_w(row, "trPr"), "tblHeader")
        index = 0
        for group in groups:
            span = group.get("colspan", 1)
            cell = _w(row, "tc")
            props = _w(cell, "tcPr")
            _w(props, "tcW", w=sum(widths[index:index+span]), type="dxa")
            if span > 1:
                _w(props, "gridSpan", val=span)
            if group.get("rowspan") == 2:
                _w(props, "vMerge", val="restart")
            _paragraph(cell, group["label"], bold=True, size=font_size, centered=True)
            index += span
        row = _w(table, "tr")
        _w(_w(row, "trPr"), "tblHeader")
        index = 0
        for group in groups:
            span = group.get("colspan", 1)
            for offset in range(span):
                cell = _w(row, "tc")
                props = _w(cell, "tcPr")
                _w(props, "tcW", w=widths[index], type="dxa")
                if group.get("rowspan") == 2:
                    _w(props, "vMerge")
                _paragraph(cell, "" if group.get("rowspan") == 2 else headers[index], bold=True, size=font_size, centered=True)
                index += 1
    for row_index, values in enumerate([headers] + rows):
        if groups and row_index == 0:
            continue
        row = _w(table, "tr")
        row_props = _w(row, "trPr")
        _w(row_props, "cantSplit")
        if row_index == 0:
            _w(row_props, "tblHeader")
        for i, value in enumerate(values):
            cell = _w(row, "tc")
            cell_props = _w(cell, "tcPr")
            _w(cell_props, "tcW", w=widths[i], type="dxa")
            if row_index == 0:
                _w(cell_props, "shd", fill="EEEEEE")
            _paragraph(cell, value, bold=row_index == 0, size=font_size)
    return table


def export_docx(context):
    if context.get("appointment"):
        return _appointment_docx(context)
    root = ET.Element("{" + W + "}document")
    body = _w(root, "body")
    landscape = context["spec"]["orientation"] == "landscape"
    profile = context["profile"]
    width = profile["width"] - profile["left"] - profile["right"]
    for line in context["header_lines"]:
        p = _paragraph(body, line, centered=context["layout"]["header_kind"] != "right", bold=True, size=22)
        if context["layout"]["header_kind"] == "right":
            props = p.find("{" + W + "}pPr")
            _w(props, "ind", left=width*4//10)
            _w(props, "jc", val="center")
    _paragraph(body, context["layout"]["title"], centered=True, bold=True, size=28, keep=True)
    _paragraph(body, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", centered=True, size=18)
    for block in context["layout"]["blocks"]:
        if block.get("table"):
            table = block["table"]
            _word_table(body, [c["label"] for c in table["columns"]], table["rows"], width, table.get("header_groups"), profile["font_size"])
        else:
            p = _paragraph(body, block["text"], size=profile["font_size"], bold=block["heading"], keep=block["heading"], centered=block["centered"])
            if block["right"]:
                _w(p.find("{" + W + "}pPr"), "jc", val="right")
    signatures = context["layout"]["signatures"]
    if signatures:
        # A single non-splitting row keeps the signature labels and signing
        # space together. Keep the final paragraph with it to avoid orphan pages.
        previous = list(body)[-1]
        if previous.tag == "{" + W + "}p":
            _w(previous.find("{" + W + "}pPr"), "keepNext")
        table = _w(body, "tbl")
        props = _w(table, "tblPr")
        _w(props, "tblW", w=width, type="dxa")
        columns = context["layout"]["signature_columns"]
        _w(props, "tblLayout", type="fixed")
        grid = _w(table, "tblGrid")
        for _ in range(columns):
            _w(grid, "gridCol", w=width//columns)
        for start in range(0, len(signatures), columns):
            row = _w(table, "tr")
            _w(_w(row, "trPr"), "cantSplit")
            for signature in signatures[start:start+columns]:
                cell = _w(row, "tc")
                _w(_w(cell, "tcPr"), "tcW", w=width//columns, type="dxa")
                _paragraph(cell, signature, bold=True, centered=True, size=22)
                _paragraph(cell, "(Ký, ghi rõ họ tên)\n\n\n", centered=True, size=20)
    sect = _w(body, "sectPr")
    _page_properties(sect, profile)
    return _pack_docx(root)


def _pack_docx(root):
    content_types = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", ET.tostring(root, encoding="utf-8", xml_declaration=True))
    return stream.getvalue()


def _page_properties(sect, profile):
    _w(sect, "pgSz", w=profile["width"], h=profile["height"], orient="landscape" if profile["width"] > profile["height"] else "portrait")
    _w(sect, "pgMar", **{k: profile[k] for k in ("top", "bottom", "left", "right", "header", "footer")})


def _two_cells(body, widths):
    table = _w(body, "tbl")
    props = _w(table, "tblPr")
    _w(props, "tblW", w=sum(widths), type="dxa")
    _w(props, "tblLayout", type="fixed")
    grid = _w(table, "tblGrid")
    row = _w(table, "tr")
    _w(_w(row, "trPr"), "cantSplit")
    cells = []
    for width in widths:
        _w(grid, "gridCol", w=width)
        cell = _w(row, "tc")
        _w(_w(cell, "tcPr"), "tcW", w=width, type="dxa")
        cells.append(cell)
    return cells


def _appointment_docx(context):
    decision = context["appointment"]
    root = ET.Element("{" + W + "}document")
    body = _w(root, "body")
    profile = context["profile"]
    width = profile["width"] - profile["left"] - profile["right"]
    left, right = _two_cells(body, (round(width*.44), width-round(width*.44)))
    _paragraph(left, "LIÊN HIỆP CÁC HỘI KH & KT VIỆT NAM", centered=True, size=18)
    _paragraph(left, "TRUNG TÂM PHÁT TRIỂN\nSỨC KHỎE BỀN VỮNG", centered=True, bold=True, size=20)
    _paragraph(left, "─────", centered=True, size=18)
    _paragraph(left, "Số: " + decision["number"], centered=True, size=22)
    _paragraph(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", centered=True, bold=True, size=20)
    _paragraph(right, "Độc lập - Tự do - Hạnh phúc", centered=True, bold=True, size=20)
    _paragraph(right, "───────────────", centered=True, size=18)
    _paragraph(right, decision["date_line"], centered=True, italic=True, size=22)
    _paragraph(body, "", size=14)
    _paragraph(body, "QUYẾT ĐỊNH", centered=True, bold=True, size=32, keep=True)
    _paragraph(body, "V/v: " + decision["subject"], centered=True, bold=True, size=28, keep=True)
    _paragraph(body, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", centered=True, size=16)
    _paragraph(body, "GIÁM ĐỐC TRUNG TÂM", centered=True, bold=True, keep=True)
    for line in decision["bases"]:
        p = _paragraph(body, line, size=profile["font_size"], justified=True, italic=True)
        _w(p.find("{" + W + "}pPr"), "ind", firstLine=680)
    _paragraph(body, "QUYẾT ĐỊNH", centered=True, bold=True, keep=True)
    for article in decision["articles"]:
        p = _paragraph(body, f'Điều {article["number"]}:', size=profile["font_size"], bold=True, keep=article["number"] == len(decision["articles"]))
        _w(list(p)[1], "tab")
        text_p = _paragraph(body, article["text"], size=profile["font_size"], justified=True)
        _w(p.find("{" + W + "}pPr"), "jc", val="both")
        props = p.find("{" + W + "}pPr")
        _w(props, "ind", left=900, hanging=900)
        _w(_w(props, "tabs"), "tab", val="left", pos=900)
        for run in list(text_p)[1:]:
            p.append(run)
        body.remove(text_p)
        if article.get("table"):
            table = article["table"]
            _word_table(body, [c["label"] for c in table["columns"]], table["rows"], width)
    left, right = _two_cells(body, (width//2, width-width//2))
    _paragraph(left, "", size=14)
    _paragraph(left, "Nơi nhận:", bold=True, italic=True, size=20)
    for line in decision["recipients"]:
        _paragraph(left, line, size=20)
    _paragraph(right, "GIÁM ĐỐC", centered=True, bold=True)
    _paragraph(right, "\n\n\n", centered=True)
    _paragraph(right, decision["signer"], centered=True, bold=True)
    sect = _w(body, "sectPr")
    _page_properties(sect, profile)
    return _pack_docx(root)


def _cell(ws, row, column, value, *, header=False, formula=False):
    cell = ws.cell(row, column)
    if isinstance(value, Decimal):
        cell.value = value
        cell.number_format = "#,##0.####"
    elif isinstance(value, (date, time)):
        cell.value = value
        cell.number_format = "dd/mm/yyyy" if isinstance(value, date) else "hh:mm"
    elif isinstance(value, (int, float)):
        cell.value = value
        cell.number_format = "#,##0.####"
    else:
        cell.value = display(value, placeholder=False)
        # Prevent Excel formula injection without changing visible text.
        if not formula:
            cell.data_type = "s"
    cell.font = Font(name="Arial", size=10, bold=header)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if header:
        cell.fill = PatternFill("solid", fgColor="E7EDF3")
    return cell


def _sheet_table(ws, row, table):
    columns, records = table["columns"], table["raw_rows"]
    keys = [col["key"] for col in columns]
    if table["key"] == "results":
        _cell(ws, row, 7, "Tiêu chí và kết quả phỏng vấn", header=True)
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=len(columns))
        for i, col in enumerate(columns[:6], 1):
            _cell(ws, row, i, col["label"], header=True)
            ws.merge_cells(start_row=row, start_column=i, end_row=row+1, end_column=i)
        row += 1
    elif table["key"] == "scale":
        for i, label in ((1, "Chức vụ"), (2, "Cấp"), (3, "Bậc")):
            _cell(ws, row, i, label, header=True)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=12)
        for i in (1, 2):
            ws.merge_cells(start_row=row, start_column=i, end_row=row+1, end_column=i)
        row += 1
    else:
        _cell(ws, row, 1, table["title"], header=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(columns))
        row += 1
    header_row = row
    for i, col in enumerate(columns, 1):
        if not ((table["key"] == "results" and i <= 6) or (table["key"] == "scale" and i <= 2)):
            _cell(ws, row, i, col["label"], header=True)
        col_width = 12 if col.get("kind") in ("decimal", "integer", "date") else max(15, min(28, len(col["label"]) + 3))
        if table["key"] == "results":
            col_width = 8 if col.get("kind") == "decimal" else 15
            if col["key"] in ("expected_salary", "proposed_salary", "candidate"):
                col_width = 18
        if col["key"] in ("item", "criterion"):
            col_width = 55
        ws.column_dimensions[get_column_letter(i)].width = col_width
    start = row + 1
    if table["key"] == "results":
        _cell(ws, start, 1, "Thang điểm", header=True)
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
        maximums = dict(zip(INTERVIEW_KEYS, (20, 20, 20, 15, 15, 10)))
        maximums.update(plan_test=25, budget_test=25, interview_total=100, test_total=50, combined=150)
        for key, maximum in maximums.items():
            _cell(ws, start, keys.index(key)+1, maximum)
        start += 1
    for offset, record in enumerate(records):
        row = start + offset
        for i, col in enumerate(columns, 1):
            _cell(ws, row, i, record.get(col["key"]))
        if table["key"] == "results":
            for inputs, total in ((INTERVIEW_KEYS, "interview_total"), (TEST_KEYS, "test_total")):
                refs = ",".join(f"{get_column_letter(keys.index(key)+1)}{row}" for key in inputs)
                _cell(ws, row, keys.index(total)+1, f'=IF(COUNT({refs})={len(inputs)},SUM({refs}),"")', formula=True)
            peers = [start + i for i, peer in enumerate(records) if peer.get("candidate_code") == record.get("candidate_code")]
            for total, average in (("interview_total", "average_interview"), ("test_total", "average_test")):
                if record.get("candidate_code"):
                    refs = ",".join(f"{get_column_letter(keys.index(total)+1)}{peer}" for peer in peers)
                    _cell(ws, row, keys.index(average)+1, f'=IF(COUNT({refs})={len(peers)},AVERAGE({refs}),"")', formula=True)
            if record.get("candidate_code"):
                left, right = (get_column_letter(keys.index(key)+1) for key in ("average_interview", "average_test"))
                _cell(ws, row, keys.index("combined")+1, f'=IF(COUNT({left}{row},{right}{row})=2,{left}{row}+{right}{row},"")', formula=True)
        for cell in ws[row]:
            edge = Side(style="thin", color="000000")
            cell.border = Border(top=edge, bottom=edge, left=edge, right=edge)
    ws.freeze_panes = f"A{start}"
    ws.print_title_rows = f"{header_row-1}:{start-1}"
    if records:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{row}"
    return max(row, header_row) + 2


def _configure_sheet(ws, landscape=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Bản soạn thảo – Trang &P / &N"
    ws.print_area = ws.dimensions


def _reference_sheet(ws, context):
    """Build each worksheet's paper layout; keep all validated source rows."""
    slug = context["spec"]["slug"]
    fields = {f["key"]: f["raw"] for s in context["sections"] for f in s["fields"]}
    tables = {s["table"]["key"]: s["table"] for s in context["sections"] if s["table"]}

    def line(row, text, last=6, column=1, bold=False):
        _cell(ws, row, column, text, header=bold)
        ws.merge_cells(start_row=row, start_column=column, end_row=row, end_column=last)

    if slug == "salary-scale":
        line(1, "THANG BẢNG LƯƠNG", 12, bold=True)
        line(2, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", 12)
        for row, key, label in ((4, "minimum_salary", "Lương tối thiểu"), (5, "step_rate", "Bậc (step)"), (6, "rounding_digits", "Làm tròn đến")):
            line(row, label, 2)
            _cell(ws, row, 3, fields.get(key))
        _cell(ws, 6, 4, "số trước dấu phẩy")
        for row, col, text in ((4, 6, "O"), (4, 7, "Cán bộ"), (4, 8, "SO"), (4, 9, "Cán bộ cao cấp"),
                               (5, 6, "M"), (5, 7, "Quản lý"), (5, 8, "L"), (5, 9, "Lãnh đạo"), (7, 12, "(VNĐ)")):
            _cell(ws, row, col, text)
        _sheet_table(ws, 8, tables["scale"])
        ws.page_setup.fitToHeight = 1
    elif slug == "candidate-score-summary":
        line(1, "BẢNG TỔNG HỢP KẾT QUẢ PHỎNG VẤN", 22, bold=True)
        line(2, "Vị trí: " + display(fields.get("position")) + " – Dự án: " + display(fields.get("project")), 22)
        line(3, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", 22)
        line(4, "Mẫu vị trí: " + display(fields.get("variant")) + " – " + display(fields.get("location")) + ", " + display(fields.get("document_date")), 22)
        _sheet_table(ws, 5, tables["results"])
    elif slug == "asset-inventory":
        ws.title = "Trang thiet bi"
        line(2, "Trung tâm Phát triển Sức khỏe Bền vững", 7)
        line(3, "VietHealth", 7)
        line(5, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", 14)
        line(6, "DANH SÁCH TRANG THIẾT BỊ VĂN PHÒNG ĐẾN NGÀY " + display(fields.get("inventory_date")), 14, bold=True)
        line(7, "Văn phòng: " + display(fields.get("location")) + " – Mã đợt: " + display(fields.get("session_no")), 14)
        source = tables["assets"]
        records = [dict(record, ordinal=i) for i, record in enumerate(source["raw_rows"], 1)]
        assets = dict(source, columns=[{"key": "ordinal", "label": "STT", "kind": "integer"}] + source["columns"], raw_rows=records)
        row = _sheet_table(ws, 9, assets)
        # Reference committee appears after the inventory, outside its grid.
        for person in tables["members"]["raw_rows"]:
            line(row, display(person.get("name")) + " – " + display(person.get("position")) + " – " + display(person.get("role")) + " – " + display(person.get("participation")), 14)
            row += 1
        ws.print_title_rows = "6:10"
    elif slug == "offboarding-handover":
        line(2, "BẢNG KIỂM BÀN GIAO NGHỈ VIỆC", 9, column=2, bold=True)
        line(3, "BẢN SOẠN THẢO – Chưa thay thế văn bản được ký / ban hành", 9, column=2)
        for row, key, label in ((4, "full_name", "Họ tên nhân viên"), (5, "department", "Phòng ban"), (6, "notice_date", "Ngày thông báo nghỉ việc")):
            line(row, label, 3, column=2)
            line(row, display(fields.get(key)), 6, column=4)
        _cell(ws, 5, 7, "Ngày bắt đầu")
        line(5, display(fields.get("join_date")), 9, column=8)
        _cell(ws, 6, 7, "Ngày làm việc cuối")
        line(6, display(fields.get("leave_date")), 9, column=8)
        line(7, "Chức vụ: " + display(fields.get("position")) + " – " + display(fields.get("location")) + ", " + display(fields.get("document_date")), 9, column=2)
        line(9, "Các mục dưới đây cần được điền đầy đủ trước khi hoàn tất thủ tục cho nghỉ việc của Người lao động:", 9, column=2)
        for col, label in ((2, "Nội dung bàn giao"), (7, "Ngày"), (8, "Ghi chú"), (9, "Xác nhận")):
            _cell(ws, 11, col, label, header=True)
        ws.merge_cells("B11:F11")
        row = 12
        for record in tables["documents"]["raw_rows"]:
            line(row, display(record.get("item")), 6, column=2)
            _cell(ws, row, 7, record.get("received_date"))
            _cell(ws, row, 8, record.get("note"))
            _cell(ws, row, 9, record.get("received"))
            row += 1
        row += 2
        line(row, "Tôi xin cam đoan đã hoàn trả toàn bộ tài sản của Trung tâm và không có nghĩa vụ về tài chính đối với Trung tâm sau khi nghỉ việc.", 9, column=2)
        row += 2
        for key, label in (("email", "Địa chỉ email"), ("address", "Địa chỉ"), ("phone", "Điện thoại di động"), ("home_phone", "Điện thoại nhà")):
            line(row, label + ":", 3, column=2)
            line(row, display(fields.get(key)), 9, column=4)
            row += 1
        line(row+1, "Nhân viên (Ký, ghi rõ họ tên)", 5, column=2)
        line(row+1, "Nhân sự tiếp nhận (Ký, ghi rõ họ tên)", 9, column=6)
        ws.print_title_rows = "11:11"
        ws.freeze_panes = "G12"
        for col in "ABCDEFGHIJ":
            ws.column_dimensions[col].width = 12
        ws.column_dimensions["H"].width = 22
    else:
        raise ValueError("Unsupported worksheet layout")

    # Source sheets use plain Office typography and black grid borders.
    for cells in ws:
        for cell in cells:
            if cell.value is not None:
                cell.font = Font(name="Times New Roman" if slug == "offboarding-handover" else "Arial", size=11 if slug == "offboarding-handover" else 10, bold=cell.font.bold)
                cell.fill = PatternFill(fill_type=None)
    ws.cell(1 if slug != "offboarding-handover" else 2, 1 if slug != "offboarding-handover" else 2).font = Font(name="Times New Roman", size=14, bold=True)


def export_xlsx(context):
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    ws = wb.active
    ws.title = "Thang bảng lương" if context["spec"]["slug"] == "salary-scale" else "Biểu mẫu"
    _reference_sheet(ws, context)
    for col, width in enumerate(context["profile"].get("column_widths", []), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    if context["spec"]["slug"] == "offboarding-handover":
        # The reference's narrow H column was a spacer. The editable export
        # uses it for notes, so reserve enough width to print their contents.
        ws.column_dimensions["H"].width = 14
    for col in range(1, max(7, ws.max_column+1)):
        letter = get_column_letter(col)
        if letter not in ws.column_dimensions:
            ws.column_dimensions[letter].width = 20
    ws.row_dimensions[1].height = 30
    # Long labels and criteria are allowed to wrap, with deliberate row heights.
    merged_widths = {}
    for merged in ws.merged_cells.ranges:
        merged_widths[(merged.min_row, merged.min_col)] = sum(ws.column_dimensions[get_column_letter(col)].width for col in range(merged.min_col, merged.max_col+1))
    for cells in ws:
        def line_count(cell):
            if cell.data_type == "f":
                return 1
            width = max(4, int(merged_widths.get((cell.row, cell.column), ws.column_dimensions[cell.column_letter].width)) - 2)
            return sum(max(1, (len(line)+width-1)//width) for line in str(cell.value or "").splitlines())
        max_lines = max((line_count(c) for c in cells if c.value is not None), default=1)
        ws.row_dimensions[cells[0].row].height = max(16 if context["spec"]["slug"] == "offboarding-handover" else 20, 15*max_lines)
    if context["spec"]["slug"] == "asset-inventory":
        inventory = next(section["table"] for section in context["sections"] if section["table"] and section["table"]["key"] == "assets")
        for title in ("Dán nhãn", "Nhãn KK"):
            labels = wb.create_sheet(title)
            for i, record in enumerate(inventory["raw_rows"], 1):
                _cell(labels, i, 1, record.get("code"), header=True)
                _cell(labels, i, 2, record.get("name"))
                _cell(labels, i, 3, record.get("location"))
                labels.row_dimensions[i].height = 42
            for col in "ABC":
                labels.column_dimensions[col].width = 30
            _configure_sheet(labels)
        guide = wb.create_sheet("Cách đọc mã")
        _cell(guide, 1, 1, "Mã thiết bị lấy nguyên từ hồ sơ tài sản.", header=True)
        _cell(guide, 2, 1, "Không suy diễn nguồn tài trợ hoặc năm mua từ mã khi chưa có quy tắc được duyệt.")
        guide.column_dimensions["A"].width = 80
        guide.row_dimensions[2].height = 42
        _configure_sheet(guide)
    _configure_sheet(ws, context["spec"]["orientation"] == "landscape")
    profile = context["profile"]
    for key in ("top", "bottom", "left", "right", "header", "footer"):
        setattr(ws.page_margins, key, profile[key]/1440)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
